"""Author regression for compound cells. Does not run an Agent/Judge or overwrite trials."""
from pathlib import Path
from copy import deepcopy
from datetime import datetime, timezone
import ast, hashlib, importlib.util, io, json, re, shutil, tempfile, unittest
from docx import Document

GT=Path(__file__).resolve().parent
spec=importlib.util.spec_from_file_location('previous_regression',GT/'validate_evaluator_v5.py')
old=importlib.util.module_from_spec(spec);spec.loader.exec_module(old)
g=old.g;FACTS=old.FACTS;TASK=old.TASK
TRIAL=Path('D:/tmp/_evaluation/20261012T031051Z')
load=old.load;sha=old.sha
FROZEN=(TRIAL/'embedded_grader.py').read_text(encoding='utf-8')

def actual():return g.rule_scores(TRIAL/'evaluation_candidate',FACTS)[0]
def fixture(edit):
    # Every on-disk mutation is confined to a newly allocated private task directory.
    with tempfile.TemporaryDirectory(prefix='.v6-regression-',dir=GT) as name:
        base=Path(name).resolve();assert base.parent==GT.resolve()
        results=base/'results';results.mkdir()
        d=Document(TRIAL/'deliverables/project_report.docx');edit(d)
        d.save(results/'project_report.docx')
        shutil.copy2(TRIAL/'deliverables/project_report.pdf',results/'project_report.pdf')
        return g.rule_scores(base,FACTS)

def assert_rules(test,s):
    for key in ['automated.c1_files','automated.c2_task_links','automated.c3_numeric_facts']:
        test.assertEqual((s[key]['status'],s[key]['score']),('pass',1),key)

class CompoundExtractionTests(unittest.TestCase):
    def test_actual_latest_delivery(self):
        s=actual();assert_rules(self,s)
        links=s['automated.c2_task_links']['evidence'][0]
        self.assertEqual(len(links['associations']),72)
        self.assertEqual(links['foreign_task_ids'],[])
        evidence=s['automated.c3_numeric_facts']['evidence'][0]
        self.assertEqual(evidence['numeric_denominator'],120)
        self.assertEqual(evidence['unparsed'],[]);self.assertEqual(evidence['differences'],[])

    def test_canonical_identifiers_and_boundaries(self):
        for raw in ['XQ-001\n统一身份接入','XQ-001统一身份接入','任务XQ-001','（ＸＱ－００１）','xq-001 / 统一身份接入']:
            self.assertEqual(g.identifier_matches(raw,'id'),['XQ-001'])
        for raw in ['XQ-0019','AXQ-001','XQ-001A','XQ-001-02','_XQ-001','XQ-CS-26']:
            self.assertEqual(g.identifier_matches(raw,'id'),[])
        self.assertEqual(g.identifier_matches('IS-001-01\n（样本01）','issue'),['IS-001-01'])
        self.assertEqual(g.identifier_matches('IS-001-010','issue'),[])

    def test_equivalent_compound_id_cells(self):
        def edit(d):
            for it in [2,3,6]:
                for row in d.tables[it].rows[1:]:
                    row.cells[0].text=row.cells[0].text.replace('\n',' / ')
            for row in d.tables[6].rows[1:]:row.cells[2].text=row.cells[2].text.replace('\n',' ')
        assert_rules(self,fixture(edit)[0])

    def test_fullwidth_and_chinese_adjacency(self):
        def edit(d):
            for it in [2,3,6]:
                for row in d.tables[it].rows[1:]:
                    row.cells[0].text=row.cells[0].text.replace('\n','').translate(str.maketrans('XQ-0123456789','ＸＱ－０１２３４５６７８９'))
        assert_rules(self,fixture(edit)[0])

    def test_id_name_field_order_and_labels(self):
        def edit(d):
            for it in [2,3]:
                t=d.tables[it];t.cell(0,0).text='任务名称 / 任务ID'
                for row in t.rows[1:]:
                    a,b=row.cells[0].text.splitlines();row.cells[0].text=b+'（'+a+'）'
        assert_rules(self,fixture(edit)[0])

    def test_actual_foreign_task_is_penalized(self):
        def edit(d):
            t=d.tables[2];r=t.add_row()
            for cell,value in zip(r.cells,['YL-001\n其他项目任务','研发','其他负责人','其他项目目标']):cell.text=value
        s,_=fixture(edit);v=s['automated.c2_task_links']
        self.assertLess(v['score'],1)
        self.assertEqual(v['evidence'][0]['foreign_task_ids'],['YL-001'])
        self.assertEqual(old.total('automated.c2_task_links',v)['task_score_cap'],.69)

    def test_wrong_issue_and_name_are_not_normalized_away(self):
        for which in ['issue','name']:
            def edit(d):
                if which=='issue':d.tables[6].cell(1,2).text='IS-002-01\n（样本01）'
                else:d.tables[2].cell(1,0).text='XQ-001\n权限矩阵'
            s,_=fixture(edit);v=s['automated.c2_task_links']
            self.assertEqual(v['status'],'partial')
            self.assertTrue(any(x['correct'] is False for x in v['evidence'][0]['associations']))
            self.assertEqual(old.total('automated.c2_task_links',v)['task_score_cap'],.69)

    def test_shared_or_unrecognized_id_requires_review(self):
        for raw in ['XQ-001 / XQ-002\n共享任务','XQ-001A\n统一身份接入']:
            def edit(d):d.tables[3].cell(1,0).text=raw
            s,_=fixture(edit)
            self.assertIsNone(s['automated.c2_task_links']['score'])
            self.assertIsNone(s['automated.c3_numeric_facts']['score'])
            self.assertEqual(s['automated.c2_task_links']['evidence'][0]['foreign_task_ids'],[])

    def test_id_annotations_keep_raw_evidence(self):
        s=actual()
        issue=next(x for x in s['automated.c2_task_links']['evidence'][0]['associations'] if x['task_id']=='XQ-001' and x['field']=='issue')
        self.assertTrue(all(x['value']=='IS-001-01' for x in issue['observations']))
        self.assertTrue(any('样本' in x['raw_value'] for x in issue['observations']))
        d=g.extract_docx(TRIAL/'deliverables/project_report.docx')
        row=next(r for r in g.rows_from_tables(d) if r.get('id')=='XQ-001' and 'done' in r)
        self.assertEqual(row['_observations']['done'][0]['raw_value'],'2,502\n2,852')

    def test_conflicting_id_columns_are_not_bound_by_first_match(self):
        d={'tables':[{'table':1,'rows':[['任务ID','对应任务编号','计划量/完成量'],['XQ-001','XQ-002','2852/2502']]}]}
        rows=g.rows_from_tables(d)
        self.assertEqual(len(rows),1)
        self.assertNotIn('id',rows[0]);self.assertTrue(rows[0]['_unparsed_association'])
        self.assertEqual(set(rows[0]['_ambiguous_id']),{'XQ-001','XQ-002'})

    def test_supported_pair_headers_and_rejected_detail_pairs(self):
        examples=[('完成量\n计划量（条）',('done','planned')),('计划量/完成量',('planned','done')),
                  ('验收\n(通过/样本)',('accepted','acceptance_total')),('验收(样本/通过)',('acceptance_total','accepted')),
                  ('通过/不通过',('accepted','rejected')),('未通过/通过',('rejected','accepted')),
                  ('任务ID\n任务名称',('id','name')),('任务(ID/名称)',('id','name')),
                  ('优先跟进问题\n(编号/确认原因)',('issue','reason'))]
        for raw,want in examples:self.assertEqual(g.header_fields(raw),want,raw)
        for raw in ['计划量(明细/补录)','计划量/补录量','完成量/完成率','完成率/样本']:
            self.assertEqual(g.header_fields(raw),(),raw)

    def test_numeric_pair_reversed_headers_and_values(self):
        def edit(d):
            t=d.tables[3];t.cell(0,3).text='计划量/完成量（条）';t.cell(0,5).text='验收(样本/通过)'
            for row in t.rows[1:]:
                a,b=row.cells[3].text.splitlines();row.cells[3].text=b+' / '+a
                a,b=row.cells[5].text.split('/');row.cells[5].text=b+'\n'+a
        assert_rules(self,fixture(edit)[0])

    def test_numeric_pair_explicit_labels_override_display_order(self):
        def edit(d):
            for row in d.tables[3].rows[1:]:
                a,b=row.cells[3].text.splitlines();row.cells[3].text='计划量：'+b+'条 / 完成量：'+a+'条'
                a,b=row.cells[5].text.split('/');row.cells[5].text='样本：'+b+' / 通过：'+a
        assert_rules(self,fixture(edit)[0])

    def test_pass_fail_counts_and_reverse_order(self):
        for reverse in [False,True]:
            def edit(d):
                t=d.tables[3];t.cell(0,5).text='不通过/通过' if reverse else '通过/不通过'
                for row in t.rows[1:]:
                    a,b=map(int,row.cells[5].text.split('/'));vals=[a,b-a]
                    if reverse:vals.reverse()
                    row.cells[5].text=' / '.join(map(str,vals))
            assert_rules(self,fixture(edit)[0])

    def test_swapped_pair_without_header_change_is_wrong(self):
        def edit(d):
            t=d.tables[3];a,b=t.cell(1,3).text.splitlines();t.cell(1,3).text=b+'\n'+a
        s,_=fixture(edit);v=s['automated.c3_numeric_facts']
        self.assertAlmostEqual(v['score'],118/120)
        self.assertEqual(len(v['evidence'][0]['confirmed_numeric_errors']),2)
        self.assertEqual(old.total('automated.c3_numeric_facts',v)['task_score_cap'],.79)

    def test_wrong_pairs_count_real_errors(self):
        for raw in ['999,999\n2,852','-1\n2,852','2502.5\n2852']:
            def edit(d):d.tables[3].cell(1,3).text=raw
            s,_=fixture(edit);v=s['automated.c3_numeric_facts']
            self.assertAlmostEqual(v['score'],119/120)
            self.assertEqual(old.total('automated.c3_numeric_facts',v)['task_score_cap'],.79)

    def test_unlabelled_or_mismatched_pairs_remain_uncertain(self):
        for raw in ['2502 2852','2502','2502/2852/99','完成量2502/完成量2852','计划量2852/2502']:
            def edit(d):d.tables[3].cell(1,3).text=raw
            s,_=fixture(edit);v=s['automated.c3_numeric_facts']
            self.assertIsNone(v['score']);self.assertEqual(v['evidence'][0]['confirmed_numeric_errors'],[])
            self.assertEqual(old.total('automated.c3_numeric_facts',v)['cap_reasons'],[])

    def test_uncertain_pair_does_not_hide_other_error(self):
        def edit(d):
            d.tables[3].cell(1,3).text='2502 2852'
            d.tables[3].cell(2,5).text='999/40'
        s,_=fixture(edit);v=s['automated.c3_numeric_facts']
        self.assertIsNone(v['score']);self.assertEqual(len(v['evidence'][0]['confirmed_numeric_errors']),1)
        self.assertEqual(old.total('automated.c3_numeric_facts',v)['task_score_cap'],.79)

    def test_conflicting_repeated_pair_is_penalized(self):
        def edit(d):
            t=d.tables[3];oldrow=[c.text for c in t.rows[1].cells];r=t.add_row()
            for c,s in zip(r.cells,oldrow):c.text=s
            r.cells[3].text='1\n2,852'
        s,_=fixture(edit);v=s['automated.c3_numeric_facts']
        self.assertAlmostEqual(v['score'],119/120)

    def test_column_row_table_order_has_no_effect(self):
        def edit(d):
            for it in [2,3,6,7,8,9]:
                t=d.tables[it];matrix=[[c.text for c in r.cells] for r in t.rows]
                matrix=[matrix[0]]+matrix[:0:-1]
                for row,values in zip(t.rows,matrix):
                    for c,value in zip(row.cells,values[::-1]):c.text=value
            body=d.element.body
            for t in reversed(d.tables):body.insert(0,t._tbl)
        assert_rules(self,fixture(edit)[0])

    def test_issue_join_uses_candidate_links_not_truth_order(self):
        d={'paragraphs':[],'images':0,'notes':{},'tables':[
          {'table':1,'rows':[['任务ID/任务名称','问题编号'],['XQ-001/统一身份接入','IS-001-01（样本01）'],['XQ-002/权限矩阵','IS-002-01']]},
          {'table':2,'rows':[['问题编号','验收(通过/样本)'],['IS-002-01','34/40'],['IS-001-01','35/40']]}
        ]}
        rows=g.rows_from_tables(d)
        self.assertEqual([(r['id'],r['accepted']) for r in rows if 'accepted' in r],[('XQ-002','34'),('XQ-001','35')])

    def test_scoring_semantic_and_native_functions_are_unchanged(self):
        names=['c2_from_review','c3_from_review','c4_from_review','c6_from_review','c7_from_review','task_score_caps','aggregate','native_probe','render_pdf']
        def funcs(code):return {n.name:ast.dump(n,include_attributes=False) for n in ast.parse(code).body if isinstance(n,ast.FunctionDef) and n.name in names}
        self.assertEqual(funcs(old.CODE),funcs(FROZEN))
        snapshot=(TRIAL/'task_snapshot.md').read_text(encoding='utf-8')
        for heading in ['Prompt','LLM Judge Rubric']:
            self.assertEqual(old.section(old.TEXT,heading),old.section(snapshot,heading))

def recheck():
    """Re-extract frozen files; reuse unchanged semantic judgments/native evidence, clearly labelled."""
    previous=load(TRIAL/'evaluation/single_run_result.json')
    scores=deepcopy(previous['checkpoints']);rules=actual()
    for k in ['automated.c1_files','automated.c2_task_links','automated.c3_numeric_facts']:scores[k]=rules[k]
    return dict(scope='Author recheck of frozen delivery using v6 extraction; unchanged v5 semantic/native evidence reused. Not a new Agent or Judge run.',
                source_run=str(TRIAL),source_review_sha256=sha(TRIAL/'judge_review.json'),
                historical_task_score=previous['task_score'],historical_uncapped_task_score=previous['uncapped_task_score'],
                historical_coverage=previous['decidable_coverage'],
                original_rule_results={k:previous['checkpoints'][k] for k in rules if k!='automated.c5_editability'},
                checkpoints=scores,**g.aggregate(scores))

def main():
    protected=[p for trial in [TRIAL,old.TRIAL] for p in [trial/'task_snapshot.md',trial/'raw_trace.jsonl',trial/'judge_review.json',trial/'evaluation/single_run_result.json',trial/'report.md',trial/'deliverables/project_report.docx',trial/'deliverables/project_report.pdf']]
    protected += [old.MAP,*list((GT.parent/'exec').iterdir())]
    before={str(p):sha(p) for p in protected}
    suite=unittest.TestSuite([unittest.defaultTestLoader.loadTestsFromTestCase(old.Tests),unittest.defaultTestLoader.loadTestsFromTestCase(CompoundExtractionTests)])
    log=io.StringIO();r=unittest.TextTestRunner(stream=log,verbosity=2).run(suite)
    unchanged=all(sha(p)==h for p,h in before.items())
    record=dict(revision=g.RUBRIC_REVISION,utc=datetime.now(timezone.utc).isoformat(),status='pass' if r.wasSuccessful() and unchanged else 'fail',
                tests_run=r.testsRun,failures=[dict(test=str(t),detail=e) for t,e in r.failures+r.errors],log=log.getvalue(),
                protected_hashes=before,protected_files_unchanged=unchanged,task_sha256=sha(TASK),
                scope='Author rule extraction/regression only; semantic criteria, penalties, weights, Prompt and mapping unchanged.',new_agent_trial='not_run',new_semantic_judge='not_run',independent_semantic_calibration='not_performed')
    if r.wasSuccessful():record['frozen_delivery_recheck']=recheck()
    (GT/'evaluator_v6_validation.json').write_text(json.dumps(record,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(log.getvalue());print(json.dumps({k:record[k] for k in ['status','tests_run','protected_files_unchanged']}))
    if 'frozen_delivery_recheck' in record:
        print(json.dumps({k:record['frozen_delivery_recheck'][k] for k in ['task_score','uncapped_task_score','decidable_coverage','group_scores','capabilities']},ensure_ascii=False))
    return 0 if record['status']=='pass' else 1
if __name__=='__main__':raise SystemExit(main())
