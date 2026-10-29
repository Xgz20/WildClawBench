from pathlib import Path
import hashlib,json,re
from docx import Document
from docx.shared import Pt

ROOT=Path(__file__).resolve().parents[2]
GT=ROOT/'gt'
TASK=ROOT.parents[3]/'tasks/general-e2e/02_Office/02_Office_task_001_multi_source_project_report.md'
def load(p): return json.loads(p.read_text(encoding='utf-8-sig'))
def save(p,data): p.write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
def replace(p,text):
    if p.runs:
        p.runs[0].text=text
        for r in p.runs[1:]: r.text=''
    else: p.add_run(text)

template=ROOT/'exec/06_报告模板.docx'
d=Document(template)
for p in d.paragraphs:
    updates={
        '范围限定与未确认信息':'范围限定',
        '保留备注限定、缺口与来源。':'按任务 ID 保留备注中的适用范围和前置条件。',
        '来源与验收说明':'来源说明',
        '附来源索引与实际检查结果。':'在注释或附录中注明对应任务/表图、来源文件、版本及具体位置，可按共用口径合并说明。',
    }
    if p.text in updates: replace(p,updates[p.text])
    if p.text=='来源说明':
        p.insert_paragraph_before('待确认事项','Heading 1')
        p.insert_paragraph_before('逐项列出缺失状态或原因，注明任务/问题 ID、来源位置和待确认责任人。')
d.save(template)

facts=load(GT/'independent_facts.json')
locations=load(GT/'source_locations.json')
ref=GT/'reference/results/project_report.docx'
d=Document(ref)
for p in d.paragraphs:
    if '指标与问题见来源索引中该任务条目' in p.text:
        replace(p,p.text.replace('指标与问题见来源索引中该任务条目','指标与问题的共同定位规则见报告末尾来源说明'))
    if p.text=='范围限定与未确认信息': replace(p,'范围限定')
    for t in facts:
        if t['qualifier'] and p.text.startswith(t['id']+'：'):
            source=next(x for x in locations if x['task_id']==t['id'] and x['source_file'].endswith('.pptx'))
            replace(p,p.text+' 来源：'+source['source_file']+'，'+source['source_locator']+'。')
        if p.text.startswith(t['id']+' 的'):
            if t['id']=='XQ-003': p.insert_paragraph_before('待确认事项','Heading 1')
            source=next(x for x in locations if x['task_id']==t['id'] and x['source_locator'].startswith('任务台账'))
            loc=source['source_locator'] if t['status'] is None else '问题记录，项目 XQ-CS-26，任务 '+t['id']+'，问题 '+t['issue']
            replace(p,p.text+' 来源：04_项目工作簿_v1.0.xlsx，'+loc+'。')
    if p.text=='来源及验收说明': replace(p,'来源说明')
    if p.text.startswith('来源逐项定位见 source_index.csv'):
        replace(p,'本报告采用的来源均为 v1.0。任务目标、安排及脚注条件已在“任务目标及前置条件”逐项标注；12 条范围限定的 PPTX 页码及备注位置见“范围限定”。以下为进展表、问题表、项目合计及两幅图示的共同数据定位规则。')

source_paragraphs=[
    '台账与正式状态：04_项目工作簿_v1.0.xlsx 的“任务台账”A5:G34，筛选项目 XQ-CS-26，以任务 ID 关联 XQ-001 至 XQ-024；据此取得名称、部门和正式状态。',
    '计划量及完成量：同一工作簿“指标明细”A5:I1444 与隐藏的“补录明细”A5:I364，分别筛选项目及任务 ID 后相加。每个任务的基础记录为 M-任务序号-001 至 048，补录为 S-任务序号-001 至 012；例如 XQ-001 对应 M-001-001 至 M-001-048 及 S-001-001 至 S-001-012。空行下的小计不计入。',
    '验收通过数与总数：同一工作簿“验收记录”A5:H1204，以项目、任务 ID 和“有效”筛选；每任务记录为 A-任务序号-001 至 040。通过数按通过结果计数，总数按有效记录计数。',
    '问题及原因：同一工作簿“问题记录”A5:I604，以项目及任务 ID 筛选，每任务记录为 IS-任务序号-01 至 20。未关闭数量按问题 ID 去重并筛选“未关闭”；优先问题和原因对应 IS-任务序号-01。安排、责任人与日期对应各部门文档附录 A 的同一任务。',
    '项目合计与业务图示：使用进展表中 24 项任务的上述指标。图 1 按责任部门汇总完成量后除以汇总计划量，包含补录；图 2 汇总各部门未关闭问题数。项目合计采用相同范围，截止日期均为 2026-09-15。',
]
for text in source_paragraphs:
    p=d.add_paragraph(text)
    p.paragraph_format.space_after=Pt(6)
d.save(ref)

m=load(GT/'checkpoint_manifest.json')
m['checkpoints']['C1'].update(evidence=['project_report.docx','project_report.pdf'],criterion='仅两种报告文件可读，五类主题有实质内容；不要求独立辅助附件')
m['checkpoints']['C2'].update(evidence=['actual DOCX tables and paragraphs','report source annotations or appendix','actual public sources'],criterion='按项目和任务 ID 正确关联；从报告内来源说明定位实际正文事实，允许按任务或共用口径合并说明')
m['checkpoints']['C4'].update(evidence=['actual candidate tables','report pending-confirmation section','final answer in raw trace when it discusses unknowns'],criterion='4项状态与4项原因在表格和待确认事项中保持未知，列明来源及待确认责任人；回答无需重复列项但不得补造')
m['checkpoints']['C7']['criterion']='同次原始轨迹证明来源、事实、表图、编辑属性、两版一致性的实际检查；不要求独立核验记录，自述和评测端检查不能作证'
m['required_deliverables']=['results/project_report.docx','results/project_report.pdf']
m['report_embedded_requirements']={'sources':'注释或附录：对应报告章节/表图或任务 ID、来源文件、版本及具体定位，可合并共用口径','unknowns':'待确认事项章节：8项缺失状态/原因、任务/问题 ID、来源及待确认责任人','self_verification':'由同次原始轨迹和最终产物判断，不要求单独核验记录'}
m['checkpoints']['C2']['semantic_thresholds']={'pass':'24 tasks correctly linked and traceable','partial':'18-23 tasks correctly linked and traceable; no systematic cross-project contamination','fail':'fewer than 18 such tasks or systematic cross-project contamination'}
m['checkpoints']['C4']['semantic_thresholds']={'pass':'all eight complete and unknown','partial':'1-4 incomplete items; no fabrication','fail':'any fabricated state/cause or 5-8 incomplete items','counting':'each item counted at most once; completeness includes location and confirmation owner'}
m['checkpoints']['C7']['state_handling']={'missing_or_incomplete_trace':'undetermined','complete_trace_confirms_no_actual_checks_including_self_claim_only':'fail','complete_trace_confirms_1_to_4_scopes':'partial','complete_trace_confirms_5_scopes_and_final_revision_recheck':'pass'}
m['packaging']['cleanup_status']='pending verification of two-output package'
rubric=TASK.read_text(encoding='utf-8-sig').split('## LLM Judge Rubric\n',1)[1].split('\n## Workspace Path',1)[0].strip()
m['rubric_sha256']=hashlib.sha256(rubric.encode()).hexdigest()
m['rubric_hash_format']='UTF-8 LLM Judge Rubric section body, LF newlines, outer whitespace stripped'
save(GT/'checkpoint_manifest.json',m)
manifest=load(GT/'material_manifest.json')
for item in manifest['files']:
    p=ROOT/'exec'/item['name']
    item.update(bytes=p.stat().st_size,sha256=hashlib.sha256(p.read_bytes()).hexdigest())
save(GT/'material_manifest.json',manifest)
print(json.dumps({'updated':'template, reference DOCX, checkpoint and material manifests','reference_paragraphs':len(d.paragraphs)},ensure_ascii=False))
