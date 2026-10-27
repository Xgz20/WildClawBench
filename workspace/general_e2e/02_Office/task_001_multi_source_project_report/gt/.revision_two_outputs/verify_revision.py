from pathlib import Path
import hashlib,json,re,shutil,subprocess,sys
from docx import Document
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parents[2]
GT=ROOT/'gt'
TMP=Path(__file__).parent
REPO=ROOT.parents[3]
TASK=REPO/'tasks/general-e2e/02_Office/02_Office_task_001_multi_source_project_report.md'
sys.path.insert(0,str(REPO))
sys.path.append('C:/Python312/Lib/site-packages')
from src.utils.task_parser import parse_task_md
import yaml

def load(p): return json.loads(p.read_text(encoding='utf-8-sig'))
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def save(p,obj): p.write_text(json.dumps(obj,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')

parsed=parse_task_md(TASK)
ns={}
exec(compile(parsed['automated_checks'],str(TASK),'exec'),ns)
facts=load(GT/'independent_facts.json')
ref=GT/'reference/results/project_report.docx'
native=TMP/'native'
if (native/'native_receipt.json').is_file():
    receipt=load(native/'native_receipt.json')
    assert receipt['input_sha256']==sha(ref) and receipt['status']=='pass'
    score=ns['result']('pass',1,'Native edit/save/reopen receipt from this revision')
else:
    score,receipt=ns['native_probe'](ref,native,180,ROOT/'exec/06_报告模板.docx')
assert score['status']=='pass',score
shutil.copy2(receipt['pdf'],GT/'reference/results/project_report.pdf')
print('WPS reference export and actual edit/save/reopen: PASS',flush=True)

template_out=TMP/'template'
proc=subprocess.run([shutil.which('pwsh.exe') or shutil.which('powershell.exe'),'-NoProfile','-ExecutionPolicy','Bypass','-File',str(native/'native_office.ps1'),'-InputPath',str(ROOT/'exec/06_报告模板.docx'),'-OutputDir',str(template_out),'-ControlPath',str(ref)],capture_output=True,timeout=90)
assert (template_out/'native_receipt.json').is_file(),(proc.returncode,proc.stdout,proc.stderr)
tr=load(template_out/'native_receipt.json')
assert tr['status']=='pass',tr
render_ref=ns['render_pdf'](GT/'reference/results/project_report.pdf',TMP/'reference_pages',150)
render_template=ns['render_pdf'](Path(tr['pdf']),TMP/'template_pages',150)
assert render_ref['status']==render_template['status']=='pass'
print('Rendered reference pages:',len(render_ref['pages']),'template:',len(render_template['pages']),flush=True)

cases=[]
candidate=TMP/'candidate'
out=candidate/'results'
out.mkdir(parents=True,exist_ok=True)
def check_files(label,docx,pdf,expected,extras=False):
    for p in out.iterdir(): p.unlink()
    if docx: shutil.copy2(ref,out/ref.name)
    if pdf=='valid': shutil.copy2(GT/'reference/results/project_report.pdf',out/'project_report.pdf')
    if pdf=='corrupt': (out/'project_report.pdf').write_bytes(b'corrupt pdf fixture')
    if extras:
        for n in ['source_index.csv','information_gaps.md','validation_record.md']:
            (out/n).write_text('irrelevant fixture',encoding='utf-8')
    scores,details=ns['rule_scores'](candidate,facts)
    assert scores['automated.c1_files']['score']==expected,(label,scores)
    assert set(details['readability'])=={'project_report.docx','project_report.pdf'}
    cases.append({'case':label,'status':'pass','c1_score':expected})
    return scores

scores=check_files('only required DOCX and PDF, no auxiliary files',True,'valid',1)
for k in ['automated.c2_task_links','automated.c3_numeric_facts','automated.c4_unknowns']:
    assert scores[k]['status']=='pass',(k,scores[k])
check_files('DOCX only',True,None,.5)
check_files('PDF only',False,'valid',.5)
check_files('both missing',False,None,0)
check_files('corrupt PDF',True,'corrupt',.5)
check_files('extra auxiliary files do not alter score',True,'valid',1,True)

fields=['sources','facts','tables_figures','editability','docx_pdf_equivalence']
trace_cases=[]
for label,available,complete,n,expected in [
    ('missing trace',False,False,0,'undetermined'),
    ('self claim without complete trace',True,False,0,'undetermined'),
    ('complete trace confirms self claim only',True,True,0,'fail'),
    ('complete trace confirms partial checks',True,True,3,'partial'),
    ('complete trace confirms all five scopes',True,True,5,'pass'),
]:
    review={'coverage':{k:i<n for i,k in enumerate(fields)}}
    actual=ns['c7_from_review'](review,available,complete)
    assert actual['status']==expected,(label,actual)
    trace_cases.append({'case':label,'status':'pass','expected_c7_state':expected})

manifest=load(GT/'checkpoint_manifest.json')
mapping=yaml.safe_load((REPO/'tools/report/data/checkpoint_capability_map7.yaml').read_text(encoding='utf-8-sig'))[manifest['task_id']]
expected={k:[v['capability']] for k,v in manifest['runtime_keys'].items() if v['capability']}
assert mapping==expected
assert set(ns['SPECS'])==set(manifest['runtime_keys'])
assert len(parsed['rubric_criteria'])==6
for key,value in ns['SPECS'].items():
    assert tuple([manifest['runtime_keys'][key]['checkpoint'],manifest['runtime_keys'][key]['within_group_weight'],manifest['runtime_keys'][key]['capability']])==value
for item in load(GT/'material_manifest.json')['files']:
    p=ROOT/'exec'/item['name']
    assert p.stat().st_size==item['bytes'] and sha(p)==item['sha256']
for item in manifest['specifications']: assert sha(Path(item['path']))==item['sha256']

document=Document(ref)
text='\n'.join(p.text for p in document.paragraphs)
pdf=PdfReader(GT/'reference/results/project_report.pdf')
pdftext='\n'.join(p.extract_text() or '' for p in pdf.pages)
assert '待确认事项' in text and '来源说明' in text
assert '待确认事项' in pdftext and '来源说明' in pdftext
for obsolete in ['source_index.csv','information_gaps.md','validation_record.md','验收说明']:
    assert obsolete not in text and obsolete not in pdftext
for item in facts:
    if item['status'] is None or item['reason'] is None:
        gap=next(p.text for p in document.paragraphs if p.text.startswith(item['id']+' 的'))
        assert '未提供' in gap and item['owner'] in gap and '04_项目工作簿' in gap
    if item['qualifier']:
        paragraph=next(p.text for p in document.paragraphs if p.text.startswith(item['id']+'：'))
        assert '05_阶段汇报_v1.0.pptx' in paragraph and '演讲者备注' in paragraph
assert len(document.tables)==2 and all(len(t.rows)==25 for t in document.tables)
print('Parser, rule regressions, C7 states, source hashes and mapping consistency: PASS',flush=True)

manifest['revision_validation']={
    'revision':'two-report-deliverables-2026-09-22',
    'task_parser':'pass','mapping_and_runtime_keys':'pass','material_hashes':'pass','specification_hashes':'pass',
    'rule_cases':cases,'c7_state_cases':trace_cases,
    'reference_rules':{k:v for k,v in scores.items() if k!='automated.c5_editability'},
    'native_editability':score,
    'office_application':{'name':receipt['engine'],'com_version':receipt['version'],'reference_pages':len(pdf.pages),'template_pages':tr['pages']},
    'reference_docx_sha256':sha(ref),'reference_pdf_sha256':sha(GT/'reference/results/project_report.pdf'),
    'template_sha256':sha(ROOT/'exec/06_报告模板.docx'),
    'visual_review':'pending full-page visual inspection',
    'scope':'Author regression and reference validation only; no actual evaluated Agent trace or independent semantic Judge verdict is claimed.',
    'independent_semantic_review':'not_ready','valid_actual_agent_trial':'not_ready',
}
# Store the native facts, not paths to temporary files that will be removed.
manifest['revision_validation']['native_editability']={
    'status':score['status'],'input_sha256':receipt['input_sha256'],'opened':receipt['opened'],
    'exported_pdf':True,'body_persisted':receipt['body_persisted'],'table_checks':receipt['table_checks'],
    'temporary_probe_artifacts':'removed after validation; this is a construction record, not Agent C7 evidence',
}
save(GT/'checkpoint_manifest.json',manifest)
print(json.dumps({'reference_pages':[p['path'] for p in render_ref['pages']],'template_pages':[p['path'] for p in render_template['pages']]},ensure_ascii=False))
