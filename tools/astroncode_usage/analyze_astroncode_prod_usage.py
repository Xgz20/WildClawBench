#!/usr/bin/env python3
"""Export and analyze AstronCode production ``agent.turn`` usage logs.

Credentials are read only from environment variables.  The fetch path is
single-threaded and uses PIT + search_after pagination so it does not create a
large scroll context or concurrent request burst on the production cluster.
"""

from __future__ import annotations

import argparse
import gzip
import hashlib
import json
import math
import os
import random
import re
import secrets
import sys
import time
import unicodedata
import zipfile
from collections import Counter, defaultdict
from collections.abc import Callable, Iterable, Iterator, Mapping, Sequence
from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, wait
from dataclasses import dataclass, field, replace
from datetime import date, datetime, timedelta, timezone
from datetime import time as datetime_time
from pathlib import Path
from threading import Lock
from typing import Any
from urllib.parse import quote, urlsplit
from zoneinfo import ZoneInfo

try:
    import pyarrow as pa
    import pyarrow.parquet as pq
    import requests
    from openpyxl import Workbook, load_workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except (
    ImportError
) as error:  # pragma: no cover - exercised only in a broken environment
    raise SystemExit(
        "缺少依赖，请先安装 requirements.txt 中的 requests、openpyxl>=3.1 和 pyarrow"
    ) from error


DEFAULT_INDEX = "oc_acode-observer_prod*"
DEFAULT_EVENT_FIELD = "event"
DEFAULT_EVENT_VALUE = "agent.turn"
DEFAULT_TIME_FIELD = "@timestamp"
DEFAULT_DAYS = 30
DEFAULT_FALLBACK_DAYS = 14
DEFAULT_MAX_RECORDS = 100_000
DEFAULT_PAGE_SIZE = 500
DEFAULT_REQUEST_INTERVAL = 0.5
DEFAULT_TIMEZONE = "Asia/Shanghai"
RAW_HEADERS = ["用户的Query", "Query长度", "时间", "所用模型", "推理强度", "任务耗时"]
EXTENDED_RAW_COLUMN_SPECS = [
    ("任务状态", "turn.status", "文本", "Agent turn 的执行状态"),
    ("Span状态", "span.status.code", "文本", "遥测 Span 状态"),
    ("事件", "event", "文本", "日志事件名称"),
    ("日志类型", "log_type", "文本", "Observer 日志类型"),
    ("主题", "topic", "文本", "日志主题"),
    ("请求模型", "gen_ai.request.model", "文本", "请求时指定的模型"),
    ("响应模型", "gen_ai.response.model", "文本", "服务实际返回的模型"),
    ("模型提供方", "gen_ai.provider.name", "文本", "模型服务提供方"),
    ("Agent名称", "gen_ai.agent.name", "文本", "执行任务的 Agent 名称"),
    ("操作名称", "gen_ai.operation.name", "文本", "生成式 AI 操作名称"),
    ("输入Token", "gen_ai.usage.input_tokens", "整数", "输入 Token 数"),
    ("输出Token", "gen_ai.usage.output_tokens", "整数", "输出 Token 数"),
    ("推理Token", "gen_ai.usage.reasoning.output_tokens", "整数", "推理输出 Token 数"),
    (
        "缓存读取Token",
        "gen_ai.usage.cache_read.input_tokens",
        "整数",
        "缓存读取的输入 Token 数",
    ),
    ("总Token", "gen_ai.usage.total_tokens", "整数", "本 turn 的总 Token 数"),
    (
        "Token估算方式",
        "acode.usage.token_estimate_method",
        "文本",
        "Token 统计或估算方式",
    ),
    ("Usage范围", "acode.usage.scope", "文本", "Token 使用量统计范围"),
    ("消息数", "turn.message_count", "整数", "本 turn 的消息数量"),
    (
        "输入图片数",
        "gen_ai.input.messages（派生）",
        "整数",
        "用户输入消息中的图片部件数量",
    ),
    (
        "模型输出",
        "gen_ai.output.messages（派生）",
        "文本",
        "模型输出文本；超长内容在 Excel 中截断",
    ),
    ("模型输出长度", "gen_ai.output.messages（派生）", "整数", "完整模型输出的字符数"),
    (
        "系统指令长度",
        "gen_ai.system_instructions（派生）",
        "整数",
        "系统指令字符数，不保存正文到 Excel",
    ),
    ("ACode Session ID", "acode.session_id", "文本", "ACode 会话标识"),
    ("Session ID", "session.id", "文本", "OpenTelemetry 会话标识"),
    ("对话ID", "gen_ai.conversation.id", "文本", "生成式 AI 对话标识"),
    ("Agent ID", "acode.agent_id", "文本", "Agent 实例标识"),
    ("Turn ID", "acode.turn_id", "文本", "Turn 标识"),
    ("Turn序号", "acode.turn_index", "整数", "Turn 在会话中的序号"),
    ("事件序号", "acode.event_sequence_no", "整数", "事件在遥测序列中的序号"),
    ("用户ID", "acode.uid", "文本", "用户标识；仅限受控本地分析"),
    ("用户角色", "user.role", "文本", "用户角色"),
    ("Trace ID", "trace_id", "文本", "链路追踪标识"),
    ("Span ID", "span_id", "文本", "Span 标识"),
    ("Originator", "acode.originator", "文本", "调用来源"),
    ("权限模式", "acode.permission_mode", "文本", "Agent 权限模式"),
    ("工作目录", "acode.cwd", "文本", "任务工作目录；可能含敏感路径"),
    ("运行时语言", "acode.runtime_tags.language", "文本", "工作区主要语言标签"),
    ("运行时库", "acode.runtime_tags.library", "文本", "工作区主要库标签"),
    ("桌面版本", "astron.desktop.version", "文本", "Astron 桌面端版本"),
    ("CLI版本", "acode.cli_version", "文本", "AstronCode CLI 版本"),
    ("OS平台", "astron.os.platform", "文本", "操作系统平台"),
    ("OS版本", "astron.os.version", "文本", "操作系统版本"),
    ("主机架构", "astron.host.arch", "文本", "CPU/主机架构"),
    ("主机名", "host.name", "文本", "客户端主机名；仅限受控本地分析"),
    ("服务名", "service_name", "文本", "遥测服务名称"),
    ("服务实例ID", "service.instance.id", "文本", "服务实例标识"),
    ("日志等级", "severity_text", "文本", "日志等级文本"),
    ("日志等级数值", "severity_number", "整数", "OpenTelemetry 日志等级数值"),
    ("Span类型", "acode.span_type", "文本", "ACode Span 类型"),
    ("Span种类", "span.kind", "文本", "OpenTelemetry Span kind"),
    ("TLS应用类型", "tls.app.type", "文本", "TLS 应用类型标签"),
    ("Transcript路径", "acode.transcript_path", "文本", "本地 transcript 路径"),
    ("采集时间", "observed_time_unix_nano", "时间", "Observer 采集时间"),
    ("Turn耗时毫秒", "turn.duration_ms", "数值", "原始 turn.duration_ms"),
    ("Span耗时微秒", "span.duration_us", "数值", "原始 span.duration_us"),
    ("原始索引", "_index", "文本", "Elasticsearch 索引名称"),
    ("原始文档ID", "_id", "文本", "Elasticsearch 文档标识"),
    (
        "去重键",
        "Trace ID + Span ID；回退 _index + _id",
        "文本",
        "Excel 去重使用的稳定键",
    ),
]
RAW_ALL_HEADERS = RAW_HEADERS + [spec[0] for spec in EXTENDED_RAW_COLUMN_SPECS]
RAW_COLUMN_SPECS = [
    ("用户的Query", "gen_ai.input.messages（派生）", "文本", "最后一条用户消息文本"),
    ("Query长度", "用户的Query（派生）", "整数", "完整 Query 字符数"),
    ("时间", "@timestamp", "时间", "事件时间，按配置时区展示"),
    ("所用模型", "请求模型；缺失时响应模型", "文本", "本 turn 使用的模型"),
    (
        "推理强度",
        "gen_ai.request.reasoning_effort 等",
        "文本",
        "Observer 未记录时为“未记录”",
    ),
    (
        "任务耗时",
        "turn.duration_ms 等（派生）",
        "秒",
        "任务耗时，优先使用 turn.duration_ms",
    ),
] + EXTENDED_RAW_COLUMN_SPECS
SNAPSHOT_SCHEMA_VERSION = "astroncode-agent-turn-snapshot-v1"
PARQUET_SCHEMA_VERSION = "astroncode-agent-turn-normalized-v1"
SCENE_CLASSIFICATION_SCHEMA_VERSION = "astroncode-scene-classifications-v1"
HYBRID_CLASSIFICATION_VERSION = "astroncode-scene-hybrid-v1"
REFINED_CLASSIFICATION_VERSION = "astroncode-scene-hybrid-v2"
SCENE_PROMPT_VERSION = "astroncode-scene-semantic-v1"
SECONDARY_SCENE_PROMPT_VERSION = "astroncode-scene-semantic-v2-secondary"
SECONDARY_SELECTOR_VERSION = "astroncode-scene-secondary-selector-v1"
SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_CATEGORY_RULES = SCRIPT_DIR / "scene_categories.json"
DEFAULT_OUTPUT_DIR = Path("outputs") / "astroncode_usage"
DEFAULT_SCENE_MODEL = "xopglm52"
DEFAULT_SCENE_BATCH_SIZE = 40
DEFAULT_SCENE_BATCH_CHAR_LIMIT = 24_000
DEFAULT_SCENE_QUERY_CHAR_LIMIT = 2_000
DEFAULT_SCENE_WORKERS = 2
DEFAULT_SCENE_REQUEST_INTERVAL = 0.5
DEFAULT_SCENE_CONFIDENCE_THRESHOLD = 0.70
DEFAULT_SECONDARY_MIN_QUERY_LENGTH = 8
LEGACY_FALLBACK_CATEGORY = "未分类（规则未命中）"
EXCEL_CELL_LIMIT = 32_767
RETRYABLE_STATUS = {429, 502, 503, 504}
SPLITTABLE_SEMANTIC_STATUS = {400, 413, 422}
MAX_SINGLE_QUERY_REJECTIONS = 10
PARQUET_BATCH_SIZE = 2_000


@dataclass(frozen=True)
class TimeWindow:
    start: datetime
    end: datetime

    def __post_init__(self) -> None:
        if self.start.tzinfo is None or self.end.tzinfo is None:
            raise ValueError("时间窗口必须包含时区")
        if self.start >= self.end:
            raise ValueError("起始时间必须早于截止时间")


@dataclass(frozen=True)
class UsageRow:
    query: str
    query_length: int
    occurred_at: datetime | None
    model: str
    reasoning_effort: str
    duration_seconds: float | None
    fingerprint: str
    details: Mapping[str, Any] = field(default_factory=dict)


@dataclass
class FetchStats:
    fetched_hits: int = 0
    exported_rows: int = 0
    duplicate_rows: int = 0
    missing_query: int = 0
    missing_time: int = 0
    missing_model: int = 0
    missing_reasoning_effort: int = 0
    missing_duration: int = 0
    truncated_queries: int = 0
    truncated_text_cells: int = 0


@dataclass(frozen=True)
class PitHandle:
    pit_id: str
    close_path: str
    close_body_key: str
    tiebreaker_fields: tuple[str, ...]


@dataclass(frozen=True)
class CategoryRule:
    name: str
    description: str
    patterns: tuple[re.Pattern[str], ...]


@dataclass(frozen=True)
class SceneClassification:
    query_hash: str
    category: str
    proposed_category: str
    method: str
    confidence: float | None
    reason: str
    model: str
    prompt_version: str
    rule_version: str
    query_length: int
    input_truncated: bool


@dataclass(frozen=True)
class SceneClassificationSet:
    classifications: Mapping[str, SceneClassification]
    classification_version: str
    rule_version: str
    model: str
    prompt_version: str
    confidence_threshold: float
    source_name: str
    secondary_prompt_version: str = ""
    secondary_selector_version: str = ""
    secondary_min_query_length: int | None = None
    base_classification_source: str = ""


@dataclass
class AnalysisStats:
    total_rows: int
    min_time: datetime | None
    max_time: datetime | None
    missing_query: int
    missing_model: int
    missing_reasoning_effort: int
    missing_duration: int
    duration_values: list[float]
    category_counts: Counter[str]
    length_counts: Counter[str]
    model_counts: Counter[str]
    timezone_name: str


@dataclass(frozen=True)
class DatasetPaths:
    snapshot: Path
    parquet: Path


@dataclass
class GroupMetrics:
    count: int = 0
    durations: list[float] = field(default_factory=list)
    total_tokens: list[float] = field(default_factory=list)
    users: set[str] = field(default_factory=set)
    sessions: set[str] = field(default_factory=set)
    known_status_count: int = 0
    success_count: int = 0


@dataclass
class CandidateAggregate:
    representative_query: str
    query_length: int
    category: str
    count: int = 0
    sessions: set[str] = field(default_factory=set)
    models: Counter[str] = field(default_factory=Counter)
    durations: list[float] = field(default_factory=list)
    total_tokens: list[float] = field(default_factory=list)
    first_time: datetime | None = None
    last_time: datetime | None = None
    image_task_count: int = 0
    input_image_count: int = 0
    classification_method: str = "规则"
    classification_confidence: float | None = None
    classification_model: str = ""
    classification_prompt_version: str = ""


@dataclass
class DatasetAnalysis:
    total_rows: int
    timezone_name: str
    schema_version: str
    min_time: datetime | None = None
    max_time: datetime | None = None
    unique_users: set[str] = field(default_factory=set)
    unique_sessions: set[str] = field(default_factory=set)
    known_turn_status_count: int = 0
    successful_turn_count: int = 0
    field_non_missing: Counter[str] = field(default_factory=Counter)
    duration_values: list[float] = field(default_factory=list)
    token_values: dict[str, list[float]] = field(
        default_factory=lambda: defaultdict(list)
    )
    category_metrics: dict[str, GroupMetrics] = field(
        default_factory=lambda: defaultdict(GroupMetrics)
    )
    classification_method_metrics: dict[str, GroupMetrics] = field(
        default_factory=lambda: defaultdict(GroupMetrics)
    )
    classification_confidence_values: list[float] = field(default_factory=list)
    length_metrics: dict[str, GroupMetrics] = field(
        default_factory=lambda: defaultdict(GroupMetrics)
    )
    model_metrics: dict[str, GroupMetrics] = field(
        default_factory=lambda: defaultdict(GroupMetrics)
    )
    version_metrics: dict[tuple[str, str, str], GroupMetrics] = field(
        default_factory=lambda: defaultdict(GroupMetrics)
    )
    effort_metrics: dict[str, GroupMetrics] = field(
        default_factory=lambda: defaultdict(GroupMetrics)
    )
    status_metrics: dict[str, GroupMetrics] = field(
        default_factory=lambda: defaultdict(GroupMetrics)
    )
    daily_metrics: dict[date, GroupMetrics] = field(
        default_factory=lambda: defaultdict(GroupMetrics)
    )
    candidates: dict[str, CandidateAggregate] = field(default_factory=dict)
    classification_version: str = "astroncode-scene-rules-only-v1"
    classification_rule_version: str = ""
    classification_model: str = ""
    classification_prompt_version: str = ""
    classification_confidence_threshold: float | None = None
    classification_source_name: str = ""
    classification_secondary_prompt_version: str = ""
    classification_secondary_selector_version: str = ""
    classification_secondary_min_query_length: int | None = None
    classification_base_source_name: str = ""


def parse_datetime_value(
    value: str | None,
    *,
    local_timezone: ZoneInfo,
    is_end: bool,
    now: datetime,
) -> datetime:
    if not value:
        return now if is_end else now - timedelta(days=DEFAULT_DAYS)

    text = value.strip()
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
        parsed_date = date.fromisoformat(text)
        if is_end:
            parsed_date += timedelta(days=1)
        return datetime.combine(parsed_date, datetime_time.min, tzinfo=local_timezone)

    normalized = text[:-1] + "+00:00" if text.endswith("Z") else text
    parsed = datetime.fromisoformat(normalized)
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=local_timezone)
    return parsed


def requested_window(
    start: str | None,
    end: str | None,
    timezone_name: str,
    now: datetime | None = None,
) -> TimeWindow:
    local_timezone = ZoneInfo(timezone_name)
    current = now or datetime.now(local_timezone)
    end_at = parse_datetime_value(
        end, local_timezone=local_timezone, is_end=True, now=current
    )
    start_at = parse_datetime_value(
        start, local_timezone=local_timezone, is_end=False, now=end_at
    )
    return TimeWindow(start=start_at, end=end_at)


def iso_utc(value: datetime) -> str:
    return value.astimezone(timezone.utc).isoformat().replace("+00:00", "Z")


def build_event_query(
    window: TimeWindow,
    *,
    event_field: str,
    event_value: str,
    time_field: str,
) -> dict[str, Any]:
    event_should: list[dict[str, Any]] = [
        {"term": {event_field: event_value}},
        {"match_phrase": {event_field: event_value}},
    ]
    if not event_field.endswith(".keyword"):
        event_should.insert(0, {"term": {f"{event_field}.keyword": event_value}})
    return {
        "bool": {
            "filter": [
                {
                    "range": {
                        time_field: {
                            "gte": iso_utc(window.start),
                            "lt": iso_utc(window.end),
                        }
                    }
                },
                {"bool": {"should": event_should, "minimum_should_match": 1}},
            ]
        }
    }


def choose_effective_window(
    requested: TimeWindow,
    *,
    count_records: Callable[[TimeWindow], int],
    max_records: int,
    fallback_days: int,
    auto_fallback: bool,
) -> tuple[TimeWindow, int, bool]:
    requested_count = count_records(requested)
    if requested_count <= max_records:
        return requested, requested_count, False
    if not auto_fallback:
        raise RuntimeError(
            f"时间窗口内约有 {requested_count:,} 条记录，超过上限 {max_records:,}；"
            "请缩短时间范围或显式提高 --max-records"
        )

    fallback_start = max(requested.start, requested.end - timedelta(days=fallback_days))
    fallback = TimeWindow(fallback_start, requested.end)
    fallback_count = count_records(fallback)
    if fallback_count > max_records:
        raise RuntimeError(
            f"自动缩短到最近 {fallback_days} 天后仍有约 {fallback_count:,} 条记录，"
            f"超过上限 {max_records:,}；请继续缩短时间范围或显式提高 --max-records"
        )
    return fallback, fallback_count, True


class ElasticsearchClient:
    """Small production-safe JSON client with one pooled connection."""

    def __init__(
        self,
        base_url: str,
        *,
        timeout_seconds: float,
        retries: int,
        ca_cert: Path | None,
    ) -> None:
        parsed = urlsplit(base_url)
        if parsed.scheme not in {"http", "https"} or not parsed.netloc:
            raise ValueError("ES_URL 必须是有效的 http/https 地址")
        if parsed.username or parsed.password:
            raise ValueError("不要把凭据写入 ES_URL，请使用 ES_USERNAME/ES_PASSWORD")

        self.base_url = base_url.rstrip("/")
        self.timeout_seconds = timeout_seconds
        self.retries = retries
        self.verify: bool | str = str(ca_cert) if ca_cert else True
        self.session = requests.Session()
        self.session.trust_env = False
        adapter = requests.adapters.HTTPAdapter(pool_connections=1, pool_maxsize=1)
        self.session.mount("http://", adapter)
        self.session.mount("https://", adapter)
        self.session.headers.update(
            {"Accept": "application/json", "Content-Type": "application/json"}
        )

        api_key = os.environ.get("ES_API_KEY")
        bearer_token = os.environ.get("ES_BEARER_TOKEN")
        username = os.environ.get("ES_USERNAME")
        password = os.environ.get("ES_PASSWORD")
        configured_methods = sum(
            (bool(api_key), bool(bearer_token), bool(username or password))
        )
        if configured_methods > 1:
            raise ValueError(
                "ES_API_KEY、ES_BEARER_TOKEN、ES_USERNAME/ES_PASSWORD 只能配置一种认证方式"
            )
        if api_key:
            self.session.headers["Authorization"] = f"ApiKey {api_key}"
        elif bearer_token:
            self.session.headers["Authorization"] = f"Bearer {bearer_token}"
        elif username or password:
            if not username or not password:
                raise ValueError("Basic Auth 必须同时设置 ES_USERNAME 和 ES_PASSWORD")
            self.session.auth = (username, password)

    def request_json(
        self,
        method: str,
        path: str,
        *,
        body: Mapping[str, Any] | None = None,
        params: Mapping[str, Any] | None = None,
    ) -> dict[str, Any]:
        last_error: Exception | None = None
        for attempt in range(self.retries + 1):
            try:
                response = self.session.request(
                    method,
                    f"{self.base_url}{path}",
                    json=body,
                    params=params,
                    timeout=(min(10.0, self.timeout_seconds), self.timeout_seconds),
                    verify=self.verify,
                )
                if response.status_code in RETRYABLE_STATUS and attempt < self.retries:
                    response.close()
                    self._backoff(attempt)
                    continue
                if not response.ok:
                    reason = _elasticsearch_error_reason(response)
                    raise RuntimeError(
                        f"Elasticsearch 请求失败：HTTP {response.status_code}，{reason}"
                    )
                payload = response.json()
                if not isinstance(payload, dict):
                    raise TypeError("Elasticsearch 返回的不是 JSON 对象")
                return payload
            except (requests.ConnectionError, requests.Timeout) as error:
                last_error = error
                if attempt >= self.retries:
                    break
                self._backoff(attempt)
            except requests.JSONDecodeError as error:
                raise RuntimeError("Elasticsearch 返回了无效 JSON") from error
        raise RuntimeError(f"无法连接 Elasticsearch：{last_error}") from last_error

    @staticmethod
    def _backoff(attempt: int) -> None:
        delay = min(8.0, 0.5 * (2**attempt)) + random.uniform(0.0, 0.2)
        time.sleep(delay)


def _elasticsearch_error_reason(response: requests.Response) -> str:
    try:
        payload = response.json()
    except requests.JSONDecodeError:
        return "响应体不是 JSON"
    error = payload.get("error") if isinstance(payload, dict) else None
    if isinstance(error, dict):
        root_cause = error.get("root_cause")
        if (
            isinstance(root_cause, list)
            and root_cause
            and isinstance(root_cause[0], dict)
        ):
            return str(
                root_cause[0].get("reason") or error.get("reason") or "未知错误"
            )[:500]
        return str(error.get("reason") or error.get("type") or "未知错误")[:500]
    return str(error or "未知错误")[:500]


def encoded_index(index: str) -> str:
    return quote(index, safe="*,-_")


def count_matching_records(
    client: ElasticsearchClient,
    *,
    index: str,
    query: Mapping[str, Any],
) -> int:
    payload = client.request_json(
        "POST", f"/{encoded_index(index)}/_count", body={"query": query}
    )
    count = payload.get("count")
    if not isinstance(count, int):
        raise TypeError("Elasticsearch _count 响应缺少整数 count")
    return count


def open_point_in_time(
    client: ElasticsearchClient,
    *,
    index: str,
    keep_alive: str,
) -> PitHandle:
    encoded = encoded_index(index)
    elasticsearch_error: RuntimeError | None = None
    try:
        payload = client.request_json(
            "POST",
            f"/{encoded}/_pit",
            body={},
            params={"keep_alive": keep_alive},
        )
        pit_id = payload.get("id")
        if isinstance(pit_id, str) and pit_id:
            return PitHandle(pit_id, "/_pit", "id", ("_shard_doc",))
        elasticsearch_error = RuntimeError("Elasticsearch PIT 响应缺少 id")
    except RuntimeError as error:
        elasticsearch_error = error

    try:
        payload = client.request_json(
            "POST",
            f"/{encoded}/_search/point_in_time",
            body={},
            params={"keep_alive": keep_alive},
        )
    except RuntimeError as opensearch_error:
        raise RuntimeError(
            "创建 PIT 失败；Elasticsearch 接口错误："
            f"{elasticsearch_error}；OpenSearch 兼容接口错误：{opensearch_error}"
        ) from opensearch_error
    pit_id = payload.get("pit_id") or payload.get("id")
    if not isinstance(pit_id, str) or not pit_id:
        raise RuntimeError("OpenSearch PIT 响应缺少 pit_id")
    return PitHandle(
        pit_id,
        "/_search/point_in_time",
        "pit_id",
        ("_index", "_id"),
    )


def iter_matching_hits(
    client: ElasticsearchClient,
    *,
    index: str,
    query: Mapping[str, Any],
    time_field: str,
    page_size: int,
    request_interval: float,
    pit_keep_alive: str,
    max_records: int,
) -> Iterator[Mapping[str, Any]]:
    pit_handle = open_point_in_time(
        client,
        index=index,
        keep_alive=pit_keep_alive,
    )
    pit_id = pit_handle.pit_id

    search_after: Sequence[Any] | None = None
    emitted = 0
    try:
        while True:
            body: dict[str, Any] = {
                "size": page_size,
                "track_total_hits": False,
                "query": query,
                "pit": {"id": pit_id, "keep_alive": pit_keep_alive},
                "sort": [
                    {time_field: "asc"},
                    *({field: "asc"} for field in pit_handle.tiebreaker_fields),
                ],
                "_source": True,
            }
            if search_after is not None:
                body["search_after"] = list(search_after)
            payload = client.request_json("POST", "/_search", body=body)
            if isinstance(payload.get("pit_id"), str):
                pit_id = payload["pit_id"]
            hits_container = payload.get("hits")
            hits = (
                hits_container.get("hits") if isinstance(hits_container, dict) else None
            )
            if not isinstance(hits, list):
                raise TypeError("Elasticsearch _search 响应缺少 hits.hits")
            if not hits:
                break
            for hit in hits:
                if not isinstance(hit, dict):
                    continue
                if emitted >= max_records:
                    raise RuntimeError(
                        f"实际返回记录超过安全上限 {max_records:,}，已停止且未写出不完整结果"
                    )
                emitted += 1
                yield hit
            last_sort = hits[-1].get("sort")
            if not isinstance(last_sort, list):
                raise TypeError("Elasticsearch 命中记录缺少 search_after 排序值")
            search_after = last_sort
            if len(hits) < page_size:
                break
            time.sleep(request_interval)
    finally:
        try:
            client.request_json(
                "DELETE",
                pit_handle.close_path,
                body={pit_handle.close_body_key: pit_id},
            )
        except RuntimeError as error:
            print(f"警告：关闭 PIT 失败：{error}", file=sys.stderr)


def value_at_path(source: Mapping[str, Any], path: str) -> Any:
    if path in source:
        return source[path]
    current: Any = source
    for part in path.split("."):
        if not isinstance(current, Mapping) or part not in current:
            return None
        current = current[part]
    return current


def first_present(containers: Iterable[Mapping[str, Any]], paths: Sequence[str]) -> Any:
    for container in containers:
        for path in paths:
            value = value_at_path(container, path)
            if value is not None and value != "":
                return value
    return None


def decode_json(value: Any, *, max_depth: int = 3) -> Any:
    decoded = value
    for _ in range(max_depth):
        if not isinstance(decoded, str):
            break
        text = decoded.strip()
        if not text or text[0] not in '[{"':
            break
        try:
            decoded = json.loads(text)
        except json.JSONDecodeError:
            break
    return decoded


def attribute_containers(source: Mapping[str, Any]) -> list[Mapping[str, Any]]:
    containers: list[Mapping[str, Any]] = [source]
    for path in ("Attributes", "attributes"):
        decoded = decode_json(value_at_path(source, path))
        if isinstance(decoded, Mapping):
            containers.append(decoded)
    return containers


def text_from_content(value: Any) -> str:
    decoded = decode_json(value)
    if isinstance(decoded, str):
        return decoded
    if isinstance(decoded, Mapping):
        for key in ("content", "text", "value"):
            candidate = decoded.get(key)
            if isinstance(candidate, str) and candidate:
                return candidate
        for key in ("parts", "content", "messages", "input"):
            if key in decoded:
                candidate = text_from_content(decoded[key])
                if candidate:
                    return candidate
        return ""
    if isinstance(decoded, list):
        parts = [text_from_content(item) for item in decoded]
        return "\n".join(part for part in parts if part)
    return ""


def query_from_messages(value: Any) -> str:
    decoded = decode_json(value)
    if isinstance(decoded, Mapping) and "messages" in decoded:
        decoded = decode_json(decoded["messages"])
    if isinstance(decoded, Mapping):
        decoded = [decoded]
    if not isinstance(decoded, list):
        return ""

    user_messages: list[str] = []
    for message in decoded:
        if not isinstance(message, Mapping) or message.get("role") != "user":
            continue
        content = message.get("parts", message.get("content", message.get("text")))
        text = text_from_content(content).strip()
        if text:
            user_messages.append(text)
    return user_messages[-1] if user_messages else ""


def image_count_from_messages(value: Any) -> int:
    decoded = decode_json(value)
    if isinstance(decoded, Mapping) and "messages" in decoded:
        decoded = decode_json(decoded["messages"])
    if isinstance(decoded, Mapping):
        decoded = [decoded]
    if not isinstance(decoded, list):
        return 0

    image_count = 0
    for message in decoded:
        if not isinstance(message, Mapping):
            continue
        parts = decode_json(message.get("parts", message.get("content")))
        if isinstance(parts, Mapping):
            parts = [parts]
        if not isinstance(parts, list):
            continue
        for part in parts:
            if not isinstance(part, Mapping):
                continue
            part_type = str(part.get("type") or "").lower()
            if "image" in part_type or part.get("image_url"):
                image_count += 1
    return image_count


def scalar_text(value: Any) -> str:
    decoded = decode_json(value)
    if decoded is None:
        return ""
    if isinstance(decoded, str):
        return decoded.strip()
    if isinstance(decoded, (int, float, bool)):
        return str(decoded)
    return json.dumps(decoded, ensure_ascii=False, separators=(",", ":"))


def optional_integer(value: Any) -> int | None:
    numeric = finite_nonnegative(value)
    if numeric is None:
        return None
    return int(numeric) if numeric.is_integer() else None


def parse_unix_timestamp(value: float) -> datetime | None:
    numeric = float(value)
    absolute = abs(numeric)
    if absolute >= 1e17:
        seconds = numeric / 1e9
    elif absolute >= 1e14:
        seconds = numeric / 1e6
    elif absolute >= 1e11:
        seconds = numeric / 1e3
    else:
        seconds = numeric
    try:
        return datetime.fromtimestamp(seconds, tz=timezone.utc)
    except (OverflowError, OSError, ValueError):
        return None


def parse_timestamp(value: Any, local_timezone: ZoneInfo) -> datetime | None:
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, (int, float)):
        parsed = parse_unix_timestamp(value)
    elif isinstance(value, str):
        text = value.strip()
        if not text:
            return None
        if re.fullmatch(r"-?\d+(?:\.\d+)?", text):
            parsed = parse_unix_timestamp(float(text))
        else:
            normalized = text[:-1] + "+00:00" if text.endswith("Z") else text
            try:
                parsed = datetime.fromisoformat(normalized)
            except ValueError:
                return None
    else:
        return None
    if parsed is None:
        return None
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    return parsed.astimezone(local_timezone).replace(tzinfo=None)


def finite_nonnegative(value: Any) -> float | None:
    try:
        numeric = float(value)
    except (TypeError, ValueError):
        return None
    return numeric if math.isfinite(numeric) and numeric >= 0 else None


def duration_seconds(
    source: Mapping[str, Any], containers: Sequence[Mapping[str, Any]]
) -> float | None:
    milliseconds = first_present(
        containers,
        (
            "turn.duration_ms",
            "task_duration_ms",
            "task.duration_ms",
            "duration_ms",
        ),
    )
    numeric = finite_nonnegative(milliseconds)
    if numeric is not None:
        return numeric / 1_000

    microseconds = first_present(
        containers, ("span.duration_us", "duration_us", "Duration")
    )
    numeric = finite_nonnegative(microseconds)
    if numeric is not None:
        return numeric / 1_000_000

    start = first_present(containers, ("Start", "start_time_unix_us"))
    end = first_present(containers, ("End", "end_time_unix_us"))
    start_numeric = finite_nonnegative(start)
    end_numeric = finite_nonnegative(end)
    if (
        start_numeric is not None
        and end_numeric is not None
        and end_numeric >= start_numeric
    ):
        return (end_numeric - start_numeric) / 1_000_000
    return None


def extract_usage_row(hit: Mapping[str, Any], timezone_name: str) -> UsageRow:
    source = hit.get("_source")
    if not isinstance(source, Mapping):
        source = {}
    containers = attribute_containers(source)

    direct_query = first_present(
        containers,
        ("user.query", "user_query", "query", "prompt", "input_text"),
    )
    query = text_from_content(direct_query).strip() if direct_query is not None else ""
    if not query:
        messages = first_present(
            containers,
            (
                "gen_ai.input.messages",
                "input.messages",
                "input",
                "messages",
            ),
        )
        query = query_from_messages(messages).strip()

    occurred_raw = first_present(
        containers,
        ("@timestamp", "timestamp", "start_time", "Start", "start_time_unix_us"),
    )
    occurred_at = parse_timestamp(occurred_raw, ZoneInfo(timezone_name))
    model_raw = first_present(
        containers,
        (
            "gen_ai.request.model",
            "gen_ai.response.model",
            "request.model",
            "model",
        ),
    )
    effort_raw = first_present(
        containers,
        (
            "gen_ai.request.reasoning_effort",
            "gen_ai.request.reasoning_effort_level",
            "acode.reasoning_effort",
            "turn.reasoning_effort",
            "turn.context.reasoning_effort",
            "turn.context.effort",
            "model_reasoning_effort",
            "reasoning_effort",
            "effort",
        ),
    )
    model = str(model_raw).strip() if model_raw is not None else "未记录"
    effort = str(effort_raw).strip() if effort_raw is not None else "未记录"
    if not model:
        model = "未记录"
    if not effort:
        effort = "未记录"

    trace_id = first_present(containers, ("TraceID", "trace.id", "trace_id"))
    span_id = first_present(containers, ("SpanID", "span.id", "span_id"))
    logical_id = f"{trace_id or ''}|{span_id or ''}".strip("|")
    fallback_id = f"{hit.get('_index', '')}|{hit.get('_id', '')}".strip("|")
    input_messages = first_present(
        containers,
        ("gen_ai.input.messages", "input.messages", "input", "messages"),
    )
    output_messages = first_present(
        containers,
        ("gen_ai.output.messages", "output.messages", "output"),
    )
    output_text = text_from_content(output_messages).strip()
    system_instructions = first_present(containers, ("gen_ai.system_instructions",))
    runtime_tags = decode_json(first_present(containers, ("acode.runtime_tags",)))
    if not isinstance(runtime_tags, Mapping):
        runtime_tags = {}
    observed_raw = first_present(containers, ("observed_time_unix_nano",))
    request_model = scalar_text(
        first_present(containers, ("gen_ai.request.model", "request.model"))
    )
    response_model = scalar_text(first_present(containers, ("gen_ai.response.model",)))
    details: dict[str, Any] = {
        "任务状态": scalar_text(first_present(containers, ("turn.status",))),
        "Span状态": scalar_text(first_present(containers, ("span.status.code",))),
        "事件": scalar_text(first_present(containers, ("event",))),
        "日志类型": scalar_text(first_present(containers, ("log_type",))),
        "主题": scalar_text(first_present(containers, ("topic",))),
        "请求模型": request_model,
        "响应模型": response_model,
        "模型提供方": scalar_text(first_present(containers, ("gen_ai.provider.name",))),
        "Agent名称": scalar_text(first_present(containers, ("gen_ai.agent.name",))),
        "操作名称": scalar_text(first_present(containers, ("gen_ai.operation.name",))),
        "输入Token": optional_integer(
            first_present(containers, ("gen_ai.usage.input_tokens",))
        ),
        "输出Token": optional_integer(
            first_present(containers, ("gen_ai.usage.output_tokens",))
        ),
        "推理Token": optional_integer(
            first_present(containers, ("gen_ai.usage.reasoning.output_tokens",))
        ),
        "缓存读取Token": optional_integer(
            first_present(containers, ("gen_ai.usage.cache_read.input_tokens",))
        ),
        "总Token": optional_integer(
            first_present(containers, ("gen_ai.usage.total_tokens",))
        ),
        "Token估算方式": scalar_text(
            first_present(containers, ("acode.usage.token_estimate_method",))
        ),
        "Usage范围": scalar_text(first_present(containers, ("acode.usage.scope",))),
        "消息数": optional_integer(first_present(containers, ("turn.message_count",))),
        "输入图片数": image_count_from_messages(input_messages),
        "模型输出": output_text,
        "模型输出长度": len(output_text),
        "系统指令长度": len(scalar_text(system_instructions)),
        "ACode Session ID": scalar_text(
            first_present(containers, ("acode.session_id",))
        ),
        "Session ID": scalar_text(first_present(containers, ("session.id",))),
        "对话ID": scalar_text(first_present(containers, ("gen_ai.conversation.id",))),
        "Agent ID": scalar_text(first_present(containers, ("acode.agent_id",))),
        "Turn ID": scalar_text(first_present(containers, ("acode.turn_id",))),
        "Turn序号": optional_integer(first_present(containers, ("acode.turn_index",))),
        "事件序号": optional_integer(
            first_present(containers, ("acode.event_sequence_no",))
        ),
        "用户ID": scalar_text(first_present(containers, ("acode.uid",))),
        "用户角色": scalar_text(first_present(containers, ("user.role",))),
        "Trace ID": scalar_text(trace_id),
        "Span ID": scalar_text(span_id),
        "Originator": scalar_text(first_present(containers, ("acode.originator",))),
        "权限模式": scalar_text(first_present(containers, ("acode.permission_mode",))),
        "工作目录": scalar_text(first_present(containers, ("acode.cwd",))),
        "运行时语言": scalar_text(runtime_tags.get("language")),
        "运行时库": scalar_text(runtime_tags.get("library")),
        "桌面版本": scalar_text(first_present(containers, ("astron.desktop.version",))),
        "CLI版本": scalar_text(first_present(containers, ("acode.cli_version",))),
        "OS平台": scalar_text(first_present(containers, ("astron.os.platform",))),
        "OS版本": scalar_text(first_present(containers, ("astron.os.version",))),
        "主机架构": scalar_text(first_present(containers, ("astron.host.arch",))),
        "主机名": scalar_text(first_present(containers, ("host.name",))),
        "服务名": scalar_text(first_present(containers, ("service_name",))),
        "服务实例ID": scalar_text(first_present(containers, ("service.instance.id",))),
        "日志等级": scalar_text(first_present(containers, ("severity_text",))),
        "日志等级数值": optional_integer(
            first_present(containers, ("severity_number",))
        ),
        "Span类型": scalar_text(first_present(containers, ("acode.span_type",))),
        "Span种类": scalar_text(first_present(containers, ("span.kind",))),
        "TLS应用类型": scalar_text(first_present(containers, ("tls.app.type",))),
        "Transcript路径": scalar_text(
            first_present(containers, ("acode.transcript_path",))
        ),
        "采集时间": parse_timestamp(observed_raw, ZoneInfo(timezone_name)),
        "Turn耗时毫秒": finite_nonnegative(
            first_present(containers, ("turn.duration_ms",))
        ),
        "Span耗时微秒": finite_nonnegative(
            first_present(containers, ("span.duration_us", "Duration"))
        ),
        "原始索引": scalar_text(hit.get("_index")),
        "原始文档ID": scalar_text(hit.get("_id")),
        "去重键": logical_id or fallback_id,
    }
    return UsageRow(
        query=query,
        query_length=len(query),
        occurred_at=occurred_at,
        model=model,
        reasoning_effort=effort,
        duration_seconds=duration_seconds(source, containers),
        fingerprint=logical_id or fallback_id,
        details=details,
    )


def safe_excel_text(value: str) -> tuple[str, bool]:
    cleaned = ILLEGAL_CHARACTERS_RE.sub("", value)
    if cleaned.startswith(("=", "+", "-", "@")):
        cleaned = "'" + cleaned
    if len(cleaned) <= EXCEL_CELL_LIMIT:
        return cleaned, False
    suffix = "…[Excel单元格截断]"
    return cleaned[: EXCEL_CELL_LIMIT - len(suffix)] + suffix, True


def style_table_header(cells: Iterable[Any]) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    alignment = Alignment(horizontal="center", vertical="center")
    for cell in cells:
        cell.fill = fill
        cell.font = font
        cell.alignment = alignment


def style_title(cell: Any) -> None:
    cell.font = Font(name="Arial", size=14, bold=True, color="1F1F1F")
    cell.alignment = Alignment(horizontal="left", vertical="center")


def parquet_schema() -> pa.Schema:
    fields = [
        ("query", pa.string()),
        ("query_length", pa.int64()),
        ("occurred_at", pa.timestamp("ms")),
        ("model", pa.string()),
        ("reasoning_effort", pa.string()),
        ("duration_seconds", pa.float64()),
        ("turn_status", pa.string()),
        ("span_status", pa.string()),
        ("event", pa.string()),
        ("log_type", pa.string()),
        ("topic", pa.string()),
        ("request_model", pa.string()),
        ("response_model", pa.string()),
        ("model_provider", pa.string()),
        ("agent_name", pa.string()),
        ("operation_name", pa.string()),
        ("input_tokens", pa.int64()),
        ("output_tokens", pa.int64()),
        ("reasoning_tokens", pa.int64()),
        ("cache_read_tokens", pa.int64()),
        ("total_tokens", pa.int64()),
        ("token_estimate_method", pa.string()),
        ("usage_scope", pa.string()),
        ("message_count", pa.int64()),
        ("input_image_count", pa.int64()),
        ("model_output_length", pa.int64()),
        ("system_instruction_length", pa.int64()),
        ("turn_index", pa.int64()),
        ("event_sequence_no", pa.int64()),
        ("user_role", pa.string()),
        ("originator", pa.string()),
        ("permission_mode", pa.string()),
        ("runtime_language", pa.string()),
        ("runtime_library", pa.string()),
        ("desktop_version", pa.string()),
        ("cli_version", pa.string()),
        ("os_platform", pa.string()),
        ("os_version", pa.string()),
        ("host_arch", pa.string()),
        ("service_name", pa.string()),
        ("severity_text", pa.string()),
        ("severity_number", pa.int64()),
        ("span_type", pa.string()),
        ("span_kind", pa.string()),
        ("tls_app_type", pa.string()),
        ("observed_at", pa.timestamp("ms")),
        ("turn_duration_ms", pa.float64()),
        ("span_duration_us", pa.float64()),
        ("source_index", pa.string()),
        ("user_hash", pa.string()),
        ("acode_session_hash", pa.string()),
        ("session_hash", pa.string()),
        ("conversation_hash", pa.string()),
        ("agent_hash", pa.string()),
        ("turn_hash", pa.string()),
        ("trace_hash", pa.string()),
        ("span_hash", pa.string()),
        ("host_hash", pa.string()),
        ("cwd_hash", pa.string()),
        ("service_instance_hash", pa.string()),
        ("transcript_path_hash", pa.string()),
        ("source_document_hash", pa.string()),
        ("dedupe_hash", pa.string()),
    ]
    return pa.schema(fields)


NORMALIZED_FIELD_SPECS: tuple[tuple[str, str, str], ...] = (
    ("query", "用户 Query", "保留原文，仅限受控本地分析"),
    ("query_length", "Query 长度", "Unicode 字符数"),
    ("occurred_at", "时间", "按配置时区存储的时间"),
    ("model", "所用模型", "请求模型缺失时回退响应模型"),
    ("reasoning_effort", "推理强度", "未采集时为“未记录”"),
    ("duration_seconds", "任务耗时", "秒"),
    ("turn_status", "Turn 状态", "Agent turn 状态"),
    ("span_status", "Span 状态", "OpenTelemetry Span 状态"),
    ("request_model", "请求模型", "gen_ai.request.model"),
    ("response_model", "响应模型", "gen_ai.response.model"),
    ("model_provider", "模型提供方", "gen_ai.provider.name"),
    ("input_tokens", "输入 Token", "整数，缺失不按 0 计"),
    ("output_tokens", "输出 Token", "整数，缺失不按 0 计"),
    ("reasoning_tokens", "推理 Token", "整数，缺失不按 0 计"),
    ("cache_read_tokens", "缓存读取 Token", "整数，缺失不按 0 计"),
    ("total_tokens", "总 Token", "整数，缺失不按 0 计"),
    ("message_count", "消息数", "turn.message_count"),
    ("input_image_count", "输入图片数", "从用户输入部件派生"),
    ("model_output_length", "模型输出长度", "仅保留长度，不在 Parquet 保留输出正文"),
    ("system_instruction_length", "系统指令长度", "仅保留长度"),
    ("user_hash", "用户哈希", "加盐 SHA-256"),
    ("session_hash", "Session 哈希", "加盐 SHA-256"),
    ("conversation_hash", "对话哈希", "加盐 SHA-256"),
    ("trace_hash", "Trace 哈希", "加盐 SHA-256"),
    ("span_hash", "Span 哈希", "加盐 SHA-256"),
    ("host_hash", "主机哈希", "加盐 SHA-256，不保留主机名原文"),
    ("cwd_hash", "工作目录哈希", "加盐 SHA-256，不保留路径原文"),
    ("transcript_path_hash", "Transcript 路径哈希", "加盐 SHA-256"),
    ("source_document_hash", "ES 文档哈希", "加盐 SHA-256，不保留 _id 原文"),
)


def _temporary_output_path(path: Path) -> Path:
    return path.with_name(f".{path.name}.partial-{os.getpid()}-{secrets.token_hex(4)}")


def restrict_file_permissions(path: Path) -> None:
    """Keep production-derived artifacts private on POSIX hosts."""
    try:
        path.chmod(0o600)
    except OSError as error:
        print(f"警告：无法将产物权限设为仅当前用户可读写：{error}", file=sys.stderr)


def _hash_identifier(value: Any, salt: bytes) -> str:
    text = scalar_text(value)
    if not text:
        return ""
    return hashlib.sha256(salt + b"\0" + text.encode("utf-8")).hexdigest()


def _hash_salt() -> tuple[bytes, str]:
    configured = os.environ.get("ASTRONCODE_HASH_SALT")
    if configured:
        salt = configured.encode("utf-8")
        mode = "environment-stable"
    else:
        salt = secrets.token_bytes(32)
        mode = "random-per-export"
    return salt, mode


def usage_row_to_record(row: UsageRow, salt: bytes) -> dict[str, Any]:
    details = row.details
    session_value = details.get("Session ID") or details.get("ACode Session ID")
    return {
        "query": row.query,
        "query_length": row.query_length,
        "occurred_at": row.occurred_at,
        "model": row.model,
        "reasoning_effort": row.reasoning_effort,
        "duration_seconds": row.duration_seconds,
        "turn_status": details.get("任务状态") or "",
        "span_status": details.get("Span状态") or "",
        "event": details.get("事件") or "",
        "log_type": details.get("日志类型") or "",
        "topic": details.get("主题") or "",
        "request_model": details.get("请求模型") or "",
        "response_model": details.get("响应模型") or "",
        "model_provider": details.get("模型提供方") or "",
        "agent_name": details.get("Agent名称") or "",
        "operation_name": details.get("操作名称") or "",
        "input_tokens": details.get("输入Token"),
        "output_tokens": details.get("输出Token"),
        "reasoning_tokens": details.get("推理Token"),
        "cache_read_tokens": details.get("缓存读取Token"),
        "total_tokens": details.get("总Token"),
        "token_estimate_method": details.get("Token估算方式") or "",
        "usage_scope": details.get("Usage范围") or "",
        "message_count": details.get("消息数"),
        "input_image_count": details.get("输入图片数") or 0,
        "model_output_length": details.get("模型输出长度") or 0,
        "system_instruction_length": details.get("系统指令长度") or 0,
        "turn_index": details.get("Turn序号"),
        "event_sequence_no": details.get("事件序号"),
        "user_role": details.get("用户角色") or "",
        "originator": details.get("Originator") or "",
        "permission_mode": details.get("权限模式") or "",
        "runtime_language": details.get("运行时语言") or "",
        "runtime_library": details.get("运行时库") or "",
        "desktop_version": details.get("桌面版本") or "",
        "cli_version": details.get("CLI版本") or "",
        "os_platform": details.get("OS平台") or "",
        "os_version": details.get("OS版本") or "",
        "host_arch": details.get("主机架构") or "",
        "service_name": details.get("服务名") or "",
        "severity_text": details.get("日志等级") or "",
        "severity_number": details.get("日志等级数值"),
        "span_type": details.get("Span类型") or "",
        "span_kind": details.get("Span种类") or "",
        "tls_app_type": details.get("TLS应用类型") or "",
        "observed_at": details.get("采集时间"),
        "turn_duration_ms": details.get("Turn耗时毫秒"),
        "span_duration_us": details.get("Span耗时微秒"),
        "source_index": details.get("原始索引") or "",
        "user_hash": _hash_identifier(details.get("用户ID"), salt),
        "acode_session_hash": _hash_identifier(details.get("ACode Session ID"), salt),
        "session_hash": _hash_identifier(session_value, salt),
        "conversation_hash": _hash_identifier(details.get("对话ID"), salt),
        "agent_hash": _hash_identifier(details.get("Agent ID"), salt),
        "turn_hash": _hash_identifier(details.get("Turn ID"), salt),
        "trace_hash": _hash_identifier(details.get("Trace ID"), salt),
        "span_hash": _hash_identifier(details.get("Span ID"), salt),
        "host_hash": _hash_identifier(details.get("主机名"), salt),
        "cwd_hash": _hash_identifier(details.get("工作目录"), salt),
        "service_instance_hash": _hash_identifier(details.get("服务实例ID"), salt),
        "transcript_path_hash": _hash_identifier(details.get("Transcript路径"), salt),
        "source_document_hash": _hash_identifier(details.get("原始文档ID"), salt),
        "dedupe_hash": _hash_identifier(row.fingerprint, salt),
    }


def _write_parquet_batch(
    writer: pq.ParquetWriter, records: list[dict[str, Any]], schema: pa.Schema
) -> None:
    if records:
        writer.write_table(pa.Table.from_pylist(records, schema=schema))
        records.clear()


def write_dataset(
    hits: Iterable[Mapping[str, Any]],
    paths: DatasetPaths,
    *,
    requested: TimeWindow,
    effective: TimeWindow,
    estimated_count: int,
    fallback_applied: bool,
    index: str,
    event_field: str,
    event_value: str,
    timezone_name: str,
) -> FetchStats:
    if paths.snapshot.resolve() == paths.parquet.resolve():
        raise ValueError("原始快照和 Parquet 路径不能相同")
    if not str(paths.snapshot).lower().endswith(".jsonl.gz"):
        raise ValueError("原始快照必须使用 .jsonl.gz 扩展名")
    if paths.parquet.suffix.lower() != ".parquet":
        raise ValueError("标准化数据必须使用 .parquet 扩展名")
    paths.snapshot.parent.mkdir(parents=True, exist_ok=True)
    paths.parquet.parent.mkdir(parents=True, exist_ok=True)
    snapshot_tmp = _temporary_output_path(paths.snapshot)
    parquet_tmp = _temporary_output_path(paths.parquet)
    salt, salt_mode = _hash_salt()
    salt_id = hashlib.sha256(salt).hexdigest()[:16]
    schema = parquet_schema().with_metadata(
        {
            b"schema_version": PARQUET_SCHEMA_VERSION.encode(),
            b"timezone": timezone_name.encode(),
            b"hash_mode": salt_mode.encode(),
            b"hash_salt_id": salt_id.encode(),
            b"requested_start": iso_utc(requested.start).encode(),
            b"requested_end": iso_utc(requested.end).encode(),
            b"effective_start": iso_utc(effective.start).encode(),
            b"effective_end": iso_utc(effective.end).encode(),
        }
    )
    stats = FetchStats()
    seen: set[str] = set()
    try:
        with gzip.open(
            snapshot_tmp, "wt", encoding="utf-8", compresslevel=6
        ) as snapshot:
            snapshot.write(
                json.dumps(
                    {
                        "record_type": "metadata",
                        "schema_version": SNAPSHOT_SCHEMA_VERSION,
                        "index": index,
                        "event_filter": {event_field: event_value},
                        "requested_start": iso_utc(requested.start),
                        "requested_end": iso_utc(requested.end),
                        "effective_start": iso_utc(effective.start),
                        "effective_end": iso_utc(effective.end),
                        "timezone": timezone_name,
                        "estimated_count": estimated_count,
                        "fallback_applied": fallback_applied,
                    },
                    ensure_ascii=False,
                    separators=(",", ":"),
                )
                + "\n"
            )
            parquet_writer = pq.ParquetWriter(
                parquet_tmp,
                schema,
                compression="zstd",
                use_dictionary=True,
                write_statistics=True,
            )
            batch: list[dict[str, Any]] = []
            try:
                for hit in hits:
                    stats.fetched_hits += 1
                    snapshot.write(
                        json.dumps(
                            {"record_type": "hit", "hit": hit},
                            ensure_ascii=False,
                            separators=(",", ":"),
                        )
                        + "\n"
                    )
                    row = extract_usage_row(hit, timezone_name)
                    if row.fingerprint and row.fingerprint in seen:
                        stats.duplicate_rows += 1
                        continue
                    if row.fingerprint:
                        seen.add(row.fingerprint)
                    stats.missing_query += int(not row.query)
                    stats.missing_time += int(row.occurred_at is None)
                    stats.missing_model += int(row.model == "未记录")
                    stats.missing_reasoning_effort += int(
                        row.reasoning_effort == "未记录"
                    )
                    stats.missing_duration += int(row.duration_seconds is None)
                    batch.append(usage_row_to_record(row, salt))
                    stats.exported_rows += 1
                    if len(batch) >= PARQUET_BATCH_SIZE:
                        _write_parquet_batch(parquet_writer, batch, schema)
                _write_parquet_batch(parquet_writer, batch, schema)
            finally:
                parquet_writer.close()
            snapshot.write(
                json.dumps(
                    {
                        "record_type": "footer",
                        "fetched_hits": stats.fetched_hits,
                        "normalized_rows": stats.exported_rows,
                        "duplicate_rows": stats.duplicate_rows,
                    },
                    ensure_ascii=False,
                    separators=(",", ":"),
                )
                + "\n"
            )
        verify_snapshot(snapshot_tmp, stats.fetched_hits)
        verify_parquet(parquet_tmp, stats.exported_rows)
        os.replace(snapshot_tmp, paths.snapshot)
        os.replace(parquet_tmp, paths.parquet)
        restrict_file_permissions(paths.snapshot)
        restrict_file_permissions(paths.parquet)
    except Exception:
        snapshot_tmp.unlink(missing_ok=True)
        parquet_tmp.unlink(missing_ok=True)
        raise
    return stats


def verify_snapshot(path: Path, expected_hits: int) -> None:
    hit_count = 0
    metadata_seen = False
    footer_seen = False
    with gzip.open(path, "rt", encoding="utf-8") as handle:
        for line_number, line in enumerate(handle, start=1):
            payload = json.loads(line)
            record_type = payload.get("record_type")
            if line_number == 1:
                metadata_seen = (
                    record_type == "metadata"
                    and payload.get("schema_version") == SNAPSHOT_SCHEMA_VERSION
                )
            elif record_type == "hit":
                hit_count += 1
            elif record_type == "footer":
                footer_seen = True
    if not metadata_seen or not footer_seen or hit_count != expected_hits:
        raise RuntimeError(
            "原始快照校验失败："
            f"metadata={metadata_seen}, footer={footer_seen}, "
            f"hits={hit_count}, expected={expected_hits}"
        )


def verify_parquet(path: Path, expected_rows: int) -> None:
    parquet_file = pq.ParquetFile(path)
    metadata = parquet_file.schema_arrow.metadata or {}
    version = metadata.get(b"schema_version", b"").decode()
    if version != PARQUET_SCHEMA_VERSION:
        raise RuntimeError(f"Parquet schema_version 不匹配：{version}")
    if parquet_file.metadata.num_rows != expected_rows:
        raise RuntimeError(
            f"Parquet 行数校验失败：期望 {expected_rows}，"
            f"实际 {parquet_file.metadata.num_rows}"
        )


def load_category_rules(path: Path) -> tuple[str, list[CategoryRule]]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    version = payload.get("version")
    raw_categories = payload.get("categories")
    if not isinstance(version, str) or not version:
        raise ValueError("场景分类配置缺少 version")
    if not isinstance(raw_categories, list) or not raw_categories:
        raise ValueError("场景分类配置缺少 categories")
    rules: list[CategoryRule] = []
    names: set[str] = set()
    for item in raw_categories:
        if not isinstance(item, dict):
            raise TypeError("场景分类项必须是对象")
        name = item.get("name")
        description = item.get("description")
        patterns = item.get("patterns")
        if not isinstance(name, str) or not name or name in names:
            raise ValueError(f"场景分类名称缺失或重复：{name}")
        if not isinstance(description, str) or not isinstance(patterns, list):
            raise TypeError(f"场景分类 {name} 的 description/patterns 无效")
        names.add(name)
        rules.append(
            CategoryRule(
                name=name,
                description=description,
                patterns=tuple(
                    re.compile(pattern, re.IGNORECASE) for pattern in patterns
                ),
            )
        )
    if rules[-1].patterns:
        raise ValueError("最后一个场景分类必须是不含 patterns 的兜底分类")
    return version, rules


def classify_query(query: str, rules: Sequence[CategoryRule]) -> str:
    if not query.strip():
        return rules[-1].name
    for rule in rules[:-1]:
        if any(pattern.search(query) for pattern in rule.patterns):
            return rule.name
    return rules[-1].name


LENGTH_BANDS: tuple[tuple[str, int, int | None], ...] = (
    ("0（空 Query）", 0, 0),
    ("1–20", 1, 20),
    ("21–50", 21, 50),
    ("51–100", 51, 100),
    ("101–200", 101, 200),
    ("201–500", 201, 500),
    ("501–1000", 501, 1000),
    (">1000", 1001, None),
)


def query_length_band(length: int) -> str:
    normalized = max(0, length)
    for label, lower, upper in LENGTH_BANDS:
        if normalized >= lower and (upper is None or normalized <= upper):
            return label
    raise AssertionError("Query 长度分档未覆盖输入")


def linear_percentile(values: Sequence[float], ratio: float) -> float | None:
    if not values:
        return None
    ordered = sorted(values)
    position = (len(ordered) - 1) * ratio
    lower = math.floor(position)
    upper = math.ceil(position)
    if lower == upper:
        return ordered[lower]
    weight = position - lower
    return ordered[lower] * (1 - weight) + ordered[upper] * weight


TOKEN_FIELDS: tuple[tuple[str, str], ...] = (
    ("input_tokens", "输入 Token"),
    ("output_tokens", "输出 Token"),
    ("reasoning_tokens", "推理 Token"),
    ("cache_read_tokens", "缓存读取 Token"),
    ("total_tokens", "总 Token"),
)

VERSION_DIMENSIONS: tuple[tuple[str, str, str], ...] = (
    ("AstronStudio", "astron.desktop.version", "desktop_version"),
    ("AstronCode", "acode.cli_version", "cli_version"),
)

QUALITY_FIELDS: tuple[tuple[str, str, str], ...] = (
    ("query", "Query", "空值不能用于产品候选集"),
    ("occurred_at", "时间", "用于日期趋势和时间范围"),
    ("model", "所用模型", "“未记录”按缺失计"),
    ("reasoning_effort", "推理强度", "“未记录”按缺失计"),
    ("desktop_version", "AstronStudio 版本", "astron.desktop.version"),
    ("cli_version", "AstronCode 版本", "acode.cli_version"),
    ("duration_seconds", "任务耗时", "缺失不按 0 计"),
    ("turn_status", "Turn 状态", "用于成功率口径"),
    ("input_tokens", "输入 Token", "缺失不按 0 计"),
    ("output_tokens", "输出 Token", "缺失不按 0 计"),
    ("reasoning_tokens", "推理 Token", "缺失不按 0 计"),
    ("cache_read_tokens", "缓存读取 Token", "缺失不按 0 计"),
    ("total_tokens", "总 Token", "缺失不按 0 计"),
    ("user_hash", "用户哈希", "用于去标识化用户统计"),
    ("session_hash", "Session 哈希", "用于去标识化会话统计"),
    ("input_image_count", "输入图片数", "0 是有效值"),
)

SUCCESS_STATUS_VALUES = {
    "completed",
    "complete",
    "success",
    "succeeded",
    "ok",
    "passed",
}

REDACTION_PATTERNS: tuple[tuple[str, re.Pattern[str], str], ...] = (
    (
        "带凭据URL",
        re.compile(r"(?i)\b(https?://)[^\s/@:]+:[^\s/@]+@"),
        r"\1[已脱敏凭据]@",
    ),
    (
        "密钥或口令",
        re.compile(
            r"(?i)\b(api[_-]?key|access[_-]?key|secret|password|passwd|pwd|token|authorization)\b\s*[:=]\s*([^\s,;]+)"
        ),
        r"\1=[已脱敏]",
    ),
    (
        "Bearer Token",
        re.compile(r"(?i)\bBearer\s+[A-Za-z0-9._~+\-/=]{8,}"),
        "Bearer [已脱敏]",
    ),
    ("OpenAI格式密钥", re.compile(r"\bsk-[A-Za-z0-9_-]{12,}\b"), "[已脱敏密钥]"),
    (
        "邮箱",
        re.compile(r"(?i)\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b"),
        "[已脱敏邮箱]",
    ),
    ("IPv4", re.compile(r"(?<!\d)(?:\d{1,3}\.){3}\d{1,3}(?!\d)"), "[已脱敏IP]"),
    ("中国身份证", re.compile(r"(?<!\d)\d{17}[0-9Xx](?!\d)"), "[已脱敏身份证]"),
    ("手机号", re.compile(r"(?<!\d)1[3-9]\d{9}(?!\d)"), "[已脱敏手机号]"),
    (
        "用户主目录",
        re.compile(r"(?i)(?:/Users/|/home/|[A-Z]:\\Users\\)[^/\\\s]+"),
        "[已脱敏用户目录]",
    ),
)


def normalize_query(query: str) -> tuple[str, str]:
    display = re.sub(r"\s+", " ", unicodedata.normalize("NFKC", query)).strip()
    return display, display.casefold()


def query_key_hash(query_key: str) -> str:
    return hashlib.sha256(query_key.encode("utf-8")).hexdigest()


CONTEXT_DEPENDENT_PATTERNS: tuple[re.Pattern[str], ...] = (
    re.compile(
        r"^(?:请)?(?:"
        r"(?:继续|接着)(?:一下|下|来|任务|执行|完成(?:任务)?|推进(?:剧情)?|"
        r"剧情推进|生成|修改|调整|处理|运行|检查|提交|更新|优化|补充|分析)?|"
        r"(?:下一步|往下)|"
        r"(?:重新|再次|再)(?:执行|运行|生成|创建|修改|调整|处理|测试|检查|"
        r"提交|开始|来|试|试试|试下|做|看|看下|看看)(?:一下|下|一遍|一次)?|"
        r"重试|开始|执行|改|提交|允许|跳过|断开|就这样|可以|可以的|好的|"
        r"好了|好|行|确认|同意|是的|对的|完成了吗|好了吗|怎么样|为什么|"
        r"是吗|然后呢"
        r")(?:吧|啊|呢|哦|哈|呀)?[。.!！?？~～]*$",
        re.IGNORECASE,
    ),
    re.compile(
        r"^(?:按|按照|基于|根据|沿用)(?:上面|上述|前面|前述|刚才|这个|该|"
        r"建议|要求|方案|设计).{0,24}(?:调整|修改|优化|实施|执行|处理|继续|"
        r"生成|重写|完善)?[。.!！?？]*$",
        re.IGNORECASE,
    ),
    re.compile(
        r"^(?:这个|那个|它|上面|上述|前面|前述|刚才的).{0,20}"
        r"(?:呢|吗|怎么|如何|调整|修改|修复|处理|继续|可以|行不行)[。.!！?？]*$",
        re.IGNORECASE,
    ),
)


SECONDARY_TASK_INTENT_PATTERNS: tuple[re.Pattern[str], ...] = (
    re.compile(
        r"生成|创建|新建|写|撰写|编写|制作|修改|调整|优化|修复|实现|开发|"
        r"分析|检查|检测|查看|打开|读取|查找|搜索|调研|总结|整理|翻译|解释|"
        r"介绍|列出|比较|对比|删除|移动|复制|重命名|部署|安装|配置|运行|"
        r"执行|测试|审查|评审|设计|计算|提取|导出|导入|上传|下载|转换|"
        r"改写|润色|发送|回复|回答|告诉|给出|推荐|规划|制定|更新|添加|"
        r"补充|移除|保存|发布|抓取|获取|调用|使用|建立|排查|定位|验证|"
        r"核对|解析|处理|完成|提交"
    ),
    re.compile(
        r"\b(?:create|write|build|implement|fix|debug|analy[sz]e|check|inspect|"
        r"open|read|find|search|research|summari[sz]e|translate|explain|list|"
        r"compare|delete|move|copy|rename|deploy|install|configure|run|execute|"
        r"test|review|design|generate|export|import|upload|download|convert|"
        r"update|add|remove|save|publish|fetch|get|use|make|show|help)\b",
        re.IGNORECASE,
    ),
)


def is_context_dependent_query(query: str) -> bool:
    display, _ = normalize_query(query)
    if not display or len(display) > 80:
        return False
    return any(pattern.fullmatch(display) for pattern in CONTEXT_DEPENDENT_PATTERNS)


def has_explicit_task_intent(query: str) -> bool:
    display, _ = normalize_query(query)
    return any(pattern.search(display) for pattern in SECONDARY_TASK_INTENT_PATTERNS)


def is_secondary_classification_candidate(
    query: str, *, minimum_length: int = DEFAULT_SECONDARY_MIN_QUERY_LENGTH
) -> bool:
    display, _ = normalize_query(query)
    return (
        len(display) >= minimum_length
        and not is_context_dependent_query(display)
        and has_explicit_task_intent(display)
    )


def redact_query(query: str) -> tuple[str, tuple[str, ...]]:
    redacted = query
    matches: list[str] = []
    for label, pattern, replacement in REDACTION_PATTERNS:
        redacted, count = pattern.subn(replacement, redacted)
        if count:
            matches.append(label)
    return redacted, tuple(matches)


def scene_classification_schema() -> pa.Schema:
    return pa.schema(
        [
            ("query_hash", pa.string()),
            ("category", pa.string()),
            ("proposed_category", pa.string()),
            ("classification_method", pa.string()),
            ("confidence", pa.float64()),
            ("reason", pa.string()),
            ("model", pa.string()),
            ("prompt_version", pa.string()),
            ("rule_version", pa.string()),
            ("query_length", pa.int64()),
            ("input_truncated", pa.bool_()),
        ]
    )


def classification_to_record(
    classification: SceneClassification,
) -> dict[str, Any]:
    return {
        "query_hash": classification.query_hash,
        "category": classification.category,
        "proposed_category": classification.proposed_category,
        "classification_method": classification.method,
        "confidence": classification.confidence,
        "reason": classification.reason,
        "model": classification.model,
        "prompt_version": classification.prompt_version,
        "rule_version": classification.rule_version,
        "query_length": classification.query_length,
        "input_truncated": classification.input_truncated,
    }


def classification_from_record(record: Mapping[str, Any]) -> SceneClassification:
    confidence_raw = record.get("confidence")
    confidence = None if confidence_raw is None else float(confidence_raw)
    return SceneClassification(
        query_hash=str(record["query_hash"]),
        category=str(record["category"]),
        proposed_category=str(record.get("proposed_category") or record["category"]),
        method=str(record["classification_method"]),
        confidence=confidence,
        reason=str(record.get("reason") or ""),
        model=str(record.get("model") or ""),
        prompt_version=str(record.get("prompt_version") or ""),
        rule_version=str(record.get("rule_version") or ""),
        query_length=int(record.get("query_length") or 0),
        input_truncated=bool(record.get("input_truncated")),
    )


def canonicalize_fallback_category(
    classification: SceneClassification, fallback_category: str
) -> SceneClassification:
    category = (
        fallback_category
        if classification.category == LEGACY_FALLBACK_CATEGORY
        else classification.category
    )
    proposed_category = (
        fallback_category
        if classification.proposed_category == LEGACY_FALLBACK_CATEGORY
        else classification.proposed_category
    )
    if (
        category == classification.category
        and proposed_category == classification.proposed_category
    ):
        return classification
    return replace(
        classification,
        category=category,
        proposed_category=proposed_category,
    )


def load_scene_classifications(
    path: Path, rules: Sequence[CategoryRule]
) -> SceneClassificationSet:
    if path.suffix.lower() != ".parquet":
        raise ValueError("场景分类输入必须使用 .parquet 扩展名")
    parquet_file = pq.ParquetFile(path)
    metadata = parquet_file.schema_arrow.metadata or {}
    schema_version = metadata.get(b"schema_version", b"").decode()
    if schema_version != SCENE_CLASSIFICATION_SCHEMA_VERSION:
        raise ValueError(f"场景分类 schema_version 不支持：{schema_version or '缺失'}")
    expected_fields = {field.name for field in scene_classification_schema()}
    actual_fields = set(parquet_file.schema_arrow.names)
    if not expected_fields.issubset(actual_fields):
        raise ValueError(
            f"场景分类 Parquet 缺少字段：{sorted(expected_fields - actual_fields)}"
        )
    category_names = {rule.name for rule in rules}
    fallback_category = rules[-1].name
    classifications: dict[str, SceneClassification] = {}
    for batch in parquet_file.iter_batches(batch_size=PARQUET_BATCH_SIZE):
        for record in batch.to_pylist():
            classification = canonicalize_fallback_category(
                classification_from_record(record), fallback_category
            )
            if classification.category not in category_names:
                raise ValueError(f"场景分类包含未知类别：{classification.category}")
            if classification.query_hash in classifications:
                raise ValueError(
                    f"场景分类包含重复 query_hash：{classification.query_hash}"
                )
            classifications[classification.query_hash] = classification
    secondary_min_query_length_raw = metadata.get(b"secondary_min_query_length")
    return SceneClassificationSet(
        classifications=classifications,
        classification_version=metadata.get(
            b"classification_version", HYBRID_CLASSIFICATION_VERSION.encode()
        ).decode(),
        rule_version=metadata.get(b"rule_version", b"").decode(),
        model=metadata.get(b"model", b"").decode(),
        prompt_version=metadata.get(b"prompt_version", b"").decode(),
        confidence_threshold=float(
            metadata.get(b"confidence_threshold", b"0.7").decode()
        ),
        source_name=path.name,
        secondary_prompt_version=metadata.get(
            b"secondary_prompt_version", b""
        ).decode(),
        secondary_selector_version=metadata.get(
            b"secondary_selector_version", b""
        ).decode(),
        secondary_min_query_length=(
            int(secondary_min_query_length_raw.decode())
            if secondary_min_query_length_raw
            else None
        ),
        base_classification_source=metadata.get(
            b"base_classification_source", b""
        ).decode(),
    )


class RequestRateLimiter:
    def __init__(self, interval_seconds: float) -> None:
        self.interval_seconds = max(0.0, interval_seconds)
        self.next_allowed = 0.0
        self.lock = Lock()

    def wait(self) -> None:
        with self.lock:
            now = time.monotonic()
            delay = max(0.0, self.next_allowed - now)
            self.next_allowed = max(now, self.next_allowed) + self.interval_seconds
        if delay:
            time.sleep(delay)


def semantic_system_prompt(rules: Sequence[CategoryRule]) -> str:
    category_lines = "\n".join(f"- {rule.name}: {rule.description}" for rule in rules)
    return (
        "你是 AstronCode 用户任务场景分类器。对每条 Query 只选择一个类别。\n"
        "分类应依据 Query 的整体语义，不应仅凭单个关键词。\n"
        "如果 Query 依赖不可见前文、自身没有独立任务语义，必须选择"
        f"“{rules[-1].name}”。\n"
        "不要引用或复述 Query，不要输出用户标识、凭据、路径等内容。\n"
        "confidence 是 0 到 1 的数字；reason 只写不超过 30 个汉字的概括。\n"
        "只输出一个 JSON 对象，不要使用 Markdown 代码块。格式："
        '{"items":[{"id":"Q001","category":"类别",'
        '"confidence":0.9,"reason":"简短理由"}]}\n'
        f"可选类别：\n{category_lines}"
    )


def secondary_semantic_system_prompt(rules: Sequence[CategoryRule]) -> str:
    category_lines = "\n".join(f"- {rule.name}: {rule.description}" for rule in rules)
    return (
        "你是 AstronCode 用户任务场景的二次分类器。输入 Query 已在第一次语义分类中被"
        "判为未分类，但已通过本地规则确认包含明确任务动作。请重新阅读整体语义，并对"
        "每条 Query 只选择一个类别。\n"
        "只要能够识别任务目标、交付物或操作对象，就优先选择最贴近的具体类别；不要因"
        "表达简短、口语化、询问工具能力或同时出现多个动作而直接放弃分类。\n"
        "只有问候、随机字符、纯对话控制、纯模型身份询问、依赖不可见前文，或确实无法"
        f"可靠映射到任一具体类别时，才选择“{rules[-1].name}”。\n"
        "不要根据不可见前文补全任务，不要仅凭单个关键词分类。不要引用或复述 Query，"
        "不要输出用户标识、凭据、路径等内容。\n"
        "confidence 是 0 到 1 的数字；reason 只写不超过 30 个汉字的概括。\n"
        "只输出一个 JSON 对象，不要使用 Markdown 代码块。格式："
        '{"items":[{"id":"Q001","category":"类别",'
        '"confidence":0.9,"reason":"简短理由"}]}\n'
        f"可选类别：\n{category_lines}"
    )


def parse_semantic_response(
    content: str,
    expected_ids: set[str],
    category_names: set[str],
) -> dict[str, tuple[str, float, str]]:
    stripped = content.strip()
    if stripped.startswith("```"):
        stripped = re.sub(r"^```(?:json)?\s*", "", stripped, flags=re.IGNORECASE)
        stripped = re.sub(r"\s*```$", "", stripped)
    start = stripped.find("{")
    end = stripped.rfind("}")
    if start < 0 or end < start:
        raise ValueError("语义分类响应不包含 JSON 对象")
    payload = json.loads(stripped[start : end + 1])
    items = payload.get("items")
    if not isinstance(items, list):
        raise TypeError("语义分类响应缺少 items 数组")
    parsed: dict[str, tuple[str, float, str]] = {}
    for item in items:
        if not isinstance(item, dict):
            raise TypeError("语义分类 items 中存在非对象")
        item_id = str(item.get("id") or "")
        category = str(item.get("category") or "")
        if item_id not in expected_ids or item_id in parsed:
            raise ValueError(f"语义分类返回未知或重复 id：{item_id}")
        if category not in category_names:
            raise ValueError(f"语义分类返回未知类别：{category}")
        confidence = float(item.get("confidence"))
        if not 0.0 <= confidence <= 1.0:
            raise ValueError(f"语义分类置信度越界：{confidence}")
        reason, _ = redact_query(str(item.get("reason") or ""))
        reason = re.sub(r"\s+", " ", reason).strip()[:120]
        parsed[item_id] = (category, confidence, reason)
    if set(parsed) != expected_ids:
        missing = sorted(expected_ids - set(parsed))
        raise ValueError(f"语义分类响应缺少 id：{missing[:5]}")
    return parsed


class SemanticSceneClient:
    def __init__(
        self,
        *,
        base_url: str,
        api_key: str,
        model: str,
        rules: Sequence[CategoryRule],
        timeout_seconds: float,
        retries: int,
        request_interval: float,
        system_prompt: str | None = None,
    ) -> None:
        if not base_url or not api_key or not model:
            raise ValueError("语义分类需要 base_url、api_key 和 model")
        self.endpoint = base_url.rstrip("/") + "/chat/completions"
        self.api_key = api_key
        self.model = model
        self.rules = tuple(rules)
        self.timeout_seconds = timeout_seconds
        self.retries = retries
        self.system_prompt = system_prompt or semantic_system_prompt(self.rules)
        self.rate_limiter = RequestRateLimiter(request_interval)
        self.single_query_rejections = 0
        self.rejection_lock = Lock()

    def classify_batch(
        self, items: Sequence[tuple[str, str, bool]]
    ) -> dict[str, tuple[str, float, str]]:
        request_items = [{"id": item_id, "query": query} for item_id, query, _ in items]
        expected_ids = {item_id for item_id, _, _ in items}
        last_error: Exception | None = None
        for attempt in range(self.retries + 1):
            try:
                self.rate_limiter.wait()
                session = requests.Session()
                session.trust_env = False
                try:
                    response = session.post(
                        self.endpoint,
                        headers={
                            "Authorization": f"Bearer {self.api_key}",
                            "Content-Type": "application/json",
                        },
                        json={
                            "model": self.model,
                            "messages": [
                                {
                                    "role": "system",
                                    "content": self.system_prompt,
                                },
                                {
                                    "role": "user",
                                    "content": json.dumps(
                                        {"items": request_items}, ensure_ascii=False
                                    ),
                                },
                            ],
                            "temperature": 0,
                            "max_tokens": min(8_000, max(2_000, len(items) * 120)),
                        },
                        timeout=self.timeout_seconds,
                    )
                finally:
                    session.close()
                if response.status_code in RETRYABLE_STATUS:
                    raise RuntimeError(
                        f"语义分类服务暂时不可用：HTTP {response.status_code}"
                    )
                response.raise_for_status()
                payload = response.json()
                content = payload["choices"][0]["message"]["content"]
                return parse_semantic_response(
                    str(content),
                    expected_ids,
                    {rule.name for rule in self.rules},
                )
            except (
                KeyError,
                TypeError,
                ValueError,
                requests.RequestException,
                RuntimeError,
            ) as error:
                last_error = error
                if (
                    isinstance(error, requests.HTTPError)
                    and error.response is not None
                    and error.response.status_code not in RETRYABLE_STATUS
                ):
                    break
                if attempt >= self.retries:
                    break
                time.sleep(min(8.0, (2**attempt) + random.random()))
        status_code = (
            last_error.response.status_code
            if isinstance(last_error, requests.HTTPError)
            and last_error.response is not None
            else None
        )
        if status_code in SPLITTABLE_SEMANTIC_STATUS:
            if len(items) > 1:
                midpoint = len(items) // 2
                parsed: dict[str, tuple[str, float, str]] = {}
                parsed.update(self.classify_batch(items[:midpoint]))
                parsed.update(self.classify_batch(items[midpoint:]))
                return parsed
            with self.rejection_lock:
                self.single_query_rejections += 1
                rejected_count = self.single_query_rejections
            if rejected_count <= MAX_SINGLE_QUERY_REJECTIONS:
                return {
                    items[0][0]: (
                        self.rules[-1].name,
                        0.0,
                        "分类服务拒绝单条请求",
                    )
                }
            raise RuntimeError(
                "语义分类服务拒绝的单条 Query 超过安全上限："
                f"{rejected_count}>{MAX_SINGLE_QUERY_REJECTIONS}"
            ) from last_error
        raise RuntimeError(
            "语义分类批次失败："
            f"items={len(items)}, chars={sum(len(item[1]) for item in items)}, "
            f"error={last_error}"
        ) from last_error


def prepare_remote_query(query: str) -> tuple[str, bool]:
    redacted, _ = redact_query(query)
    if len(redacted) <= DEFAULT_SCENE_QUERY_CHAR_LIMIT:
        return redacted, False
    return redacted[:DEFAULT_SCENE_QUERY_CHAR_LIMIT], True


def build_semantic_batches(
    queries: Sequence[tuple[str, str]],
    *,
    batch_size: int,
    batch_char_limit: int,
) -> list[list[tuple[str, str, str, bool]]]:
    batches: list[list[tuple[str, str, str, bool]]] = []
    current: list[tuple[str, str, str, bool]] = []
    current_chars = 0
    for query_hash, query in queries:
        prepared, truncated = prepare_remote_query(query)
        estimated_chars = len(prepared) + 80
        if current and (
            len(current) >= batch_size
            or current_chars + estimated_chars > batch_char_limit
        ):
            batches.append(current)
            current = []
            current_chars = 0
        item_id = f"Q{len(current) + 1:03d}"
        current.append((item_id, query_hash, prepared, truncated))
        current_chars += estimated_chars
    if current:
        batches.append(current)
    return batches


def write_scene_classification_parquet(
    classifications: Mapping[str, SceneClassification],
    output_path: Path,
    *,
    rule_version: str,
    model: str,
    confidence_threshold: float,
    classification_version: str = HYBRID_CLASSIFICATION_VERSION,
    prompt_version: str = SCENE_PROMPT_VERSION,
    extra_metadata: Mapping[str, str] | None = None,
) -> None:
    ensure_suffix(output_path, ".parquet", "场景分类输出路径")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    metadata = {
        b"schema_version": SCENE_CLASSIFICATION_SCHEMA_VERSION.encode(),
        b"classification_version": classification_version.encode(),
        b"rule_version": rule_version.encode(),
        b"model": model.encode(),
        b"prompt_version": prompt_version.encode(),
        b"confidence_threshold": str(confidence_threshold).encode(),
    }
    if extra_metadata:
        metadata.update(
            {
                str(key).encode(): str(value).encode()
                for key, value in extra_metadata.items()
            }
        )
    schema = scene_classification_schema().with_metadata(metadata)
    records = [
        classification_to_record(classifications[key])
        for key in sorted(classifications)
    ]
    table = pa.Table.from_pylist(records, schema=schema)
    temporary_path = output_path.with_suffix(output_path.suffix + ".tmp")
    pq.write_table(table, temporary_path, compression="zstd")
    os.replace(temporary_path, output_path)
    restrict_file_permissions(output_path)


def classify_scenes_to_parquet(
    input_path: Path,
    output_path: Path,
    checkpoint_path: Path,
    rules_path: Path,
    *,
    base_url: str,
    api_key: str,
    model: str,
    confidence_threshold: float,
    batch_size: int,
    batch_char_limit: int,
    workers: int,
    request_interval: float,
    timeout_seconds: float,
    retries: int,
    resume: bool,
) -> dict[str, int]:
    ensure_suffix(input_path, ".parquet", "标准化数据输入路径")
    ensure_suffix(output_path, ".parquet", "场景分类输出路径")
    if output_path.exists() and not resume:
        raise ValueError(f"场景分类输出已存在，拒绝覆盖：{output_path}")
    rule_version, rules = load_category_rules(rules_path)
    fallback_category = rules[-1].name
    parquet_file = pq.ParquetFile(input_path)
    unique_queries: dict[str, str] = {}
    for batch in parquet_file.iter_batches(
        columns=["query"], batch_size=PARQUET_BATCH_SIZE
    ):
        for record in batch.to_pylist():
            display, query_key = normalize_query(str(record.get("query") or ""))
            if query_key:
                unique_queries.setdefault(query_key_hash(query_key), display)

    classifications: dict[str, SceneClassification] = {}
    pending_remote: list[tuple[str, str]] = []
    for query_hash, query in sorted(unique_queries.items()):
        rule_category = classify_query(query, rules)
        if rule_category != fallback_category:
            classifications[query_hash] = SceneClassification(
                query_hash=query_hash,
                category=rule_category,
                proposed_category=rule_category,
                method="规则",
                confidence=1.0,
                reason="规则优先级命中",
                model="",
                prompt_version=SCENE_PROMPT_VERSION,
                rule_version=rule_version,
                query_length=len(query),
                input_truncated=False,
            )
        elif is_context_dependent_query(query):
            classifications[query_hash] = SceneClassification(
                query_hash=query_hash,
                category=fallback_category,
                proposed_category=fallback_category,
                method="上下文依赖",
                confidence=1.0,
                reason="缺少独立任务语义",
                model="",
                prompt_version=SCENE_PROMPT_VERSION,
                rule_version=rule_version,
                query_length=len(query),
                input_truncated=False,
            )
        else:
            pending_remote.append((query_hash, query))

    checkpoint_metadata = {
        "record_type": "metadata",
        "schema_version": SCENE_CLASSIFICATION_SCHEMA_VERSION,
        "classification_version": HYBRID_CLASSIFICATION_VERSION,
        "rule_version": rule_version,
        "model": model,
        "prompt_version": SCENE_PROMPT_VERSION,
        "confidence_threshold": confidence_threshold,
    }
    checkpoint_path.parent.mkdir(parents=True, exist_ok=True)
    if resume and checkpoint_path.exists():
        with checkpoint_path.open("r", encoding="utf-8") as checkpoint_file:
            for line_number, line in enumerate(checkpoint_file, start=1):
                payload = json.loads(line)
                if line_number == 1:
                    if payload != checkpoint_metadata:
                        raise ValueError("场景分类 checkpoint 配置与本次运行不一致")
                    continue
                classification = classification_from_record(payload)
                if classification.query_hash in unique_queries:
                    classifications[classification.query_hash] = classification
    else:
        if checkpoint_path.exists():
            raise ValueError(
                f"checkpoint 已存在，使用 --resume 继续：{checkpoint_path}"
            )
        with checkpoint_path.open("x", encoding="utf-8") as checkpoint_file:
            checkpoint_file.write(
                json.dumps(checkpoint_metadata, ensure_ascii=False) + "\n"
            )
        restrict_file_permissions(checkpoint_path)

    pending_remote = [item for item in pending_remote if item[0] not in classifications]
    batches = build_semantic_batches(
        pending_remote,
        batch_size=batch_size,
        batch_char_limit=batch_char_limit,
    )
    client = SemanticSceneClient(
        base_url=base_url,
        api_key=api_key,
        model=model,
        rules=rules,
        timeout_seconds=timeout_seconds,
        retries=retries,
        request_interval=request_interval,
    )

    completed_batches = 0
    with ThreadPoolExecutor(max_workers=workers) as executor:
        future_map: dict[Any, list[tuple[str, str, str, bool]]] = {}
        query_by_hash = dict(pending_remote)
        batch_iterator = iter(batches)

        def submit_next_batch() -> bool:
            try:
                batch = next(batch_iterator)
            except StopIteration:
                return False
            future = executor.submit(
                client.classify_batch,
                [
                    (item_id, prepared, truncated)
                    for item_id, _, prepared, truncated in batch
                ],
            )
            future_map[future] = batch
            return True

        for _ in range(min(workers, len(batches))):
            submit_next_batch()
        while future_map:
            completed_futures, _ = wait(tuple(future_map), return_when=FIRST_COMPLETED)
            for future in completed_futures:
                bound_batch = future_map.pop(future)
                parsed = future.result()
                new_records: list[SceneClassification] = []
                for item_id, query_hash, _, truncated in bound_batch:
                    proposed_category, confidence, reason = parsed[item_id]
                    if proposed_category == fallback_category:
                        final_category = fallback_category
                        method = "大模型未判定"
                    elif confidence < confidence_threshold:
                        final_category = fallback_category
                        method = "大模型低置信度"
                    else:
                        final_category = proposed_category
                        method = "大模型"
                    classification = SceneClassification(
                        query_hash=query_hash,
                        category=final_category,
                        proposed_category=proposed_category,
                        method=method,
                        confidence=confidence,
                        reason=reason,
                        model=model,
                        prompt_version=SCENE_PROMPT_VERSION,
                        rule_version=rule_version,
                        query_length=len(query_by_hash[query_hash]),
                        input_truncated=truncated,
                    )
                    classifications[query_hash] = classification
                    new_records.append(classification)
                with checkpoint_path.open("a", encoding="utf-8") as checkpoint_file:
                    for classification in new_records:
                        checkpoint_file.write(
                            json.dumps(
                                classification_to_record(classification),
                                ensure_ascii=False,
                            )
                            + "\n"
                        )
                completed_batches += 1
                if completed_batches % 10 == 0 or completed_batches == len(batches):
                    print(
                        f"语义分类进度：{completed_batches:,}/{len(batches):,} 批，"
                        f"已完成 {len(classifications):,}/{len(unique_queries):,} 个唯一 Query",
                        file=sys.stderr,
                        flush=True,
                    )
                submit_next_batch()

    if set(classifications) != set(unique_queries):
        missing = len(set(unique_queries) - set(classifications))
        raise RuntimeError(f"场景分类未覆盖全部唯一 Query：缺少 {missing:,} 条")
    write_scene_classification_parquet(
        classifications,
        output_path,
        rule_version=rule_version,
        model=model,
        confidence_threshold=confidence_threshold,
    )
    if checkpoint_path.exists():
        checkpoint_path.unlink()
    method_counts = Counter(item.method for item in classifications.values())
    return {
        "unique_queries": len(unique_queries),
        "rule_queries": method_counts["规则"],
        "context_dependent_queries": method_counts["上下文依赖"],
        "semantic_queries": sum(
            count
            for method, count in method_counts.items()
            if method.startswith("大模型")
        ),
        "unclassified_queries": sum(
            count
            for item, count in Counter(
                classification.category for classification in classifications.values()
            ).items()
            if item == fallback_category
        ),
    }


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def refine_scene_classifications_to_parquet(
    input_path: Path,
    base_classifications_path: Path,
    output_path: Path,
    checkpoint_path: Path,
    rules_path: Path,
    *,
    base_url: str,
    api_key: str,
    model: str,
    confidence_threshold: float,
    minimum_query_length: int,
    batch_size: int,
    batch_char_limit: int,
    workers: int,
    request_interval: float,
    timeout_seconds: float,
    retries: int,
    resume: bool,
) -> dict[str, int]:
    ensure_suffix(input_path, ".parquet", "标准化数据输入路径")
    ensure_suffix(base_classifications_path, ".parquet", "基础场景分类输入路径")
    ensure_suffix(output_path, ".parquet", "场景分类输出路径")
    if output_path.exists() and not resume:
        raise ValueError(f"场景分类输出已存在，拒绝覆盖：{output_path}")

    rule_version, rules = load_category_rules(rules_path)
    fallback_category = rules[-1].name
    base_set = load_scene_classifications(base_classifications_path, rules)

    parquet_file = pq.ParquetFile(input_path)
    unique_queries: dict[str, str] = {}
    for batch in parquet_file.iter_batches(
        columns=["query"], batch_size=PARQUET_BATCH_SIZE
    ):
        for record in batch.to_pylist():
            display, query_key = normalize_query(str(record.get("query") or ""))
            if query_key:
                unique_queries.setdefault(query_key_hash(query_key), display)

    expected_hashes = set(unique_queries)
    actual_hashes = set(base_set.classifications)
    if actual_hashes != expected_hashes:
        missing = len(expected_hashes - actual_hashes)
        extra = len(actual_hashes - expected_hashes)
        raise ValueError(
            "基础场景分类与标准化数据不一致："
            f"缺少 {missing:,} 个 Query，多出 {extra:,} 个 Query"
        )

    classifications: dict[str, SceneClassification] = {}
    pending_secondary: list[tuple[str, str]] = []
    new_context_queries = 0
    preserved_queries = 0
    for query_hash, query in sorted(unique_queries.items()):
        base = canonicalize_fallback_category(
            base_set.classifications[query_hash], fallback_category
        )
        base = replace(base, rule_version=rule_version)
        if base.method != "大模型未判定":
            classifications[query_hash] = base
            preserved_queries += 1
            continue
        if is_context_dependent_query(query):
            classifications[query_hash] = SceneClassification(
                query_hash=query_hash,
                category=fallback_category,
                proposed_category=fallback_category,
                method="上下文依赖",
                confidence=1.0,
                reason="缺少独立任务语义（二次识别）",
                model="",
                prompt_version=SECONDARY_SELECTOR_VERSION,
                rule_version=rule_version,
                query_length=len(query),
                input_truncated=False,
            )
            new_context_queries += 1
            continue
        if is_secondary_classification_candidate(
            query, minimum_length=minimum_query_length
        ):
            pending_secondary.append((query_hash, query))
            continue
        classifications[query_hash] = base
        preserved_queries += 1

    secondary_candidate_queries = len(pending_secondary)
    checkpoint_metadata = {
        "record_type": "metadata",
        "schema_version": SCENE_CLASSIFICATION_SCHEMA_VERSION,
        "classification_version": REFINED_CLASSIFICATION_VERSION,
        "rule_version": rule_version,
        "model": model,
        "prompt_version": SECONDARY_SCENE_PROMPT_VERSION,
        "confidence_threshold": confidence_threshold,
        "secondary_selector_version": SECONDARY_SELECTOR_VERSION,
        "secondary_min_query_length": minimum_query_length,
        "base_classification_source": base_classifications_path.name,
        "base_classification_sha256": file_sha256(base_classifications_path),
    }
    checkpoint_path.parent.mkdir(parents=True, exist_ok=True)
    if resume and checkpoint_path.exists():
        pending_hashes = {query_hash for query_hash, _ in pending_secondary}
        with checkpoint_path.open("r", encoding="utf-8") as checkpoint_file:
            for line_number, line in enumerate(checkpoint_file, start=1):
                payload = json.loads(line)
                if line_number == 1:
                    if payload != checkpoint_metadata:
                        raise ValueError("二次分类 checkpoint 配置与本次运行不一致")
                    continue
                classification = canonicalize_fallback_category(
                    classification_from_record(payload), fallback_category
                )
                if classification.query_hash in pending_hashes:
                    classifications[classification.query_hash] = classification
    else:
        if checkpoint_path.exists():
            raise ValueError(
                f"checkpoint 已存在，使用 --resume 继续：{checkpoint_path}"
            )
        with checkpoint_path.open("x", encoding="utf-8") as checkpoint_file:
            checkpoint_file.write(
                json.dumps(checkpoint_metadata, ensure_ascii=False) + "\n"
            )
        restrict_file_permissions(checkpoint_path)

    pending_secondary = [
        item for item in pending_secondary if item[0] not in classifications
    ]
    batches = build_semantic_batches(
        pending_secondary,
        batch_size=batch_size,
        batch_char_limit=batch_char_limit,
    )
    client = SemanticSceneClient(
        base_url=base_url,
        api_key=api_key,
        model=model,
        rules=rules,
        timeout_seconds=timeout_seconds,
        retries=retries,
        request_interval=request_interval,
        system_prompt=secondary_semantic_system_prompt(rules),
    )

    completed_batches = 0
    with ThreadPoolExecutor(max_workers=workers) as executor:
        future_map: dict[Any, list[tuple[str, str, str, bool]]] = {}
        query_by_hash = dict(pending_secondary)
        batch_iterator = iter(batches)

        def submit_next_batch() -> bool:
            try:
                batch = next(batch_iterator)
            except StopIteration:
                return False
            future = executor.submit(
                client.classify_batch,
                [
                    (item_id, prepared, truncated)
                    for item_id, _, prepared, truncated in batch
                ],
            )
            future_map[future] = batch
            return True

        for _ in range(min(workers, len(batches))):
            submit_next_batch()
        while future_map:
            completed_futures, _ = wait(tuple(future_map), return_when=FIRST_COMPLETED)
            for future in completed_futures:
                bound_batch = future_map.pop(future)
                parsed = future.result()
                new_records: list[SceneClassification] = []
                for item_id, query_hash, _, truncated in bound_batch:
                    proposed_category, confidence, reason = parsed[item_id]
                    if proposed_category == fallback_category:
                        final_category = fallback_category
                        method = "大模型二次未判定"
                    elif confidence < confidence_threshold:
                        final_category = fallback_category
                        method = "大模型二次低置信度"
                    else:
                        final_category = proposed_category
                        method = "大模型二次分类"
                    classification = SceneClassification(
                        query_hash=query_hash,
                        category=final_category,
                        proposed_category=proposed_category,
                        method=method,
                        confidence=confidence,
                        reason=reason,
                        model=model,
                        prompt_version=SECONDARY_SCENE_PROMPT_VERSION,
                        rule_version=rule_version,
                        query_length=len(query_by_hash[query_hash]),
                        input_truncated=truncated,
                    )
                    classifications[query_hash] = classification
                    new_records.append(classification)
                with checkpoint_path.open("a", encoding="utf-8") as checkpoint_file:
                    for classification in new_records:
                        checkpoint_file.write(
                            json.dumps(
                                classification_to_record(classification),
                                ensure_ascii=False,
                            )
                            + "\n"
                        )
                completed_batches += 1
                if completed_batches % 10 == 0 or completed_batches == len(batches):
                    print(
                        f"二次分类进度：{completed_batches:,}/{len(batches):,} 批，"
                        f"已完成 {len(classifications):,}/{len(unique_queries):,} 个唯一 Query",
                        file=sys.stderr,
                        flush=True,
                    )
                submit_next_batch()

    if set(classifications) != set(unique_queries):
        missing = len(set(unique_queries) - set(classifications))
        raise RuntimeError(f"二次分类未覆盖全部唯一 Query：缺少 {missing:,} 条")

    model_versions = list(dict.fromkeys(item for item in (base_set.model, model) if item))
    prompt_versions = list(
        dict.fromkeys(
            item
            for item in (base_set.prompt_version, SECONDARY_SCENE_PROMPT_VERSION)
            if item
        )
    )
    write_scene_classification_parquet(
        classifications,
        output_path,
        rule_version=rule_version,
        model=" + ".join(model_versions),
        confidence_threshold=confidence_threshold,
        classification_version=REFINED_CLASSIFICATION_VERSION,
        prompt_version=" + ".join(prompt_versions),
        extra_metadata={
            "secondary_prompt_version": SECONDARY_SCENE_PROMPT_VERSION,
            "secondary_selector_version": SECONDARY_SELECTOR_VERSION,
            "secondary_min_query_length": str(minimum_query_length),
            "base_classification_source": base_classifications_path.name,
        },
    )
    if checkpoint_path.exists():
        checkpoint_path.unlink()

    method_counts = Counter(item.method for item in classifications.values())
    category_counts = Counter(item.category for item in classifications.values())
    return {
        "unique_queries": len(unique_queries),
        "preserved_queries": preserved_queries,
        "new_context_queries": new_context_queries,
        "secondary_candidate_queries": secondary_candidate_queries,
        "secondary_classified_queries": method_counts["大模型二次分类"],
        "secondary_low_confidence_queries": method_counts["大模型二次低置信度"],
        "secondary_unjudged_queries": method_counts["大模型二次未判定"],
        "unclassified_queries": category_counts[fallback_category],
    }


def is_present(field_name: str, value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, str):
        stripped = value.strip()
        if not stripped:
            return False
        if field_name in {"model", "reasoning_effort"} and stripped == "未记录":
            return False
    return True


def status_is_success(value: Any) -> bool:
    return str(value or "").strip().casefold() in SUCCESS_STATUS_VALUES


def update_group(metrics: GroupMetrics, record: Mapping[str, Any]) -> None:
    metrics.count += 1
    duration = finite_nonnegative(record.get("duration_seconds"))
    if duration is not None:
        metrics.durations.append(duration)
    total_tokens = finite_nonnegative(record.get("total_tokens"))
    if total_tokens is not None:
        metrics.total_tokens.append(total_tokens)
    user_hash = str(record.get("user_hash") or "")
    session_hash = str(record.get("session_hash") or "")
    if user_hash:
        metrics.users.add(user_hash)
    if session_hash:
        metrics.sessions.add(session_hash)
    status = str(record.get("turn_status") or "").strip()
    if status:
        metrics.known_status_count += 1
        metrics.success_count += int(status_is_success(status))


def analyze_parquet(
    input_path: Path,
    rules: Sequence[CategoryRule],
    scene_classifications: SceneClassificationSet | None = None,
) -> DatasetAnalysis:
    if input_path.suffix.lower() != ".parquet":
        raise ValueError("标准化数据输入必须使用 .parquet 扩展名")
    parquet_file = pq.ParquetFile(input_path)
    metadata = parquet_file.schema_arrow.metadata or {}
    schema_version = metadata.get(b"schema_version", b"").decode()
    if schema_version != PARQUET_SCHEMA_VERSION:
        raise ValueError(f"Parquet schema_version 不支持：{schema_version or '缺失'}")
    timezone_name = metadata.get(b"timezone", DEFAULT_TIMEZONE.encode()).decode()
    required = {field.name for field in parquet_schema()}
    actual = set(parquet_file.schema_arrow.names)
    if not required.issubset(actual):
        raise ValueError(f"Parquet 缺少字段：{sorted(required - actual)}")
    analysis = DatasetAnalysis(
        total_rows=0,
        timezone_name=timezone_name,
        schema_version=schema_version,
    )
    if scene_classifications is not None:
        analysis.classification_version = scene_classifications.classification_version
        analysis.classification_rule_version = scene_classifications.rule_version
        analysis.classification_model = scene_classifications.model
        analysis.classification_prompt_version = scene_classifications.prompt_version
        analysis.classification_confidence_threshold = (
            scene_classifications.confidence_threshold
        )
        analysis.classification_source_name = scene_classifications.source_name
        analysis.classification_secondary_prompt_version = (
            scene_classifications.secondary_prompt_version
        )
        analysis.classification_secondary_selector_version = (
            scene_classifications.secondary_selector_version
        )
        analysis.classification_secondary_min_query_length = (
            scene_classifications.secondary_min_query_length
        )
        analysis.classification_base_source_name = (
            scene_classifications.base_classification_source
        )
    for batch in parquet_file.iter_batches(batch_size=PARQUET_BATCH_SIZE):
        for record in batch.to_pylist():
            analysis.total_rows += 1
            query = str(record.get("query") or "")
            display_query, query_key = normalize_query(query)
            try:
                query_length = int(record.get("query_length"))
            except (TypeError, ValueError):
                query_length = len(query)
            occurred_at = record.get("occurred_at")
            if isinstance(occurred_at, datetime):
                analysis.min_time = (
                    occurred_at
                    if analysis.min_time is None
                    else min(analysis.min_time, occurred_at)
                )
                analysis.max_time = (
                    occurred_at
                    if analysis.max_time is None
                    else max(analysis.max_time, occurred_at)
                )
            user_hash = str(record.get("user_hash") or "")
            session_hash = str(record.get("session_hash") or "")
            if user_hash:
                analysis.unique_users.add(user_hash)
            if session_hash:
                analysis.unique_sessions.add(session_hash)

            for field_name in parquet_file.schema_arrow.names:
                if is_present(field_name, record.get(field_name)):
                    analysis.field_non_missing[field_name] += 1

            duration = finite_nonnegative(record.get("duration_seconds"))
            if duration is not None:
                analysis.duration_values.append(duration)
            for token_field, _ in TOKEN_FIELDS:
                token_value = finite_nonnegative(record.get(token_field))
                if token_value is not None:
                    analysis.token_values[token_field].append(token_value)

            if scene_classifications is None:
                category = classify_query(query, rules)
                if not query_key:
                    classification_method = "空Query"
                elif category == rules[-1].name:
                    classification_method = "规则未命中"
                else:
                    classification_method = "规则"
                classification_confidence = None
                classification_model = ""
                classification_prompt_version = ""
            else:
                if not query_key:
                    category = rules[-1].name
                    classification_method = "空Query"
                    classification_confidence = 1.0
                    classification_model = ""
                    classification_prompt_version = scene_classifications.prompt_version
                else:
                    classification = scene_classifications.classifications.get(
                        query_key_hash(query_key)
                    )
                    if classification is None:
                        raise RuntimeError(
                            "场景分类数据未覆盖当前 Parquet 中的全部唯一 Query"
                        )
                    category = classification.category
                    classification_method = classification.method
                    classification_confidence = classification.confidence
                    classification_model = classification.model
                    classification_prompt_version = classification.prompt_version
            length_band = query_length_band(query_length)
            model = str(record.get("model") or "未记录").strip() or "未记录"
            effort = str(record.get("reasoning_effort") or "未记录").strip() or "未记录"
            update_group(analysis.category_metrics[category], record)
            update_group(
                analysis.classification_method_metrics[classification_method], record
            )
            if (
                classification_method.startswith("大模型")
                and classification_confidence is not None
            ):
                analysis.classification_confidence_values.append(
                    classification_confidence
                )
            update_group(analysis.length_metrics[length_band], record)
            update_group(analysis.model_metrics[model], record)
            for product, source_field, parquet_field in VERSION_DIMENSIONS:
                version = str(record.get(parquet_field) or "").strip() or "未记录"
                update_group(
                    analysis.version_metrics[(product, source_field, version)], record
                )
            update_group(analysis.effort_metrics[effort], record)

            for status_type, field_name in (
                ("Turn状态", "turn_status"),
                ("Span状态", "span_status"),
            ):
                status = str(record.get(field_name) or "").strip() or "未记录"
                update_group(
                    analysis.status_metrics[f"{status_type}\t{status}"], record
                )
            turn_status = str(record.get("turn_status") or "").strip()
            if turn_status:
                analysis.known_turn_status_count += 1
                analysis.successful_turn_count += int(status_is_success(turn_status))

            if isinstance(occurred_at, datetime):
                update_group(analysis.daily_metrics[occurred_at.date()], record)

            if not query_key:
                continue
            candidate = analysis.candidates.get(query_key)
            if candidate is None:
                candidate = CandidateAggregate(
                    representative_query=display_query,
                    query_length=query_length,
                    category=category,
                    classification_method=classification_method,
                    classification_confidence=classification_confidence,
                    classification_model=classification_model,
                    classification_prompt_version=classification_prompt_version,
                )
                analysis.candidates[query_key] = candidate
            candidate.count += 1
            if session_hash:
                candidate.sessions.add(session_hash)
            candidate.models[model] += 1
            if duration is not None:
                candidate.durations.append(duration)
            total_tokens = finite_nonnegative(record.get("total_tokens"))
            if total_tokens is not None:
                candidate.total_tokens.append(total_tokens)
            if isinstance(occurred_at, datetime):
                candidate.first_time = (
                    occurred_at
                    if candidate.first_time is None
                    else min(candidate.first_time, occurred_at)
                )
                candidate.last_time = (
                    occurred_at
                    if candidate.last_time is None
                    else max(candidate.last_time, occurred_at)
                )
            image_count = optional_integer(record.get("input_image_count")) or 0
            candidate.image_task_count += int(image_count > 0)
            candidate.input_image_count += image_count
    return analysis


def add_sheet_title(sheet: Any, title: str) -> None:
    sheet.sheet_view.showGridLines = False
    sheet["A2"] = title
    style_title(sheet["A2"])
    sheet.row_dimensions[2].height = 24


def average(values: Sequence[float]) -> float | None:
    return sum(values) / len(values) if values else None


def share(value: int, total: int) -> float | None:
    return value / total if total else None


def success_rate(metrics: GroupMetrics) -> float | None:
    return share(metrics.success_count, metrics.known_status_count)


def configure_table_sheet(
    sheet: Any,
    *,
    last_row: int,
    last_column: int,
    widths: Sequence[float],
    freeze: bool = True,
) -> None:
    sheet.sheet_view.showGridLines = False
    if freeze and last_row >= 5:
        sheet.freeze_panes = "A5"
    if last_row >= 4:
        sheet.auto_filter.ref = f"A4:{get_column_letter(last_column)}{last_row}"
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=1, max_row=last_row, max_col=last_column):
        for cell in row:
            if cell.row not in {2, 4}:
                cell.font = Font(name="Arial", size=10, bold=cell.font.bold)
            cell.alignment = Alignment(
                horizontal=cell.alignment.horizontal,
                vertical="center",
                wrap_text=cell.alignment.wrap_text,
            )


def group_row(label: str, metrics: GroupMetrics, total: int) -> list[Any]:
    return [
        label,
        metrics.count,
        share(metrics.count, total),
        len(metrics.users),
        len(metrics.sessions),
        len(metrics.durations),
        linear_percentile(metrics.durations, 0.50),
        linear_percentile(metrics.durations, 0.90),
        len(metrics.total_tokens),
        linear_percentile(metrics.total_tokens, 0.50),
        linear_percentile(metrics.total_tokens, 0.90),
        success_rate(metrics),
    ]


def format_group_sheet(sheet: Any, start_row: int = 5) -> None:
    for row in sheet.iter_rows(min_row=start_row, max_row=sheet.max_row):
        row[2].number_format = "0.0%"
        for cell in (row[6], row[7]):
            cell.number_format = '0.000" 秒"'
        for cell in (row[9], row[10]):
            cell.number_format = "#,##0"
        row[11].number_format = "0.0%"


def write_analysis_workbook(
    analysis: DatasetAnalysis,
    output_path: Path,
    *,
    input_path: Path,
    category_version: str,
    rules: Sequence[CategoryRule],
) -> None:
    ensure_suffix(output_path, ".xlsx", "分析结果输出路径")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    overview = workbook.active
    overview.title = "分析总览"
    overview.sheet_properties.tabColor = "1F4E78"
    add_sheet_title(overview, "AstronCode 生产使用数据分析")
    overview.append([])
    overview.append(["指标", "值", "口径"])
    style_table_header(overview[4])
    overview_rows = [
        ("去重后样本数", analysis.total_rows, "agent.turn 记录"),
        ("候选 Query 数", len(analysis.candidates), "非空 Query 规范化后精确去重"),
        ("去标识用户数", len(analysis.unique_users), "加盐哈希后去重"),
        ("去标识会话数", len(analysis.unique_sessions), "加盐哈希后去重"),
        ("数据起始时间", analysis.min_time, analysis.timezone_name),
        ("数据截止时间", analysis.max_time, analysis.timezone_name),
        ("有效耗时记录数", len(analysis.duration_values), "缺失不按 0 计"),
        (
            "任务耗时 P90",
            linear_percentile(analysis.duration_values, 0.90),
            "线性插值，单位秒",
        ),
        (
            "任务耗时 P95",
            linear_percentile(analysis.duration_values, 0.95),
            "线性插值，单位秒",
        ),
        (
            "任务耗时 P99",
            linear_percentile(analysis.duration_values, 0.99),
            "线性插值，单位秒",
        ),
        ("Turn 状态已记录数", analysis.known_turn_status_count, "空值不计入成功率分母"),
        (
            "Turn 成功率",
            share(analysis.successful_turn_count, analysis.known_turn_status_count),
            "completed/success/ok 等状态",
        ),
        ("场景分类版本", analysis.classification_version, "单标签混合分类"),
        ("分类规则版本", category_version, "高精度规则优先"),
        (
            "分类数据文件",
            analysis.classification_source_name or "未使用（纯规则）",
            "本地分类结果，可复用且无需重新访问模型",
        ),
        (
            "语义分类模型",
            analysis.classification_model or "未使用",
            analysis.classification_prompt_version or "纯规则分析",
        ),
        (
            "语义置信度阈值",
            analysis.classification_confidence_threshold,
            "低于阈值保持未分类",
        ),
        (
            "二次分类提示词版本",
            analysis.classification_secondary_prompt_version or "未使用",
            "仅处理初次大模型未判定且具有明确任务意图的 Query",
        ),
        (
            "二次分类候选策略",
            analysis.classification_secondary_selector_version or "未使用",
            (
                f"规范化长度不少于 {analysis.classification_secondary_min_query_length} 个字符，"
                "包含明确任务动作，且不属于上下文依赖"
                if analysis.classification_secondary_min_query_length is not None
                else "未使用"
            ),
        ),
        (
            "二次分类命中记录数",
            (
                analysis.classification_method_metrics.get(
                    "大模型二次分类", GroupMetrics()
                ).count
            ),
            "按 agent.turn 记录计数，仅统计达到置信度阈值的具体分类",
        ),
        (
            "大模型分类记录数",
            sum(
                metrics.count
                for name, metrics in analysis.classification_method_metrics.items()
                if name.startswith("大模型")
            ),
            "按 agent.turn 记录计数",
        ),
        (
            "上下文依赖记录数",
            analysis.classification_method_metrics["上下文依赖"].count,
            "不根据不可见前文猜测场景",
        ),
        ("Parquet schema 版本", analysis.schema_version, "标准化数据口径"),
        ("标准化数据文件", input_path.name, "本地文件名"),
    ]
    for row in overview_rows:
        overview.append(row)
    for row_index in range(5, 5 + len(overview_rows)):
        label = overview.cell(row_index, 1).value
        if label in {"数据起始时间", "数据截止时间"}:
            overview.cell(row_index, 2).number_format = "yyyy-mm-dd hh:mm:ss"
        elif "耗时 P" in str(label):
            overview.cell(row_index, 2).number_format = '0.000" 秒"'
        elif label in {"Turn 成功率", "语义置信度阈值"}:
            overview.cell(row_index, 2).number_format = "0.0%"
        elif isinstance(overview.cell(row_index, 2).value, int):
            overview.cell(row_index, 2).number_format = "#,##0"
    configure_table_sheet(
        overview,
        last_row=overview.max_row,
        last_column=3,
        widths=(28, 72, 48),
        freeze=False,
    )

    descriptions = {rule.name: rule.description for rule in rules}
    group_headers = [
        "维度值",
        "任务数",
        "占比",
        "用户数",
        "会话数",
        "有效耗时数",
        "耗时P50",
        "耗时P90",
        "有效总Token记录数",
        "总Token P50",
        "总Token P90",
        "Turn成功率",
    ]

    scene_sheet = workbook.create_sheet("场景分类")
    add_sheet_title(scene_sheet, "场景分类分布")
    scene_sheet.append([])
    scene_sheet.append(group_headers + ["分类说明"])
    style_table_header(scene_sheet[4])
    ordered_categories = sorted(
        descriptions,
        key=lambda name: (
            -analysis.category_metrics[name].count,
            list(descriptions).index(name),
        ),
    )
    for name in ordered_categories:
        scene_sheet.append(
            group_row(name, analysis.category_metrics[name], analysis.total_rows)
            + [descriptions[name]]
        )
    format_group_sheet(scene_sheet)
    for row in scene_sheet.iter_rows(min_row=5, max_row=scene_sheet.max_row):
        row[12].alignment = Alignment(vertical="top", wrap_text=True)
    configure_table_sheet(
        scene_sheet,
        last_row=scene_sheet.max_row,
        last_column=13,
        widths=(28, 12, 10, 12, 12, 15, 14, 14, 17, 15, 15, 14, 58),
    )

    classification_method_descriptions = {
        "规则": "高精度规则直接命中。",
        "规则未命中": "纯规则分析中未命中任何规则。",
        "上下文依赖": "短指令依赖不可见前文，保持未分类。",
        "大模型": "规则未命中后由大模型语义分类，且置信度达到阈值。",
        "大模型低置信度": "大模型给出候选类别，但置信度低于阈值，保持未分类。",
        "大模型未判定": "大模型判断 Query 缺少足够独立语义，保持未分类。",
        "大模型二次分类": "初次未判定但具有明确任务意图，二次语义分类达到阈值。",
        "大模型二次低置信度": "二次分类给出候选类别，但置信度低于阈值，保持未分类。",
        "大模型二次未判定": "二次分类仍无法可靠映射到具体场景，保持未分类。",
        "空Query": "Query 为空，保持未分类。",
    }
    classification_sheet = workbook.create_sheet("分类方法")
    add_sheet_title(classification_sheet, "场景分类方法分布")
    classification_sheet.append([])
    classification_sheet.append(group_headers + ["方法说明"])
    style_table_header(classification_sheet[4])
    for name, metrics in sorted(
        analysis.classification_method_metrics.items(),
        key=lambda item: (-item[1].count, item[0]),
    ):
        classification_sheet.append(
            group_row(name, metrics, analysis.total_rows)
            + [classification_method_descriptions.get(name, "")]
        )
    format_group_sheet(classification_sheet)
    for row in classification_sheet.iter_rows(
        min_row=5, max_row=classification_sheet.max_row
    ):
        row[12].alignment = Alignment(vertical="top", wrap_text=True)
    configure_table_sheet(
        classification_sheet,
        last_row=classification_sheet.max_row,
        last_column=13,
        widths=(28, 12, 10, 12, 12, 15, 14, 14, 17, 15, 15, 14, 58),
    )

    length_sheet = workbook.create_sheet("Query长度")
    add_sheet_title(length_sheet, "Query 长度分布")
    length_sheet.append([])
    length_sheet.append(["长度梯度", "下界（含）", "上界（含）"] + group_headers[1:])
    style_table_header(length_sheet[4])
    for label, lower, upper in LENGTH_BANDS:
        length_sheet.append(
            [label, lower, upper]
            + group_row(label, analysis.length_metrics[label], analysis.total_rows)[1:]
        )
    for row in length_sheet.iter_rows(min_row=5, max_row=length_sheet.max_row):
        row[4].number_format = "0.0%"
        row[8].number_format = '0.000" 秒"'
        row[9].number_format = '0.000" 秒"'
        row[11].number_format = "#,##0"
        row[12].number_format = "#,##0"
        row[13].number_format = "0.0%"
    configure_table_sheet(
        length_sheet,
        last_row=length_sheet.max_row,
        last_column=14,
        widths=(18, 13, 13, 12, 10, 12, 12, 15, 14, 14, 17, 15, 15, 14),
    )

    duration_sheet = workbook.create_sheet("耗时分布")
    add_sheet_title(duration_sheet, "任务耗时分布")
    duration_sheet.append([])
    duration_sheet.append(
        [
            "维度",
            "维度值",
            "任务数",
            "有效耗时数",
            "覆盖率",
            "平均值",
            "P50",
            "P90",
            "P95",
            "P99",
        ]
    )
    style_table_header(duration_sheet[4])
    duration_rows: list[tuple[str, str, int, Sequence[float]]] = [
        ("全部", "全部任务", analysis.total_rows, analysis.duration_values)
    ]
    duration_rows.extend(
        ("模型", name, metrics.count, metrics.durations)
        for name, metrics in sorted(
            analysis.model_metrics.items(), key=lambda item: (-item[1].count, item[0])
        )
    )
    for dimension, name, count, values in duration_rows:
        duration_sheet.append(
            [
                dimension,
                name,
                count,
                len(values),
                share(len(values), count),
                average(values),
                linear_percentile(values, 0.50),
                linear_percentile(values, 0.90),
                linear_percentile(values, 0.95),
                linear_percentile(values, 0.99),
            ]
        )
    for row in duration_sheet.iter_rows(min_row=5, max_row=duration_sheet.max_row):
        row[4].number_format = "0.0%"
        for cell in row[5:10]:
            cell.number_format = '0.000" 秒"'
    configure_table_sheet(
        duration_sheet,
        last_row=duration_sheet.max_row,
        last_column=10,
        widths=(14, 32, 12, 16, 12, 14, 14, 14, 14, 14),
    )

    for sheet_name, title, metrics_map in (
        ("模型分布", "模型使用分布", analysis.model_metrics),
    ):
        sheet = workbook.create_sheet(sheet_name)
        add_sheet_title(sheet, title)
        sheet.append([])
        sheet.append(group_headers)
        style_table_header(sheet[4])
        for name, metrics in sorted(
            metrics_map.items(), key=lambda item: (-item[1].count, item[0])
        ):
            sheet.append(group_row(name, metrics, analysis.total_rows))
        format_group_sheet(sheet)
        configure_table_sheet(
            sheet,
            last_row=sheet.max_row,
            last_column=12,
            widths=(32, 12, 10, 12, 12, 15, 14, 14, 17, 15, 15, 14),
        )

    version_sheet = workbook.create_sheet("版本分布")
    add_sheet_title(version_sheet, "AstronStudio 与 AstronCode 版本分布")
    version_sheet.append([])
    version_sheet.append(["产品", "来源字段"] + group_headers)
    style_table_header(version_sheet[4])
    for product, source_field, _ in VERSION_DIMENSIONS:
        version_rows = [
            (key[2], metrics)
            for key, metrics in analysis.version_metrics.items()
            if key[0] == product and key[1] == source_field
        ]
        for version, metrics in sorted(
            version_rows, key=lambda item: (-item[1].count, item[0])
        ):
            version_sheet.append(
                [product, source_field]
                + group_row(version, metrics, analysis.total_rows)
            )
    for row in version_sheet.iter_rows(min_row=5, max_row=version_sheet.max_row):
        row[4].number_format = "0.0%"
        for cell in (row[8], row[9]):
            cell.number_format = '0.000" 秒"'
        for cell in (row[11], row[12]):
            cell.number_format = "#,##0"
        row[13].number_format = "0.0%"
    configure_table_sheet(
        version_sheet,
        last_row=version_sheet.max_row,
        last_column=14,
        widths=(18, 28, 28, 12, 10, 12, 12, 15, 14, 14, 17, 15, 15, 14),
    )

    effort_sheet = workbook.create_sheet("推理强度")
    add_sheet_title(effort_sheet, "推理强度分布")
    effort_sheet.append([])
    effort_sheet.append(group_headers)
    style_table_header(effort_sheet[4])
    for name, metrics in sorted(
        analysis.effort_metrics.items(), key=lambda item: (-item[1].count, item[0])
    ):
        effort_sheet.append(group_row(name, metrics, analysis.total_rows))
    format_group_sheet(effort_sheet)
    configure_table_sheet(
        effort_sheet,
        last_row=effort_sheet.max_row,
        last_column=12,
        widths=(32, 12, 10, 12, 12, 15, 14, 14, 17, 15, 15, 14),
    )

    token_sheet = workbook.create_sheet("Token分布")
    add_sheet_title(token_sheet, "Token 使用分布")
    token_sheet.append([])
    token_sheet.append(
        ["Token类型", "有效记录数", "覆盖率", "平均值", "P50", "P90", "P95", "P99"]
    )
    style_table_header(token_sheet[4])
    for field_name, label in TOKEN_FIELDS:
        values = analysis.token_values[field_name]
        token_sheet.append(
            [
                label,
                len(values),
                share(len(values), analysis.total_rows),
                average(values),
                linear_percentile(values, 0.50),
                linear_percentile(values, 0.90),
                linear_percentile(values, 0.95),
                linear_percentile(values, 0.99),
            ]
        )
    for row in token_sheet.iter_rows(min_row=5, max_row=token_sheet.max_row):
        row[2].number_format = "0.0%"
        for cell in row[3:8]:
            cell.number_format = "#,##0"
    configure_table_sheet(
        token_sheet,
        last_row=token_sheet.max_row,
        last_column=8,
        widths=(24, 18, 12, 16, 14, 14, 14, 14),
    )

    status_sheet = workbook.create_sheet("状态分布")
    add_sheet_title(status_sheet, "任务状态分布")
    status_sheet.append([])
    status_sheet.append(["状态类型", "状态值", "任务数", "占比"])
    style_table_header(status_sheet[4])
    for key, metrics in sorted(
        analysis.status_metrics.items(),
        key=lambda item: (item[0].split("\t", 1)[0], -item[1].count, item[0]),
    ):
        status_type, value = key.split("\t", 1)
        status_sheet.append(
            [
                status_type,
                value,
                metrics.count,
                share(metrics.count, analysis.total_rows),
            ]
        )
    for row in status_sheet.iter_rows(min_row=5, max_row=status_sheet.max_row):
        row[3].number_format = "0.0%"
    configure_table_sheet(
        status_sheet,
        last_row=status_sheet.max_row,
        last_column=4,
        widths=(18, 30, 14, 12),
    )

    daily_sheet = workbook.create_sheet("日期趋势")
    add_sheet_title(daily_sheet, "日期趋势")
    daily_sheet.append([])
    daily_sheet.append(
        [
            "日期",
            "任务数",
            "占比",
            "用户数",
            "会话数",
            "耗时P50",
            "耗时P90",
            "总Token P50",
            "总Token P90",
            "Turn状态已记录数",
            "Turn成功率",
        ]
    )
    style_table_header(daily_sheet[4])
    for day, metrics in sorted(analysis.daily_metrics.items()):
        daily_sheet.append(
            [
                datetime.combine(day, datetime_time.min),
                metrics.count,
                share(metrics.count, analysis.total_rows),
                len(metrics.users),
                len(metrics.sessions),
                linear_percentile(metrics.durations, 0.50),
                linear_percentile(metrics.durations, 0.90),
                linear_percentile(metrics.total_tokens, 0.50),
                linear_percentile(metrics.total_tokens, 0.90),
                metrics.known_status_count,
                success_rate(metrics),
            ]
        )
    for row in daily_sheet.iter_rows(min_row=5, max_row=daily_sheet.max_row):
        row[0].number_format = "yyyy-mm-dd"
        row[2].number_format = "0.0%"
        row[5].number_format = '0.000" 秒"'
        row[6].number_format = '0.000" 秒"'
        row[7].number_format = "#,##0"
        row[8].number_format = "#,##0"
        row[10].number_format = "0.0%"
    configure_table_sheet(
        daily_sheet,
        last_row=daily_sheet.max_row,
        last_column=11,
        widths=(15, 12, 10, 12, 12, 14, 14, 15, 15, 20, 14),
    )

    quality_sheet = workbook.create_sheet("数据质量")
    add_sheet_title(quality_sheet, "字段完整性")
    quality_sheet.append([])
    quality_sheet.append(["字段", "标签", "已记录数", "缺失数", "覆盖率", "说明"])
    style_table_header(quality_sheet[4])
    for field_name, label, note in QUALITY_FIELDS:
        present = analysis.field_non_missing[field_name]
        quality_sheet.append(
            [
                field_name,
                label,
                present,
                analysis.total_rows - present,
                share(present, analysis.total_rows),
                note,
            ]
        )
    for row in quality_sheet.iter_rows(min_row=5, max_row=quality_sheet.max_row):
        row[4].number_format = "0.0%"
        row[5].alignment = Alignment(vertical="top", wrap_text=True)
    configure_table_sheet(
        quality_sheet,
        last_row=quality_sheet.max_row,
        last_column=6,
        widths=(28, 24, 16, 14, 12, 52),
    )

    methods_sheet = workbook.create_sheet("口径说明")
    methods_sheet.sheet_properties.tabColor = "A6A6A6"
    add_sheet_title(methods_sheet, "分析口径和字段说明")
    methods_sheet.append([])
    methods_sheet.append(["项目", "内容"])
    style_table_header(methods_sheet[4])
    method_rows = [
        ("数据源", input_path.name),
        ("样本单位", "去重后的 agent.turn"),
        (
            "去重",
            "Trace ID + Span ID；缺失时回退 _index + _id，在 Parquet 中仅保留加盐哈希",
        ),
        ("缺失值", "缺失的耗时、Token 和状态不按 0 计"),
        ("分位数", "线性插值，仅使用有效非负数"),
        (
            "场景分类",
            (
                f"{analysis.classification_version}；规则 {category_version} 优先，"
                "规则未命中且非上下文依赖的唯一 Query 使用大模型语义分类；"
                "初次未判定但具有明确任务意图的 Query 可进入二次分类，"
                "产品可人工改分类"
            ),
        ),
        (
            "分类复用",
            analysis.classification_source_name
            or "本次未加载语义分类结果，使用纯规则分类",
        ),
        (
            "语义分类",
            (
                f"模型 {analysis.classification_model}；提示词 "
                f"{analysis.classification_prompt_version}；置信度阈值 "
                f"{analysis.classification_confidence_threshold:.0%}"
                if analysis.classification_model
                and analysis.classification_confidence_threshold is not None
                else "未使用"
            ),
        ),
        (
            "二次分类",
            (
                f"提示词 {analysis.classification_secondary_prompt_version}；候选策略 "
                f"{analysis.classification_secondary_selector_version}；最短 Query "
                f"{analysis.classification_secondary_min_query_length} 个字符；"
                "与初次分类使用相同置信度阈值"
                if analysis.classification_secondary_prompt_version
                and analysis.classification_secondary_min_query_length is not None
                else "未使用"
            ),
        ),
        (
            "上下文依赖",
            "“继续”“按照上面调整”等缺少独立语义的短指令保持未分类，不猜测不可见前文",
        ),
        (
            "远程内容保护",
            "初次分类仅发送规则未命中的唯一 Query；二次分类仅重发符合候选策略的初次未判定 Query；发送前自动脱敏并截断至 2,000 字符，分类缓存不保存 Query 正文",
        ),
        (
            "版本分布",
            "AstronStudio 使用 astron.desktop.version；AstronCode 使用 acode.cli_version；缺失值单列为“未记录”",
        ),
        (
            "Query 去重",
            "Unicode NFKC + 合并空白 + 不区分大小写；对唯一 Query 分类后映射回全部 agent.turn",
        ),
        ("隐私", "Query 仅保留在受控本地 Parquet；产品候选表使用自动脱敏 Query"),
    ]
    for row in method_rows:
        methods_sheet.append(row)
    methods_sheet.append([])
    methods_sheet.append(["Parquet字段", "中文含义", "说明"])
    dictionary_header_row = methods_sheet.max_row
    style_table_header(methods_sheet[dictionary_header_row])
    for field_name, label, note in NORMALIZED_FIELD_SPECS:
        methods_sheet.append([field_name, label, note])
    for row in methods_sheet.iter_rows(min_row=5, max_row=methods_sheet.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    methods_sheet.sheet_view.showGridLines = False
    methods_sheet.freeze_panes = f"A{dictionary_header_row + 1}"
    for column, width in {"A": 30, "B": 32, "C": 66}.items():
        methods_sheet.column_dimensions[column].width = width

    workbook.properties.title = "AstronCode 生产使用数据分析"
    workbook.properties.subject = (
        "场景、Query 长度、耗时、模型、版本、Token、状态和数据质量"
    )
    workbook.save(output_path)
    restrict_file_permissions(output_path)
    verify_analysis_workbook(output_path, analysis)


CANDIDATE_HEADERS = [
    "候选ID",
    "Query",
    "Query长度",
    "长度梯度",
    "最终自动分类",
    "分类方式",
    "分类置信度",
    "分类模型",
    "提示词版本",
    "产品确认分类",
    "出现次数",
    "会话数",
    "首次时间",
    "最后时间",
    "主要模型",
    "耗时P50（秒）",
    "耗时P90（秒）",
    "总Token P50",
    "总Token P90",
    "含图片任务数",
    "输入图片总数",
    "脱敏状态",
    "自动脱敏项",
    "产品选择状态",
    "不选原因",
    "产品备注",
    "目标评测用例ID",
]
CANDIDATE_COLUMN_WIDTHS = (
    22,
    84,
    12,
    14,
    28,
    18,
    14,
    24,
    28,
    28,
    12,
    12,
    20,
    20,
    28,
    16,
    16,
    16,
    16,
    16,
    16,
    18,
    24,
    18,
    28,
    34,
    24,
)


def candidate_id(query_key: str) -> str:
    return "ACQ-" + hashlib.sha256(query_key.encode("utf-8")).hexdigest()[:16]


def dominant_model(models: Counter[str]) -> str:
    if not models:
        return "未记录"
    return min(models.items(), key=lambda item: (-item[1], item[0]))[0]


def candidate_sort_key(item: tuple[str, CandidateAggregate]) -> tuple[Any, ...]:
    query_key, candidate = item
    return (
        -candidate.count,
        -len(candidate.sessions),
        -candidate.query_length,
        query_key,
    )


def write_candidate_workbook(
    analysis: DatasetAnalysis,
    output_path: Path,
    *,
    input_path: Path,
    category_version: str,
    rules: Sequence[CategoryRule],
) -> None:
    ensure_suffix(output_path, ".xlsx", "产品候选表输出路径")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("候选Query")
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A5"
    sheet.sheet_properties.tabColor = "1F4E78"
    for index, width in enumerate(CANDIDATE_COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.append([])
    title = WriteOnlyCell(sheet, value="AstronCode 评测集产品候选 Query")
    style_title(title)
    sheet.append([title])
    sheet.append([])
    header_cells = [WriteOnlyCell(sheet, value=value) for value in CANDIDATE_HEADERS]
    style_table_header(header_cells)
    sheet.append(header_cells)
    editable_fill = PatternFill("solid", fgColor="FFF2CC")
    query_count = 0
    truncated_count = 0
    for query_key, candidate in sorted(
        analysis.candidates.items(), key=candidate_sort_key
    ):
        redacted, redaction_labels = redact_query(candidate.representative_query)
        redacted, truncated = safe_excel_text(redacted)
        truncated_count += int(truncated)
        privacy_status = "已自动脱敏" if redaction_labels else "未发现敏感模式"
        if truncated:
            privacy_status += "；Excel已截断"
        values: list[Any] = [
            candidate_id(query_key),
            redacted,
            candidate.query_length,
            query_length_band(candidate.query_length),
            candidate.category,
            candidate.classification_method,
            candidate.classification_confidence,
            candidate.classification_model,
            candidate.classification_prompt_version,
            "",
            candidate.count,
            len(candidate.sessions),
            candidate.first_time,
            candidate.last_time,
            dominant_model(candidate.models),
            linear_percentile(candidate.durations, 0.50),
            linear_percentile(candidate.durations, 0.90),
            linear_percentile(candidate.total_tokens, 0.50),
            linear_percentile(candidate.total_tokens, 0.90),
            candidate.image_task_count,
            candidate.input_image_count,
            privacy_status,
            "、".join(redaction_labels),
            "未评审",
            "",
            "",
            "",
        ]
        cells: list[WriteOnlyCell] = []
        for column_index, value in enumerate(values, start=1):
            cell = WriteOnlyCell(sheet, value=value)
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(
                vertical="top" if column_index == 2 else "center"
            )
            if column_index == 2:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            if column_index in {10, 24, 25, 26, 27}:
                cell.fill = editable_fill
            if column_index == 7:
                cell.number_format = "0.0%"
            elif column_index in {13, 14}:
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            elif column_index in {16, 17}:
                cell.number_format = '0.000" 秒"'
            elif column_index in {3, 11, 12, 18, 19, 20, 21}:
                cell.number_format = "#,##0"
            cells.append(cell)
        sheet.append(cells)
        query_count += 1

    last_row = max(4, 4 + query_count)
    sheet.auto_filter.ref = f"A4:AA{last_row}"

    if query_count:
        category_formula = '"' + ",".join(rule.name for rule in rules) + '"'
        category_validation = DataValidation(
            type="list", formula1=category_formula, allow_blank=True
        )
        category_validation.error = "请从场景分类列表中选择"
        category_validation.errorTitle = "分类无效"
        sheet.data_validations.append(category_validation)
        category_validation.add(f"J5:J{last_row}")
        status_validation = DataValidation(
            type="list",
            formula1='"未评审,入选,待定,不选"',
            allow_blank=False,
        )
        status_validation.error = "请选择未评审、入选、待定或不选"
        status_validation.errorTitle = "状态无效"
        sheet.data_validations.append(status_validation)
        status_validation.add(f"X5:X{last_row}")

    guide = workbook.create_sheet("使用说明")
    guide.sheet_view.showGridLines = False
    guide.sheet_properties.tabColor = "A6A6A6"
    guide.column_dimensions["A"].width = 28
    guide.column_dimensions["B"].width = 88
    guide.append([])
    guide_title = WriteOnlyCell(guide, value="产品筛选说明")
    style_title(guide_title)
    guide.append([guide_title])
    guide.append([])
    guide_header = [WriteOnlyCell(guide, value=value) for value in ("项目", "内容")]
    style_table_header(guide_header)
    guide.append(guide_header)
    guide_rows = [
        ("标准化数据", input_path.name),
        ("候选数", len(analysis.candidates)),
        ("场景分类版本", analysis.classification_version),
        ("场景规则版本", category_version),
        (
            "场景分类数据",
            analysis.classification_source_name or "未使用（纯规则）",
        ),
        (
            "语义分类模型",
            analysis.classification_model or "未使用",
        ),
        (
            "未分类口径",
            "空 Query、上下文依赖短指令、两轮大模型仍未判定、不满足二次分类条件或低于置信度阈值的 Query 保持未分类",
        ),
        (
            "黄色列",
            "产品确认分类、产品选择状态、不选原因、产品备注、目标评测用例 ID 可编辑",
        ),
        (
            "Query 去重",
            "Unicode NFKC + 合并空白 + 不区分大小写；唯一 Query 分类后映射回全部记录，不做语义聚类",
        ),
        (
            "自动脱敏",
            "覆盖凭据 URL、密钥、Bearer Token、邮箱、IPv4、身份证、手机号和用户主目录；入选前仍需人工复核",
        ),
        (
            "Excel 截断",
            f"{truncated_count} 条候选 Query 超过 32,767 字符；候选表保留可见前缀并标记“Excel已截断”，完整原文仅在受控 Parquet 中",
        ),
        (
            "建议筛选顺序",
            "先按出现次数、会话数和场景分布查看，再结合耗时、Token 和图片输入判断评测价值",
        ),
    ]
    for key, value in guide_rows:
        cells = [WriteOnlyCell(guide, value=key), WriteOnlyCell(guide, value=value)]
        cells[0].font = Font(name="Arial", size=10, bold=True)
        cells[1].font = Font(name="Arial", size=10)
        cells[1].alignment = Alignment(vertical="top", wrap_text=True)
        guide.append(cells)

    workbook.properties.title = "AstronCode 评测集产品候选 Query"
    workbook.properties.subject = "去重脱敏后的产品评审候选数据"
    workbook.save(output_path)
    restrict_file_permissions(output_path)
    verify_candidate_workbook(output_path, query_count)


def verify_analysis_workbook(path: Path, analysis: DatasetAnalysis) -> None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        expected_sheets = [
            "分析总览",
            "场景分类",
            "分类方法",
            "Query长度",
            "耗时分布",
            "模型分布",
            "版本分布",
            "推理强度",
            "Token分布",
            "状态分布",
            "日期趋势",
            "数据质量",
            "口径说明",
        ]
        if workbook.sheetnames != expected_sheets:
            raise RuntimeError(f"分析工作簿工作表不匹配：{workbook.sheetnames}")
        scene_total = sum(
            row[1]
            for row in workbook["场景分类"].iter_rows(
                min_row=5, min_col=1, max_col=2, values_only=True
            )
            if row[0]
        )
        classification_method_total = sum(
            row[1]
            for row in workbook["分类方法"].iter_rows(
                min_row=5, min_col=1, max_col=2, values_only=True
            )
            if row[0]
        )
        length_total = sum(
            row[3]
            for row in workbook["Query长度"].iter_rows(
                min_row=5, min_col=1, max_col=4, values_only=True
            )
            if row[0]
        )
        model_total = sum(
            row[1]
            for row in workbook["模型分布"].iter_rows(
                min_row=5, min_col=1, max_col=2, values_only=True
            )
            if row[0]
        )
        version_totals: Counter[str] = Counter()
        for product, _, _, count in workbook["版本分布"].iter_rows(
            min_row=5, min_col=1, max_col=4, values_only=True
        ):
            if product:
                version_totals[str(product)] += int(count)
        if (
            scene_total,
            classification_method_total,
            length_total,
            model_total,
            version_totals["AstronStudio"],
            version_totals["AstronCode"],
        ) != (
            analysis.total_rows,
            analysis.total_rows,
            analysis.total_rows,
            analysis.total_rows,
            analysis.total_rows,
            analysis.total_rows,
        ):
            raise RuntimeError(
                "分析工作簿汇总无法与 Parquet 样本数对账："
                f"scene={scene_total}, classification_method={classification_method_total}, "
                f"length={length_total}, model={model_total}, "
                f"studio_version={version_totals['AstronStudio']}, "
                f"cli_version={version_totals['AstronCode']}, "
                f"parquet={analysis.total_rows}"
            )
    finally:
        workbook.close()


def verify_candidate_workbook(path: Path, expected_rows: int) -> None:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if workbook.sheetnames != ["候选Query", "使用说明"]:
            raise RuntimeError(f"产品候选工作簿工作表不匹配：{workbook.sheetnames}")
        sheet = workbook["候选Query"]
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=4, max_row=4))]
        if headers != CANDIDATE_HEADERS:
            raise RuntimeError("产品候选表头不匹配")
    finally:
        workbook.close()
    row_count = 0
    formula_found = False
    with zipfile.ZipFile(path) as archive:  # noqa: SIM117
        with archive.open("xl/worksheets/sheet1.xml") as worksheet_xml:
            tail = b""
            while True:
                chunk = worksheet_xml.read(1024 * 1024)
                if not chunk:
                    break
                data = tail + chunk
                stable = data[:-32] if len(data) > 32 else b""
                row_count += len(re.findall(rb"<row\b", stable))
                formula_found = formula_found or bool(re.search(rb"<f(?:\s|>)", stable))
                tail = data[-32:]
            row_count += len(re.findall(rb"<row\b", tail))
            formula_found = formula_found or bool(re.search(rb"<f(?:\s|>)", tail))
    actual_rows = max(0, row_count - 4)
    if actual_rows != expected_rows:
        raise RuntimeError(
            f"产品候选行数校验失败：期望 {expected_rows}，实际 {actual_rows}"
        )
    if formula_found:
        raise RuntimeError("产品候选表存在未预期公式")


def window_label(window: TimeWindow) -> tuple[str, str]:
    timezone_value = ZoneInfo(DEFAULT_TIMEZONE)
    start = window.start.astimezone(timezone_value).strftime("%Y%m%d")
    end = (window.end.astimezone(timezone_value) - timedelta(microseconds=1)).strftime(
        "%Y%m%d"
    )
    return start, end


def default_output_path(
    prefix: str, window: TimeWindow, output_dir: Path, suffix: str
) -> Path:
    start, end = window_label(window)
    return output_dir / f"{prefix}_{start}_{end}{suffix}"


def ensure_suffix(path: Path, suffix: str, label: str) -> None:
    if not str(path).lower().endswith(suffix.lower()):
        raise ValueError(f"{label}必须使用 {suffix} 扩展名：{path}")


def validate_fetch_args(args: argparse.Namespace) -> None:
    if not args.es_url:
        raise ValueError("缺少 Elasticsearch 地址，请传入 --es-url 或设置 ES_URL")
    if not 50 <= args.page_size <= 2_000:
        raise ValueError("--page-size 必须在 50 到 2000 之间")
    if args.request_interval < 0.05:
        raise ValueError("--request-interval 不能小于 0.05 秒")
    if args.max_records < 1:
        raise ValueError("--max-records 必须大于 0")
    if args.fallback_days < 1:
        raise ValueError("--fallback-days 必须大于 0")
    if args.timeout <= 0 or args.retries < 0:
        raise ValueError("--timeout 必须大于 0，--retries 不能小于 0")


def fetch_to_dataset(
    args: argparse.Namespace,
) -> tuple[DatasetPaths, FetchStats, TimeWindow]:
    validate_fetch_args(args)
    requested = requested_window(args.start, args.end, args.timezone)
    client = ElasticsearchClient(
        args.es_url,
        timeout_seconds=args.timeout,
        retries=args.retries,
        ca_cert=args.ca_cert,
    )

    def count_for(window: TimeWindow) -> int:
        query = build_event_query(
            window,
            event_field=args.event_field,
            event_value=args.event_value,
            time_field=args.time_field,
        )
        return count_matching_records(client, index=args.index, query=query)

    effective, estimated_count, fallback_applied = choose_effective_window(
        requested,
        count_records=count_for,
        max_records=args.max_records,
        fallback_days=args.fallback_days,
        auto_fallback=not args.no_auto_fallback,
    )
    if fallback_applied:
        print(
            f"命中数超过 {args.max_records:,}，已按规则缩短为最近 {args.fallback_days} 天",
            file=sys.stderr,
        )
    paths = DatasetPaths(
        snapshot=args.snapshot_output
        or default_output_path(
            "astroncode_agent_turn_snapshot",
            effective,
            args.output_dir,
            ".jsonl.gz",
        ),
        parquet=args.parquet_output
        or default_output_path(
            "astroncode_usage_normalized", effective, args.output_dir, ".parquet"
        ),
    )
    query = build_event_query(
        effective,
        event_field=args.event_field,
        event_value=args.event_value,
        time_field=args.time_field,
    )
    hits = iter_matching_hits(
        client,
        index=args.index,
        query=query,
        time_field=args.time_field,
        page_size=args.page_size,
        request_interval=args.request_interval,
        pit_keep_alive=args.pit_keep_alive,
        max_records=args.max_records,
    )
    stats = write_dataset(
        hits,
        paths,
        requested=requested,
        effective=effective,
        estimated_count=estimated_count,
        fallback_applied=fallback_applied,
        index=args.index,
        event_field=args.event_field,
        event_value=args.event_value,
        timezone_name=args.timezone,
    )
    return paths, stats, effective


def derived_analysis_paths(input_path: Path) -> tuple[Path, Path]:
    stem = input_path.stem
    if "_normalized_" in stem:
        analysis_stem = stem.replace("_normalized_", "_analysis_", 1)
        candidates_stem = stem.replace("_normalized_", "_candidates_", 1)
    else:
        analysis_stem = f"{stem}_analysis"
        candidates_stem = f"{stem}_candidates"
    return (
        input_path.with_name(f"{analysis_stem}.xlsx"),
        input_path.with_name(f"{candidates_stem}.xlsx"),
    )


def analyze_to_workbooks(
    input_path: Path,
    analysis_output: Path,
    candidates_output: Path,
    category_rules_path: Path,
    scene_classifications_path: Path | None = None,
) -> DatasetAnalysis:
    ensure_suffix(input_path, ".parquet", "标准化数据输入路径")
    ensure_suffix(analysis_output, ".xlsx", "分析结果输出路径")
    ensure_suffix(candidates_output, ".xlsx", "产品候选表输出路径")
    resolved = {
        input_path.resolve(),
        analysis_output.resolve(),
        candidates_output.resolve(),
    }
    if len(resolved) != 3:
        raise ValueError("输入 Parquet、分析 Excel 和产品候选 Excel 路径必须不同")
    category_version, rules = load_category_rules(category_rules_path)
    scene_classifications = (
        load_scene_classifications(scene_classifications_path, rules)
        if scene_classifications_path is not None
        else None
    )
    if (
        scene_classifications is not None
        and scene_classifications.rule_version != category_version
    ):
        raise ValueError(
            "场景分类数据的规则版本与当前分类规则不一致："
            f"{scene_classifications.rule_version} != {category_version}"
        )
    analysis = analyze_parquet(input_path, rules, scene_classifications)
    analysis.classification_rule_version = category_version
    write_analysis_workbook(
        analysis,
        analysis_output,
        input_path=input_path,
        category_version=category_version,
        rules=rules,
    )
    write_candidate_workbook(
        analysis,
        candidates_output,
        input_path=input_path,
        category_version=category_version,
        rules=rules,
    )
    return analysis


def add_fetch_arguments(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("--es-url", default=os.environ.get("ES_URL"))
    parser.add_argument("--index", default=os.environ.get("ES_INDEX", DEFAULT_INDEX))
    parser.add_argument("--start", help="起始时间；ISO 8601 或 YYYY-MM-DD")
    parser.add_argument(
        "--end", help="截止时间；日期写法包含该日，时间写法按该时刻排除"
    )
    parser.add_argument(
        "--timezone", default=DEFAULT_TIMEZONE, help="日期输入和 Excel 展示时区"
    )
    parser.add_argument("--event-field", default=DEFAULT_EVENT_FIELD)
    parser.add_argument("--event-value", default=DEFAULT_EVENT_VALUE)
    parser.add_argument("--time-field", default=DEFAULT_TIME_FIELD)
    parser.add_argument("--page-size", type=int, default=DEFAULT_PAGE_SIZE)
    parser.add_argument(
        "--request-interval", type=float, default=DEFAULT_REQUEST_INTERVAL
    )
    parser.add_argument("--max-records", type=int, default=DEFAULT_MAX_RECORDS)
    parser.add_argument("--fallback-days", type=int, default=DEFAULT_FALLBACK_DAYS)
    parser.add_argument("--no-auto-fallback", action="store_true")
    parser.add_argument("--pit-keep-alive", default="2m")
    parser.add_argument("--timeout", type=float, default=30.0)
    parser.add_argument("--retries", type=int, default=4)
    parser.add_argument("--ca-cert", type=Path)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--snapshot-output", type=Path)
    parser.add_argument("--parquet-output", type=Path)


def add_analysis_arguments(
    parser: argparse.ArgumentParser, *, input_required: bool
) -> None:
    if input_required:
        parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--analysis-output", type=Path)
    parser.add_argument("--candidates-output", type=Path)
    parser.add_argument("--category-rules", type=Path, default=DEFAULT_CATEGORY_RULES)
    parser.add_argument(
        "--scene-classifications",
        type=Path,
        help="可复用的混合场景分类 Parquet；省略时使用纯规则分类",
    )


def derived_scene_classification_path(input_path: Path) -> Path:
    stem = input_path.stem
    if "_normalized_" in stem:
        stem = stem.replace("_normalized_", "_scene_classifications_", 1)
    else:
        stem = f"{stem}_scene_classifications"
    return input_path.with_name(f"{stem}_v3.parquet")


def derived_refined_scene_classification_path(input_path: Path) -> Path:
    stem = input_path.stem
    if "_normalized_" in stem:
        stem = stem.replace("_normalized_", "_scene_classifications_", 1)
    else:
        stem = f"{stem}_scene_classifications"
    return input_path.with_name(f"{stem}_v4.parquet")


def add_scene_classification_arguments(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--checkpoint", type=Path)
    parser.add_argument("--category-rules", type=Path, default=DEFAULT_CATEGORY_RULES)
    parser.add_argument(
        "--base-url",
        default=os.environ.get("SCENE_LLM_BASE_URL")
        or os.environ.get("OPENROUTER_BASE_URL"),
    )
    parser.add_argument(
        "--api-key-env",
        default=os.environ.get("SCENE_LLM_API_KEY_ENV", "OPENROUTER_API_KEY"),
        help="保存分类服务密钥的环境变量名；不会把密钥写入参数或产物",
    )
    parser.add_argument(
        "--model", default=os.environ.get("SCENE_LLM_MODEL", DEFAULT_SCENE_MODEL)
    )
    parser.add_argument("--confidence-threshold", type=float, default=0.70)
    parser.add_argument("--batch-size", type=int, default=DEFAULT_SCENE_BATCH_SIZE)
    parser.add_argument(
        "--batch-char-limit", type=int, default=DEFAULT_SCENE_BATCH_CHAR_LIMIT
    )
    parser.add_argument("--workers", type=int, default=DEFAULT_SCENE_WORKERS)
    parser.add_argument(
        "--request-interval",
        type=float,
        default=DEFAULT_SCENE_REQUEST_INTERVAL,
    )
    parser.add_argument("--timeout", type=float, default=90.0)
    parser.add_argument("--retries", type=int, default=3)
    parser.add_argument("--resume", action="store_true")
    parser.add_argument(
        "--allow-remote-query-content",
        action="store_true",
        help="确认允许把自动脱敏且截断后的规则未命中 Query 发送给分类模型",
    )


def add_scene_refinement_arguments(parser: argparse.ArgumentParser) -> None:
    add_scene_classification_arguments(parser)
    parser.add_argument(
        "--base-classifications",
        type=Path,
        required=True,
        help="第一次混合分类生成的场景分类 Parquet",
    )
    parser.add_argument(
        "--secondary-min-query-length",
        type=int,
        default=DEFAULT_SECONDARY_MIN_QUERY_LENGTH,
        help="二次分类候选 Query 的最短规范化字符数",
    )


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="低并发导出并分析 AstronCode 生产 agent.turn 使用数据"
    )
    subparsers = parser.add_subparsers(dest="command", required=True)

    fetch_parser = subparsers.add_parser(
        "fetch", help="从 Elasticsearch 导出原始快照和标准化 Parquet"
    )
    add_fetch_arguments(fetch_parser)

    analyze_parser = subparsers.add_parser(
        "analyze", help="基于本地 Parquet 生成汇总和产品候选 Excel"
    )
    add_analysis_arguments(analyze_parser, input_required=True)

    classify_parser = subparsers.add_parser(
        "classify-scenes", help="对本地 Parquet 生成可复用的混合场景分类"
    )
    add_scene_classification_arguments(classify_parser)

    refine_parser = subparsers.add_parser(
        "refine-scenes", help="复用第一次分类缓存，对未判定 Query 做二次分类"
    )
    add_scene_refinement_arguments(refine_parser)

    run_parser = subparsers.add_parser("run", help="连续执行导出和分析")
    add_fetch_arguments(run_parser)
    add_analysis_arguments(run_parser, input_required=False)
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        if args.command == "fetch":
            paths, stats, effective = fetch_to_dataset(args)
            print(f"原始快照：{paths.snapshot}")
            print(f"标准化数据：{paths.parquet}")
            print(
                f"实际窗口：{iso_utc(effective.start)} 至 {iso_utc(effective.end)}（截止不含）"
            )
            print(
                f"ES 返回 {stats.fetched_hits:,} 条，去重后标准化 {stats.exported_rows:,} 条"
            )
            return 0

        if args.command == "analyze":
            default_analysis, default_candidates = derived_analysis_paths(args.input)
            analysis_output = args.analysis_output or default_analysis
            candidates_output = args.candidates_output or default_candidates
            analysis = analyze_to_workbooks(
                args.input,
                analysis_output,
                candidates_output,
                args.category_rules,
                args.scene_classifications,
            )
            print(f"分析汇总：{analysis_output}")
            print(f"产品候选：{candidates_output}")
            print(
                f"样本数：{analysis.total_rows:,}，候选 Query：{len(analysis.candidates):,}"
            )
            return 0

        if args.command == "classify-scenes":
            if not args.allow_remote_query_content:
                raise ValueError(
                    "远程语义分类需要显式传入 --allow-remote-query-content"
                )
            if not 0.0 <= args.confidence_threshold <= 1.0:
                raise ValueError("--confidence-threshold 必须在 0 到 1 之间")
            if args.batch_size < 1 or args.batch_char_limit < 1:
                raise ValueError("--batch-size 和 --batch-char-limit 必须为正数")
            if args.workers < 1 or args.workers > 4:
                raise ValueError("--workers 必须在 1 到 4 之间")
            api_key = os.environ.get(args.api_key_env)
            if not api_key:
                raise ValueError(f"环境变量 {args.api_key_env} 未配置")
            output_path = args.output or derived_scene_classification_path(args.input)
            checkpoint_path = args.checkpoint or Path(f"{output_path}.checkpoint.jsonl")
            stats = classify_scenes_to_parquet(
                args.input,
                output_path,
                checkpoint_path,
                args.category_rules,
                base_url=args.base_url or "",
                api_key=api_key,
                model=args.model,
                confidence_threshold=args.confidence_threshold,
                batch_size=args.batch_size,
                batch_char_limit=args.batch_char_limit,
                workers=args.workers,
                request_interval=args.request_interval,
                timeout_seconds=args.timeout,
                retries=args.retries,
                resume=args.resume,
            )
            print(f"场景分类数据：{output_path}")
            print(
                "唯一 Query：{unique_queries:,}，规则：{rule_queries:,}，"
                "上下文依赖：{context_dependent_queries:,}，语义分类："
                "{semantic_queries:,}，最终未分类：{unclassified_queries:,}".format(
                    **stats
                )
            )
            return 0

        if args.command == "refine-scenes":
            if not args.allow_remote_query_content:
                raise ValueError(
                    "远程语义分类需要显式传入 --allow-remote-query-content"
                )
            if not 0.0 <= args.confidence_threshold <= 1.0:
                raise ValueError("--confidence-threshold 必须在 0 到 1 之间")
            if args.secondary_min_query_length < 1:
                raise ValueError("--secondary-min-query-length 必须为正数")
            if args.batch_size < 1 or args.batch_char_limit < 1:
                raise ValueError("--batch-size 和 --batch-char-limit 必须为正数")
            if args.workers < 1 or args.workers > 4:
                raise ValueError("--workers 必须在 1 到 4 之间")
            api_key = os.environ.get(args.api_key_env)
            if not api_key:
                raise ValueError(f"环境变量 {args.api_key_env} 未配置")
            output_path = args.output or derived_refined_scene_classification_path(
                args.input
            )
            checkpoint_path = args.checkpoint or Path(f"{output_path}.checkpoint.jsonl")
            stats = refine_scene_classifications_to_parquet(
                args.input,
                args.base_classifications,
                output_path,
                checkpoint_path,
                args.category_rules,
                base_url=args.base_url or "",
                api_key=api_key,
                model=args.model,
                confidence_threshold=args.confidence_threshold,
                minimum_query_length=args.secondary_min_query_length,
                batch_size=args.batch_size,
                batch_char_limit=args.batch_char_limit,
                workers=args.workers,
                request_interval=args.request_interval,
                timeout_seconds=args.timeout,
                retries=args.retries,
                resume=args.resume,
            )
            print(f"二次分类数据：{output_path}")
            print(
                "唯一 Query：{unique_queries:,}，复用：{preserved_queries:,}，"
                "新增上下文依赖：{new_context_queries:,}，二次候选："
                "{secondary_candidate_queries:,}，二次命中："
                "{secondary_classified_queries:,}，二次低置信度："
                "{secondary_low_confidence_queries:,}，二次未判定："
                "{secondary_unjudged_queries:,}，最终未分类："
                "{unclassified_queries:,}".format(**stats)
            )
            return 0

        paths, fetch_stats, effective = fetch_to_dataset(args)
        analysis_output = args.analysis_output or default_output_path(
            "astroncode_usage_analysis", effective, args.output_dir, ".xlsx"
        )
        candidates_output = args.candidates_output or default_output_path(
            "astroncode_usage_candidates", effective, args.output_dir, ".xlsx"
        )
        analysis = analyze_to_workbooks(
            paths.parquet,
            analysis_output,
            candidates_output,
            args.category_rules,
            args.scene_classifications,
        )
        print(f"原始快照：{paths.snapshot}")
        print(f"标准化数据：{paths.parquet}")
        print(f"分析汇总：{analysis_output}")
        print(f"产品候选：{candidates_output}")
        print(
            f"ES 返回 {fetch_stats.fetched_hits:,} 条，去重后分析 "
            f"{analysis.total_rows:,} 条，候选 Query {len(analysis.candidates):,} 条"
        )
        return 0
    except (
        OSError,
        TypeError,
        ValueError,
        RuntimeError,
        json.JSONDecodeError,
        re.error,
        pa.ArrowException,
    ) as error:
        print(f"执行失败：{error}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
