#!/usr/bin/env python3
"""跨轮次对比报告：同一 (模型, Harness) 单元在多个 round 之间的维度对比。

与 generate_eval_report.py 的分工：
- generate_eval_report.py 出**单轮**正式报告，前提是所有 unit 属于同一 round
  （它显式校验并在跨 round 时退出），unit 标识为 `<model>@<harness>`。
- 本脚本出**跨轮**对比报告：同一 unit 在不同 round 各占一行，标识为
  `<model>@<harness> (<round>)`，并支持把用例范围收敛到可比子集。

典型场景：round4 相比 round3 新增了 7 个自定义用例（tag: custom），要和
round3 做同集对比时必须先剔除这些用例，否则分母不同、分数不可比。

用法：
    # 取两轮交集（推荐，最严格）
    python3 tools/report/scripts/generate_round_compare_report.py \
        --rounds eval_out/all_suite/round3_t3600_new eval_out/all_suite/round4_t3600_new \
        --unit xopglm52@astroncode --task-scope intersection

    # 按 tag 排除
    python3 ... --task-scope exclude-tags --exclude-tags custom

    # 显式指定用例清单
    python3 ... --task-scope task-file --task-file ids.txt
"""
from __future__ import annotations

import argparse
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

_SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(_SCRIPT_DIR))

# 直接复用单轮报告脚本的加载与出表逻辑，避免口径漂移
from generate_eval_report import (  # noqa: E402
    DEFAULT_ENTITIES_PATH,
    DIFFICULTY_ORDER,
    MODALITY_ORDER,
    MODALITY_ZH,
    SUITE_DIR_RE,
    UnitResult,
    build_suite_zh_map,
    build_task_order,
    discover_units,
    find_tasks_dir,
    load_all_task_meta,
    load_capability_map,
    style_header_row,
    set_widths,
    apply_pct_format,
    write_capability_sheet,
    write_dimension_sheet_transposed,
)
import report_entities  # noqa: E402


# ===========================================================================
# 用例范围解析
# ===========================================================================

_TAGS_BLOCK_RE = re.compile(r"^tags:\s*$((?:\n\s*-\s*\S+)+)", re.M)
_TAGS_INLINE_RE = re.compile(r"^tags:\s*\[([^\]]*)\]\s*$", re.M)


def task_tags(md_path: Path) -> set[str]:
    """读取任务 md frontmatter 的 tags（支持列表块与内联数组两种写法）。"""
    try:
        head = md_path.read_text(encoding="utf-8")[:2000]
    except OSError:
        return set()
    tags: set[str] = set()
    block = _TAGS_BLOCK_RE.search(head)
    if block:
        tags |= {line.strip().lstrip("-").strip()
                 for line in block.group(1).splitlines() if line.strip()}
    inline = _TAGS_INLINE_RE.search(head)
    if inline:
        tags |= {t.strip().strip("\"'") for t in inline.group(1).split(",") if t.strip()}
    return {t for t in tags if t}


def collect_tagged_tasks(tasks_dir: Path, wanted: set[str]) -> set[str]:
    """返回 tags 命中 wanted 的 task_id 集合（扫描主目录与 extension/）。"""
    hit: set[str] = set()
    for base in (tasks_dir, tasks_dir / "extension",
                 tasks_dir / "cn", tasks_dir / "extension" / "cn"):
        if not base.is_dir():
            continue
        for suite_dir in sorted(base.iterdir()):
            if not suite_dir.is_dir() or not SUITE_DIR_RE.match(suite_dir.name):
                continue
            for md in sorted(suite_dir.glob("*.md")):
                if task_tags(md) & wanted:
                    hit.add(md.stem)
    return hit


# ===========================================================================
# 跨轮 unit 加载
# ===========================================================================

def load_round_unit(round_root: Path, model: str, harness: str,
                    registry, pricing_date: date | None) -> UnitResult:
    """在指定 round 下定位 (model, harness) 单元并加载。"""
    specs = [s for s in discover_units(round_root) if s[0] == model and s[1] == harness]
    if not specs:
        available = sorted({f"{m}@{h}" for m, h, _ in discover_units(round_root)})
        sys.exit(f"错误：round {round_root.name} 中不存在单元 {model}@{harness}；"
                 f"可用：{available}")
    if len(specs) > 1:
        sys.exit(f"错误：round {round_root.name} 中 {model}@{harness} 匹配到多个目录")
    _, _, unit_dir = specs[0]
    return UnitResult(model, harness, unit_dir, registry, pricing_date)


def relabel_unit(unit: UnitResult, round_label: str) -> UnitResult:
    """给 unit 打上轮次标签。

    维度表用 `unit_display` 展示、用 `unit` 作内部 key（见
    write_dimension_sheet_transposed 的 view_values），两者都必须带上轮次，
    否则同一 (model, harness) 在多轮间会互相覆盖。
    """
    unit.round_label = round_label
    unit.unit = f"{unit.model}@{unit.harness} ({round_label})"
    unit.unit_display = f"{unit.model_display}@{unit.harness_display} ({round_label})"
    return unit


def restrict_tasks(unit: UnitResult, keep: set[str]) -> None:
    """把 unit 的用例收敛到 keep 集合内。

    total_pct / avg_pct 均由 self.tasks 现算（无缓存），故过滤后分数自动重算。
    """
    unit.tasks = [t for t in unit.tasks if t.task_id in keep]
    unit.task_map = {t.task_id: t for t in unit.tasks}


def resolve_scope(args, units: list[UnitResult], tasks_dir: Path | None) -> tuple[set[str], str]:
    """返回 (保留的 task_id 集合, 口径说明)。"""
    per_round = {u.round_label: {t.task_id for t in u.tasks} for u in units}
    union: set[str] = set()
    inter: set[str] | None = None
    for ids in per_round.values():
        union |= ids
        inter = ids if inter is None else (inter & ids)
    inter = inter or set()

    if args.task_scope == "all":
        return union, f"全部用例（并集 {len(union)} 例，各轮用例数可能不同、分母不可比）"

    if args.task_scope == "intersection":
        return inter, f"各轮交集 {len(inter)} 例（严格同集对比）"

    if args.task_scope == "exclude-tags":
        if tasks_dir is None:
            sys.exit("错误：--task-scope exclude-tags 需要任务定义目录，请用 --tasks-dir 指定")
        excluded = collect_tagged_tasks(tasks_dir, set(args.exclude_tags))
        keep = union - excluded
        hit = sorted(excluded & union)
        return keep, (f"排除 tag={','.join(args.exclude_tags)} 的 {len(hit)} 例后剩 {len(keep)} 例")

    if args.task_scope == "task-file":
        if not args.task_file:
            sys.exit("错误：--task-scope task-file 需要 --task-file")
        wanted = {line.strip() for line in Path(args.task_file).read_text(
            encoding="utf-8").splitlines() if line.strip() and not line.startswith("#")}
        keep = union & wanted
        missing = wanted - union
        if missing:
            print(f"[警告] --task-file 中 {len(missing)} 个 ID 在结果中不存在，已忽略："
                  f"{sorted(missing)[:5]}", file=sys.stderr)
        return keep, f"按清单指定 {len(keep)} 例"

    sys.exit(f"错误：未知 --task-scope {args.task_scope}")


# ===========================================================================
# 轮次差异 Sheet
# ===========================================================================

def write_round_diff_sheet(wb, units: list[UnitResult],
                           dim_groups: list[tuple[str, list[tuple[str, set[str]]]]],
                           scope_note: str, base_label: str) -> None:
    """逐维度列出各轮取值与相对基准轮的分差（后续轮 − 基准轮）。"""
    ws = wb.create_sheet("轮次差异")
    ws.append([f"对比口径：{scope_note}；基准轮：{base_label}（分差 = 对比轮 − 基准轮）"])
    ws.cell(1, 1).font = Font(bold=True, color="1F4E78")

    base = next((u for u in units if u.round_label == base_label), None)
    others = [u for u in units if u.round_label != base_label]
    if base is None:
        return

    ws.append([])
    header = ["维度", "取值", "例数", f"{base_label}"]
    for u in others:
        header += [u.round_label, "分差"]
    ws.append(header)
    style_header_row(ws)
    header_row = ws.max_row

    def emit(dim_name: str, label: str, ids: set[str] | None) -> None:
        if ids is not None and not ids:
            return
        n = len(ids) if ids is not None else len(base.tasks)
        bv = base.total_pct if ids is None else base.avg_pct(ids)
        row = [dim_name, label, n, round(bv, 1) if bv is not None else "-"]
        for u in others:
            ov = u.total_pct if ids is None else u.avg_pct(ids)
            row.append(round(ov, 1) if ov is not None else "-")
            row.append(round(ov - bv, 1) if (ov is not None and bv is not None) else "-")
        ws.append(row)

    emit("总分", "总平均分", None)
    for dim_name, groups in dim_groups:
        for label, ids in groups:
            emit(dim_name, label, ids)

    widths = {1: 16, 2: 26, 3: 8}
    set_widths(ws, widths, default=14)
    ws.freeze_panes = f"A{header_row + 1}"
    for row in range(header_row + 1, ws.max_row + 1):
        apply_pct_format(ws, row, range(4, 4 + 1 + 2 * len(others)))


def write_compare_metadata_sheet(wb, units: list[UnitResult], round_roots: list[Path],
                                 scope_note: str, keep: set[str], args) -> None:
    """隐藏元数据 Sheet：可反查 raw ID、各轮路径、用例范围与生成参数。"""
    ws = wb.create_sheet("_对比元数据")
    ws.append(["类型", "键", "值"])
    style_header_row(ws)
    ws.append(["对比单元", "raw_unit", args.unit])
    for u in units:
        ws.append(["对比单元", f"{u.round_label}.display", u.unit_display])
        ws.append(["对比单元", f"{u.round_label}.task_count", len(u.tasks)])
        ws.append(["对比单元", f"{u.round_label}.total_pct", round(u.total_pct, 4)])
    for i, root in enumerate(round_roots):
        ws.append(["轮次", f"round[{i}]{'（基准）' if i == 0 else ''}", str(root)])
    ws.append(["用例范围", "task_scope", args.task_scope])
    if args.task_scope == "exclude-tags":
        ws.append(["用例范围", "exclude_tags", ",".join(args.exclude_tags)])
    if args.task_file:
        ws.append(["用例范围", "task_file", args.task_file])
    ws.append(["用例范围", "说明", scope_note])
    ws.append(["用例范围", "生效用例数", len(keep)])
    ws.append(["配置", "dimensions", args.dimensions])
    ws.append(["配置", "entities", args.entities])
    ws.append(["配置", "pricing_date",
               args.pricing_date.isoformat() if args.pricing_date else "-"])
    ws.append(["配置", "生成时间", datetime.now().isoformat(timespec="seconds")])
    for tid in sorted(keep):
        ws.append(["生效用例", tid, ""])
    set_widths(ws, {1: 14, 2: 30, 3: 60})
    ws.sheet_state = "hidden"


def main() -> int:
    ap = argparse.ArgumentParser(
        description="跨轮次对比报告（同一单元在多个 round 间的维度对比）")
    ap.add_argument("--rounds", nargs="+", required=True,
                    help="2 个及以上 round 根目录，第 1 个为基准轮")
    ap.add_argument("--unit", required=True, metavar="MODEL@HARNESS",
                    help="对比单元原始 ID，如 xopglm52@astroncode")
    ap.add_argument("--task-scope", default="intersection",
                    choices=("intersection", "exclude-tags", "task-file", "all"),
                    help="用例范围（默认 intersection：各轮交集）")
    ap.add_argument("--exclude-tags", nargs="+", default=["custom"],
                    help="--task-scope exclude-tags 时要排除的 tag（默认 custom）")
    ap.add_argument("--task-file", help="--task-scope task-file 时的 task_id 清单文件")
    ap.add_argument("--dimensions", default="category,capability,difficulty",
                    help="逗号分隔：category,capability,difficulty,modality")
    ap.add_argument("--entities", default=str(DEFAULT_ENTITIES_PATH))
    ap.add_argument("--pricing-date", type=date.fromisoformat,
                    help="定价快照日期 YYYY-MM-DD（对比报告不出成本列，可省略）")
    ap.add_argument("--tasks-dir", help="任务定义目录（默认自动查找 <repo>/tasks）")
    ap.add_argument("--capability-map", help="检查点能力映射 YAML")
    ap.add_argument("-o", "--output-dir", help="输出目录（默认最后一轮的 report-workspace/output）")
    args = ap.parse_args()

    if len(args.rounds) < 2:
        ap.error("--rounds 至少需要 2 个 round 目录")
    if "@" not in args.unit:
        ap.error("--unit 格式应为 MODEL@HARNESS")
    model, _, harness = args.unit.partition("@")

    round_roots = []
    for raw in args.rounds:
        p = Path(raw).expanduser().resolve()
        if not p.is_dir():
            ap.error(f"round 目录不存在：{p}")
        round_roots.append(p)

    registry = None
    if args.pricing_date:
        registry = report_entities.load_registry(Path(args.entities))

    units: list[UnitResult] = []
    for root in round_roots:
        u = load_round_unit(root, model, harness, registry, args.pricing_date)
        relabel_unit(u, root.name)
        units.append(u)
        print(f"已加载 {u.unit}：{len(u.tasks)} 个任务，总均分 {u.total_pct:.1f}%")

    tasks_dir = find_tasks_dir(args.tasks_dir)
    if tasks_dir is None and args.task_scope == "exclude-tags":
        sys.exit("错误：未找到 tasks/ 目录，exclude-tags 无法判定 tag")
    task_meta = load_all_task_meta(tasks_dir)

    keep, scope_note = resolve_scope(args, units, tasks_dir)
    if not keep:
        sys.exit("错误：用例范围为空，无法对比")
    for u in units:
        before = len(u.tasks)
        restrict_tasks(u, keep)
        print(f"  {u.unit}：{before} → {len(u.tasks)} 例，"
              f"过滤后总均分 {u.total_pct:.1f}%")

    counts = {len(u.tasks) for u in units}
    if len(counts) != 1:
        print(f"[警告] 各轮用例数不一致 {sorted(counts)}，分母不同、分数不可直接比较",
              file=sys.stderr)

    # 出表顺序固定按 --rounds 传入顺序（时间线），不按分数排序
    order = build_task_order(units)
    suites = sorted({s for s, _ in order})
    suite_zh = build_suite_zh_map(task_meta)

    def meta_groups(field: str, known_order: list[str],
                    label_map: dict[str, str] | None = None) -> list[tuple[str, set[str]]]:
        values = {task_meta.get(tid, {}).get(field, "") for _, tid in order}
        values.discard("")
        ordered = [v for v in known_order if v in values] + sorted(values - set(known_order))
        return [((label_map or {}).get(v, v),
                 {tid for _, tid in order if task_meta.get(tid, {}).get(field) == v})
                for v in ordered]

    category_groups = [(suite_zh.get(s, s), {tid for su, tid in order if su == s})
                       for s in suites]
    difficulty_groups = meta_groups("difficulty", DIFFICULTY_ORDER)
    modality_groups = meta_groups("modality", MODALITY_ORDER, MODALITY_ZH)

    wanted = {d.strip() for d in args.dimensions.split(",") if d.strip()}
    wb = Workbook()
    wb.remove(wb.active)

    dim_groups_for_diff: list[tuple[str, list[tuple[str, set[str]]]]] = []
    # 控制变量视图针对"固定 Harness 比模型 / 固定模型比 Harness"，跨轮场景不适用，
    # 故 target_model / target_harness 一律传 None 跳过。
    if "category" in wanted:
        write_dimension_sheet_transposed(wb, "分类对比", units, category_groups, None, None)
        dim_groups_for_diff.append(("分类", category_groups))
    if "capability" in wanted:
        cap_map = load_capability_map(args.capability_map)
        if cap_map:
            write_capability_sheet(wb, units, cap_map, None, None)
        else:
            print("[警告] 未加载到能力映射，跳过 Agent 能力维度", file=sys.stderr)
    if "difficulty" in wanted:
        write_dimension_sheet_transposed(wb, "难度对比", units, difficulty_groups, None, None)
        dim_groups_for_diff.append(("难度", difficulty_groups))
    if "modality" in wanted:
        write_dimension_sheet_transposed(wb, "模态对比", units, modality_groups, None, None)
        dim_groups_for_diff.append(("模态", modality_groups))

    if not wb.sheetnames:
        sys.exit("错误：--dimensions 未产出任何 Sheet")

    write_round_diff_sheet(wb, units, dim_groups_for_diff, scope_note, units[0].round_label)
    write_compare_metadata_sheet(wb, units, round_roots, scope_note, keep, args)

    out_dir = Path(args.output_dir) if args.output_dir else (
        round_roots[-1] / "report-workspace" / "output")
    out_dir.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_unit = f"{model}@{harness}".replace("/", "_")
    out_path = out_dir / f"compare_{safe_unit}_{len(units)}rounds_{ts}.xlsx"
    wb.save(out_path)
    print(f"\n✅ 跨轮对比报告已生成：{out_path}")
    print(f"   口径：{scope_note}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
