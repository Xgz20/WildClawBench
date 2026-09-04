from types import SimpleNamespace
import unittest

from openpyxl import Workbook

from tools.report.scripts import generate_eval_report


class ToolCompareSheetTest(unittest.TestCase):
    def test_tool_search_row_and_hit_rate_are_rendered(self):
        metrics = {
            "total": 3,
            "success": 3,
            "failure": 0,
            "format_error": 0,
            "unclear": 0,
            "search_total": 2,
            "search_hit": 1,
            "search_miss": 1,
            "search_unresolved": 0,
            "by_tool": {
                "bash": {
                    "total": 1,
                    "success": 1,
                    "failure": 0,
                    "format_error": 0,
                    "unclear": 0,
                },
                "tool_search": {
                    "total": 2,
                    "success": 2,
                    "failure": 0,
                    "format_error": 0,
                    "unclear": 0,
                    "search_total": 2,
                    "search_hit": 1,
                    "search_miss": 1,
                    "search_unresolved": 0,
                },
            },
        }
        unit = SimpleNamespace(
            harness="astroncode",
            harness_display="AstronCode",
            model_display="Test Model",
            tool_metrics_total=lambda: metrics,
        )
        workbook = Workbook()

        generate_eval_report.write_tool_compare_sheet(workbook, [unit])

        sheet = workbook["工具调用对比"]
        self.assertEqual(
            [cell.value for cell in sheet[2]],
            [
                "模型",
                "工具",
                "调用数",
                "成功",
                "失败",
                "不确定",
                "格式错误",
                "成功率",
                "格式准确率",
                "检索命中率（tool_search）",
            ],
        )
        rows = {sheet.cell(row, 2).value: row for row in range(3, sheet.max_row + 1)}
        search_row = rows["tool_search"]
        self.assertEqual(sheet.cell(search_row, 3).value, 2)
        self.assertEqual(sheet.cell(search_row, 10).value, 50.0)
        self.assertEqual(sheet.cell(rows["（全部工具）"], 10).value, 50.0)
        self.assertEqual(sheet.cell(rows["bash"], 10).value, "-")


if __name__ == "__main__":
    unittest.main()
