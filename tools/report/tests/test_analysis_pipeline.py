from __future__ import annotations

import importlib.util
import json
import sqlite3
import subprocess
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook
from src.utils.cli_args import build_run_batch_parser
from src.utils.run_selection import write_rerun_metadata


REPORT_DIR = Path(__file__).resolve().parents[1]
MANIFEST_SCRIPT = REPORT_DIR / "skills/low-score-analysis/scripts/generate_failed_tasks_manifest.py"
UTILS_SCRIPT = REPORT_DIR / "skills/low-score-analysis/scripts/utils.py"
EXCEL_SCRIPT = REPORT_DIR / "scripts/generate_eval_report.py"
REPORT_ENTITIES_SCRIPT = REPORT_DIR / "scripts/report_entities.py"
VALIDITY_SCRIPT = REPORT_DIR / "skills/validate-eval-results/scripts/validate_eval_results.py"
AUDIT_SCRIPT = REPORT_DIR / "skills/audit-eval-report/scripts/audit_eval_report.py"
LEADER_EXTRACT_SCRIPT = REPORT_DIR / "skills/eval-report/scripts/extract_leader_report_data.py"
DEEPSEEK_FIXTURE = REPORT_DIR.parent.parent / "tests/fixtures/deepseek_harness"


def load_module(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


manifest = load_module("analysis_manifest", MANIFEST_SCRIPT)
analysis_utils = load_module("analysis_utils", UTILS_SCRIPT)
excel_report = load_module("excel_report", EXCEL_SCRIPT)
validity_check = load_module("validity_check", VALIDITY_SCRIPT)
report_audit = load_module("report_audit", AUDIT_SCRIPT)


class AnalysisPipelineTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.round_dir = Path(self.temp_dir.name) / "round3"
        self.unit_dir = self.round_dir / "model-x" / "harness-y"
        self.suite_dir = self.unit_dir / "01_suite"
        self.paths = {}
        for task_id, score in (
            ("task_50", 0.5),
            ("task_70", 0.7),
            ("task_almost_full", 0.9996),
            ("task_float_edge", 0.9999999995),
            ("task_full", 1.0),
            ("task_unscored", None),
        ):
            run_dir = self.suite_dir / task_id / "run_001"
            run_dir.mkdir(parents=True)
            if score is not None:
                (run_dir / "score.json").write_text(
                    json.dumps({"overall_score": score, "check": score}), encoding="utf-8"
                )
            (run_dir / "execution_status.json").write_text(
                json.dumps({"status": "completed"}), encoding="utf-8"
            )
            (run_dir / "usage.json").write_text(
                json.dumps({"request_count": 1, "total_tokens": 100}), encoding="utf-8"
            )
            (run_dir / "chat_openclaw.jsonl").write_text(
                "\n".join(json.dumps({"type": "event", "payload": {"index": index}})
                          for index in range(5)),
                encoding="utf-8",
            )
            self.paths[task_id] = run_dir

        self.tasks_dir = Path(self.temp_dir.name) / "tasks"
        task_suite = self.tasks_dir / "01_suite"
        task_suite.mkdir(parents=True)
        for task_id in self.paths:
            (task_suite / f"{task_id}.md").write_text(
                "---\n"
                f"id: {task_id}\n"
                "name: Test task\n"
                "category: 01_suite\n"
                "difficulty: L2\n"
                "modality: pure-text\n"
                "timeout_seconds: 300\n"
                "grading_type: automated\n"
                "---\n\n## Prompt\nTest\n",
                encoding="utf-8",
            )

        self.records = manifest.scan_unit("model-x", "harness-y", self.unit_dir, None)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def selection(self, **overrides):
        values = {
            "select_all": False,
            "imperfect": False,
            "threshold": None,
            "score_min": None,
            "score_max": None,
            "task_ids": [],
            "task_paths": [],
        }
        values.update(overrides)
        return manifest.build_selection(**values)

    def test_entity_registry_loads_names_and_dated_pricing(self) -> None:
        report_entities = load_module("report_entities_registry", REPORT_ENTITIES_SCRIPT)
        registry = report_entities.load_registry(REPORT_DIR / "data/entities.yaml")

        self.assertEqual(registry.model_display("xopglm52"), "GLM-5.2")
        self.assertEqual(registry.harness_display("astroncode"), "AstronCode")
        self.assertEqual(
            registry.harness_display("deepseek-harness"),
            "DeepSeek Harness",
        )
        self.assertEqual(
            registry.pricing_profile("gpt-5.5", date(2026, 7, 30)).profile_id,
            "2026-07-30-openai",
        )
        self.assertAlmostEqual(float(registry.cny_per_usd(date(2026, 7, 30))), 6.77)

    def test_entity_registry_covers_all_cli_harnesses(self) -> None:
        report_entities = load_module("report_entities_harnesses", REPORT_ENTITIES_SCRIPT)
        registry = report_entities.load_registry(REPORT_DIR / "data/entities.yaml")
        parser = build_run_batch_parser(default_model="test-model", default_parallel=1)
        backend_action = next(
            action for action in parser._actions if action.dest == "agent_backend"
        )
        expected_displays = {
            "openclaw": "OpenClaw",
            "astronclaw": "AstronClaw",
            "claudecode": "Claude Code",
            "codex": "Codex",
            "hermesagent": "Hermes Agent",
            "astroncode": "AstronCode",
            "opencode": "OpenCode",
            "deepseek-harness": "DeepSeek Harness",
        }

        self.assertTrue(set(backend_action.choices).issubset(registry.harnesses))
        self.assertEqual(
            {
                harness_id: registry.harness_display(harness_id)
                for harness_id in backend_action.choices
            },
            expected_displays,
        )
        self.assertTrue(
            all(
                registry.harnesses[harness_id].get("family")
                for harness_id in backend_action.choices
            )
        )

    def test_entity_registry_unknown_id_falls_back_with_warning(self) -> None:
        report_entities = load_module("report_entities_fallback", REPORT_ENTITIES_SCRIPT)
        registry = report_entities.load_registry(REPORT_DIR / "data/entities.yaml")

        with self.assertLogs(level="WARNING"):
            self.assertEqual(registry.model_display("model-new"), "model-new")

    def test_entity_registry_rejects_unknown_schema_version(self) -> None:
        report_entities = load_module("report_entities_schema", REPORT_ENTITIES_SCRIPT)
        path = Path(self.temp_dir.name) / "entities.yaml"
        path.write_text(
            "schema_version: 99\nmodels: {}\nharnesses: {}\n",
            encoding="utf-8",
        )

        with self.assertRaisesRegex(ValueError, "schema_version"):
            report_entities.load_registry(path)

    def test_cost_normalizes_astroncode_and_opencode_usage(self) -> None:
        report_entities = load_module("report_entities_usage", REPORT_ENTITIES_SCRIPT)
        astron = report_entities.normalize_billable_usage({
            "input_tokens": 1000,
            "cache_read_tokens": 600,
            "cache_write_tokens": 0,
            "output_tokens": 100,
            "total_tokens": 1100,
        })
        opencode = report_entities.normalize_billable_usage({
            "input_tokens": 400,
            "cache_read_tokens": 600,
            "cache_write_tokens": 0,
            "output_tokens": 100,
            "total_tokens": 1100,
        })

        self.assertEqual(astron, report_entities.BillableUsage(400, 600, 0, 100))
        self.assertEqual(opencode, report_entities.BillableUsage(400, 600, 0, 100))

    def test_cost_converts_glm_and_spark_from_cny(self) -> None:
        report_entities = load_module("report_entities_cny", REPORT_ENTITIES_SCRIPT)
        registry = report_entities.load_registry(REPORT_DIR / "data/entities.yaml")
        usage = report_entities.BillableUsage(400, 600, 0, 100)

        glm = report_entities.estimate_cost_usd(
            registry, "xopglm52", date(2026, 7, 30), usage,
            request_input_tokens=None,
        )
        spark = report_entities.estimate_cost_usd(
            registry, "xsparkx2agent", date(2026, 7, 30), usage,
            request_input_tokens=None,
        )

        self.assertAlmostEqual(float(glm.usd), 0.0072 / 6.77, places=10)
        self.assertAlmostEqual(float(spark.usd), 0.00358 / 6.77, places=10)

    def test_raw_usage_cost_status_prevents_unavailable_cost_from_becoming_zero(self) -> None:
        result = excel_report._estimate_run_cost(
            "xopdeepseekv4flash0731",
            "deepseek-harness",
            None,
            {
                "cost_usd": 0.0,
                "cost_status": "unavailable",
                "cost_reason": "DSH sessions do not include provider cost",
            },
            registry=None,
            pricing_date=None,
        )

        self.assertIsNone(result.usd)
        self.assertEqual(result.status, "unavailable")
        self.assertEqual(result.reason, "DSH sessions do not include provider cost")

    def test_not_applicable_cost_does_not_require_registry_price(self) -> None:
        report_entities = load_module("report_entities_not_applicable", REPORT_ENTITIES_SCRIPT)
        registry = report_entities.load_registry(REPORT_DIR / "data/entities.yaml")
        result = excel_report._estimate_run_cost(
            "xopdeepseekv4flash0731",
            "deepseek-harness",
            None,
            {"cost_usd": 0.0, "cost_status": "not_applicable"},
            registry=registry,
            pricing_date=date(2026, 8, 15),
        )

        self.assertEqual(float(result.usd), 0.0)
        self.assertEqual(result.status, "not_applicable")

    def test_deepseek_harness_extractor_reads_each_request_usage(self) -> None:
        report_entities = load_module("report_entities_deepseek_requests", REPORT_ENTITIES_SCRIPT)

        requests = report_entities.extract_deepseek_harness_requests(DEEPSEEK_FIXTURE)

        self.assertCountEqual(
            requests,
            [
                report_entities.RequestUsage(10, 2, 4, 1),
                report_entities.RequestUsage(3, 0, 2, 0),
                report_entities.RequestUsage(5, 0, 1, 3),
            ],
        )

    def test_deepseek_harness_report_cost_uses_model_tiers(self) -> None:
        report_entities = load_module("report_entities_deepseek_cost", REPORT_ENTITIES_SCRIPT)
        registry = report_entities.load_registry(REPORT_DIR / "data/entities.yaml")
        result = excel_report._estimate_run_cost(
            "gpt-5.6-sol",
            "deepseek-harness",
            DEEPSEEK_FIXTURE,
            {"cost_status": "unavailable", "cost_usd": 0.0},
            registry=registry,
            pricing_date=date(2026, 8, 15),
        )

        expected = (177.25 + 75 + 73.75) / 1_000_000
        self.assertAlmostEqual(float(result.usd), expected)
        self.assertEqual(result.status, "estimated")

    def test_deepseek_harness_report_cost_is_unavailable_without_model_price(self) -> None:
        report_entities = load_module("report_entities_deepseek_unpriced", REPORT_ENTITIES_SCRIPT)
        registry = report_entities.load_registry(REPORT_DIR / "data/entities.yaml")
        result = excel_report._estimate_run_cost(
            "xopdeepseekv4flash0731",
            "deepseek-harness",
            DEEPSEEK_FIXTURE,
            {"cost_status": "unavailable", "cost_usd": 0.0},
            registry=registry,
            pricing_date=date(2026, 8, 15),
        )

        self.assertIsNone(result.usd)
        self.assertEqual(result.status, "unavailable")
        self.assertIn("定价档案", result.reason)

    def test_gpt_cost_uses_each_request_context_tier(self) -> None:
        report_entities = load_module("report_entities_gpt", REPORT_ENTITIES_SCRIPT)
        registry = report_entities.load_registry(REPORT_DIR / "data/entities.yaml")
        requests = [
            report_entities.RequestUsage(100000, 100000, 1000),
            report_entities.RequestUsage(200000, 100000, 1000),
        ]

        result = report_entities.estimate_request_costs_usd(
            registry, "gpt-5.5", date(2026, 7, 30), requests
        )
        expected = (
            100000 * 5 + 100000 * 0.5 + 1000 * 30
            + 200000 * 10 + 100000 * 1 + 1000 * 45
        ) / 1_000_000
        self.assertAlmostEqual(float(result.usd), expected)

    def test_cost_is_unavailable_without_cache_write_price(self) -> None:
        report_entities = load_module("report_entities_cache_write", REPORT_ENTITIES_SCRIPT)
        registry = report_entities.load_registry(REPORT_DIR / "data/entities.yaml")

        result = report_entities.estimate_cost_usd(
            registry,
            "xopglm52",
            date(2026, 7, 30),
            report_entities.BillableUsage(400, 600, 1, 100),
            request_input_tokens=None,
        )

        self.assertIsNone(result.usd)
        self.assertEqual(result.status, "unavailable")
        self.assertIn("缓存写入", result.reason)

    def test_request_readers_normalize_astroncode_and_opencode(self) -> None:
        report_entities = load_module("report_entities_readers", REPORT_ENTITIES_SCRIPT)
        astron_dir = Path(self.temp_dir.name) / "astron-run"
        astron_dir.mkdir()
        astron_events = [
            {
                "type": "event_msg",
                "payload": {
                    "type": "token_count",
                    "info": {
                        "total_token_usage": {"input_tokens": 100, "output_tokens": 10},
                        "last_token_usage": {
                            "input_tokens": 100,
                            "cached_input_tokens": 20,
                            "output_tokens": 10,
                        },
                    },
                },
            },
            {
                "type": "event_msg",
                "payload": {
                    "type": "token_count",
                    "info": {
                        "total_token_usage": {"input_tokens": 100, "output_tokens": 10},
                        "last_token_usage": {
                            "input_tokens": 100,
                            "cached_input_tokens": 20,
                            "output_tokens": 10,
                        },
                    },
                },
            },
            {
                "type": "event_msg",
                "payload": {
                    "type": "token_count",
                    "info": {
                        "total_token_usage": {"input_tokens": 220, "output_tokens": 22},
                        "last_token_usage": {
                            "input_tokens": 120,
                            "cached_input_tokens": 30,
                            "output_tokens": 12,
                        },
                    },
                },
            },
        ]
        (astron_dir / "chat.jsonl").write_text(
            "\n".join(json.dumps(event) for event in astron_events),
            encoding="utf-8",
        )

        opencode_dir = Path(self.temp_dir.name) / "opencode-run"
        database_dir = opencode_dir / "opencode_data"
        database_dir.mkdir(parents=True)
        with sqlite3.connect(database_dir / "opencode.db") as connection:
            connection.execute("CREATE TABLE part (id TEXT PRIMARY KEY, data TEXT NOT NULL)")
            for index, tokens in enumerate((
                {"input": 80, "output": 7, "reasoning": 3,
                 "cache": {"read": 20, "write": 0}},
                {"input": 90, "output": 10, "reasoning": 2,
                 "cache": {"read": 30, "write": 0}},
            )):
                connection.execute(
                    "INSERT INTO part (id, data) VALUES (?, ?)",
                    (str(index), json.dumps({"type": "step-finish", "tokens": tokens})),
                )

        expected = [
            report_entities.RequestUsage(80, 20, 10),
            report_entities.RequestUsage(90, 30, 12),
        ]
        self.assertEqual(report_entities.extract_astroncode_requests(astron_dir), expected)
        self.assertEqual(report_entities.extract_opencode_requests(opencode_dir), expected)

    def test_excel_uses_display_names_and_preserves_raw_identity(self) -> None:
        workbook = load_workbook(self.generate_comparison_excel(), data_only=True)
        overview = workbook["总览"]
        headers = [cell.value for cell in overview[1]]
        models = {
            overview.cell(row, headers.index("模型") + 1).value
            for row in range(2, overview.max_row + 1)
        }
        harnesses = {
            str(overview.cell(row, headers.index("Harness") + 1).value).split(" (", 1)[0]
            for row in range(2, overview.max_row + 1)
        }

        self.assertEqual(models, {"GPT-5.5", "Spark-X2-300B"})
        self.assertEqual(harnesses, {"AstronCode", "OpenCode"})
        self.assertEqual(workbook["_报告元数据"].sheet_state, "hidden")
        self.assertIn(
            "xsparkx2agent@astroncode",
            {
                workbook["_报告元数据"].cell(row, 2).value
                for row in range(2, workbook["_报告元数据"].max_row + 1)
            },
        )

    def test_excel_recomputes_nonzero_costs(self) -> None:
        workbook = load_workbook(self.generate_comparison_excel(), data_only=True)
        overview = workbook["总览"]
        headers = [cell.value for cell in overview[1]]
        cost_column = headers.index("总成本(USD)") + 1

        self.assertTrue(
            all(
                overview.cell(row, cost_column).value > 0
                for row in range(2, overview.max_row + 1)
            )
        )

    def test_summary_dimensions_use_only_active_tasks(self) -> None:
        class FakeUnit:
            def __init__(self, label, scores):
                self.unit_display = label
                self.tasks = [SimpleNamespace(task_id=task_id) for task_id in scores]
                self.scores = scores

            def avg_pct(self, task_ids):
                values = [self.scores[task_id] for task_id in task_ids
                          if task_id in self.scores]
                return sum(values) / len(values) * 100 if values else None

        units = [FakeUnit("model@harness", {"active_l1": 0.8, "active_l2": 0.6})]
        task_meta = {
            "active_l1": {"difficulty": "L1"},
            "active_l2": {"difficulty": "L2"},
            "unrelated_l3": {"difficulty": "L3"},
        }

        rows = excel_report._build_dim_comparison(
            units, task_meta, "difficulty", ["L1", "L2", "L3"]
        )

        self.assertEqual([(row["name"], row["task_count"]) for row in rows],
                         [("L1", 1), ("L2", 1)])
        self.assertEqual(rows[0]["scores"]["model@harness"], 80.0)

    def test_dimension_sheets_append_two_controlled_views(self) -> None:
        workbook = load_workbook(self.generate_comparison_excel())
        expected_sheets = (
            "分类对比",
            "Agent能力对比",
            "Agent能力对比·去污染",
            "难度对比",
            "模态对比",
        )
        for title in expected_sheets:
            sheet = workbook[title]
            first_column = [
                sheet.cell(row, 1).value for row in range(1, sheet.max_row + 1)
            ]
            self.assertEqual(sheet["A1"].value, "模型@Harness")
            self.assertIn("固定 AstronCode：模型对比", first_column)
            self.assertIn("固定 Spark-X2-300B：Harness 对比", first_column)

            model_title_row = first_column.index("固定 AstronCode：模型对比") + 1
            harness_title_row = first_column.index(
                "固定 Spark-X2-300B：Harness 对比"
            ) + 1
            self.assertEqual(sheet.cell(model_title_row + 1, 1).value, "模型")
            self.assertEqual(sheet.cell(harness_title_row + 1, 1).value, "Harness")
            self.assertEqual(
                {
                    sheet.cell(row, 1).value
                    for row in range(model_title_row + 2, harness_title_row - 2)
                    if sheet.cell(row, 1).value
                },
                {"GPT-5.5", "Spark-X2-300B"},
            )
            self.assertEqual(
                {
                    sheet.cell(row, 1).value
                    for row in range(harness_title_row + 2, sheet.max_row + 1)
                    if sheet.cell(row, 1).value
                },
                {"AstronCode", "OpenCode"},
            )

    def test_controlled_views_preserve_overview_and_apply_styles(self) -> None:
        workbook = load_workbook(self.generate_comparison_excel())
        self.assertEqual(
            [cell.value for cell in workbook["总览"][1]],
            [
                "模型", "Harness", "总平均分", "用例数", "正常完成数",
                "执行错误数", "超时数", "评测异常数", "完成率", "总tokens",
                "总请求数", "总耗时(s)", "总成本(USD)",
            ],
        )
        self.assertEqual(
            workbook.sheetnames[:6],
            [
                "总览", "分类对比", "Agent能力对比", "Agent能力对比·去污染",
                "难度对比", "模态对比",
            ],
        )
        overview = workbook["总览"]
        target_overview_cell = next(
            overview.cell(row, 1)
            for row in range(2, overview.max_row + 1)
            if overview.cell(row, 1).value == "Spark-X2-300B"
            and str(overview.cell(row, 2).value).startswith("AstronCode")
        )
        self.assertTrue(target_overview_cell.font.bold)
        self.assertEqual(target_overview_cell.fill.fill_type, "solid")
        for title in workbook.sheetnames[1:6]:
            sheet = workbook[title]
            first_column = [
                sheet.cell(row, 1).value for row in range(1, sheet.max_row + 1)
            ]
            model_title_row = first_column.index("固定 AstronCode：模型对比") + 1
            harness_title_row = first_column.index(
                "固定 Spark-X2-300B：Harness 对比"
            ) + 1
            target_model_cell = next(
                sheet.cell(row, 1)
                for row in range(model_title_row + 2, harness_title_row)
                if sheet.cell(row, 1).value == "Spark-X2-300B"
            )
            target_harness_cell = next(
                sheet.cell(row, 1)
                for row in range(harness_title_row + 2, sheet.max_row + 1)
                if sheet.cell(row, 1).value == "AstronCode"
            )
            self.assertTrue(target_model_cell.font.bold)
            self.assertEqual(target_model_cell.fill.fill_type, "solid")
            self.assertTrue(target_harness_cell.font.bold)
            self.assertEqual(target_harness_cell.fill.fill_type, "solid")
            self.assertGreaterEqual(len(sheet.conditional_formatting), 3)

    def test_controlled_views_wrap_headers_and_multiline_values(self) -> None:
        workbook = load_workbook(self.generate_comparison_excel())
        for title in (
            "分类对比",
            "Agent能力对比",
            "Agent能力对比·去污染",
            "难度对比",
            "模态对比",
        ):
            sheet = workbook[title]
            first_column = [
                sheet.cell(row, 1).value for row in range(1, sheet.max_row + 1)
            ]
            for view_title in (
                "固定 AstronCode：模型对比",
                "固定 Spark-X2-300B：Harness 对比",
            ):
                header_row = first_column.index(view_title) + 2
                self.assertGreaterEqual(
                    sheet.row_dimensions[header_row].height or 0, 36
                )
                self.assertTrue(
                    all(
                        sheet.cell(header_row, column).alignment.wrap_text
                        for column in range(1, sheet.max_column + 1)
                    )
                )

        controlled_workbook = Workbook()
        agent_sheet = controlled_workbook.active
        units = [
            SimpleNamespace(
                unit="model-a@harness-a", model="model-a", harness="harness-a",
                model_display="Model A", harness_display="Harness A",
            ),
            SimpleNamespace(
                unit="model-b@harness-a", model="model-b", harness="harness-a",
                model_display="Model B", harness_display="Harness A",
            ),
            SimpleNamespace(
                unit="model-a@harness-b", model="model-a", harness="harness-b",
                model_display="Model A", harness_display="Harness B",
            ),
        ]
        values = {
            "model-a@harness-a": [50.0, "验证交付 50%\n工具调用 40%\n代码生成 30%"],
            "model-b@harness-a": [60.0, "验证交付 60%\n工具调用 50%\n代码生成 40%"],
            "model-a@harness-b": [55.0, "验证交付 55%\n工具调用 45%\n代码生成 35%"],
        }
        excel_report.append_controlled_views(
            agent_sheet,
            units,
            ["总平均分", "模型强项"],
            values,
            "model-a",
            "harness-a",
        )
        multiline_cells = [
            cell
            for row in agent_sheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and "\n" in cell.value
        ]
        self.assertTrue(multiline_cells)
        self.assertTrue(all(cell.alignment.wrap_text for cell in multiline_cells))
        self.assertTrue(
            all(
                (agent_sheet.row_dimensions[cell.row].height or 0) >= 45
                for cell in multiline_cells
            )
        )

    def add_valid_run(self, task_id: str, name: str, score: float) -> Path:
        run_dir = self.suite_dir / task_id / name
        run_dir.mkdir()
        (run_dir / "score.json").write_text(
            json.dumps({"overall_score": score, "check": score}), encoding="utf-8"
        )
        (run_dir / "execution_status.json").write_text(
            json.dumps({"status": "completed"}), encoding="utf-8"
        )
        (run_dir / "usage.json").write_text(
            json.dumps({"request_count": 1, "total_tokens": 100}), encoding="utf-8"
        )
        (run_dir / "chat_openclaw.jsonl").write_text(
            "\n".join(json.dumps({"type": "event", "payload": {"index": index}})
                      for index in range(5)),
            encoding="utf-8",
        )
        return run_dir

    def create_comparison_fixture(self) -> Path:
        result_root = Path(self.temp_dir.name) / "comparison-round"
        self.comparison_tasks_dir = Path(self.temp_dir.name) / "comparison-tasks"
        comparison_suite = self.comparison_tasks_dir / "01_suite"
        comparison_suite.mkdir(parents=True)
        (comparison_suite / "task_cost.md").write_text(
            "---\n"
            "id: task_cost\n"
            "name: Cost task\n"
            "category: 01_suite\n"
            "difficulty: L2\n"
            "modality: pure-text\n"
            "timeout_seconds: 300\n"
            "grading_type: automated\n"
            "---\n\n## Prompt\nTest\n",
            encoding="utf-8",
        )
        for model in ("gpt-5.5", "xsparkx2agent"):
            for harness in ("astroncode", "opencode"):
                run_dir = result_root / model / harness / "01_suite/task_cost/run_001"
                run_dir.mkdir(parents=True)
                (run_dir / "score.json").write_text(
                    json.dumps({"overall_score": 0.8}), encoding="utf-8"
                )
                (run_dir / "execution_status.json").write_text(
                    json.dumps({
                        "status": "completed",
                        "elapsed_time": 10,
                        "harness_version": "1.0.0",
                    }),
                    encoding="utf-8",
                )
                usage = {
                    "input_tokens": 1000 if harness == "astroncode" else 400,
                    "cache_read_tokens": 600,
                    "cache_write_tokens": 0,
                    "output_tokens": 100,
                    "total_tokens": 1100,
                    "request_count": 1,
                    "elapsed_time": 10,
                    "cost_usd": 0,
                }
                (run_dir / "usage.json").write_text(json.dumps(usage), encoding="utf-8")
                if model == "gpt-5.5" and harness == "astroncode":
                    event = {
                        "type": "event_msg",
                        "payload": {
                            "type": "token_count",
                            "info": {
                                "total_token_usage": {
                                    "input_tokens": 1000,
                                    "output_tokens": 100,
                                },
                                "last_token_usage": {
                                    "input_tokens": 1000,
                                    "cached_input_tokens": 600,
                                    "output_tokens": 100,
                                },
                            },
                        },
                    }
                    (run_dir / "chat.jsonl").write_text(
                        json.dumps(event) + "\n", encoding="utf-8"
                    )
                elif model == "gpt-5.5":
                    database_dir = run_dir / "opencode_data"
                    database_dir.mkdir()
                    with sqlite3.connect(database_dir / "opencode.db") as connection:
                        connection.execute(
                            "CREATE TABLE part (id TEXT PRIMARY KEY, data TEXT NOT NULL)"
                        )
                        connection.execute(
                            "INSERT INTO part (id, data) VALUES (?, ?)",
                            (
                                "1",
                                json.dumps({
                                    "type": "step-finish",
                                    "tokens": {
                                        "input": 400,
                                        "output": 100,
                                        "reasoning": 0,
                                        "cache": {"read": 600, "write": 0},
                                    },
                                }),
                            ),
                        )
                else:
                    (run_dir / "chat_openclaw.jsonl").write_text(
                        "\n".join(
                            json.dumps({"type": "event", "payload": {"index": index}})
                            for index in range(5)
                        ),
                        encoding="utf-8",
                    )
        return result_root

    def generate_comparison_excel(self, result_root: Path | None = None) -> Path:
        result_root = result_root or self.create_comparison_fixture()
        output_dir = result_root / "comparison-output"
        subprocess.run(
            [
                sys.executable,
                str(EXCEL_SCRIPT),
                "--result-root",
                str(result_root),
                "--models",
                "gpt-5.5",
                "xsparkx2agent",
                "--harnesses",
                "astroncode",
                "opencode",
                "--target-model",
                "xsparkx2agent",
                "--target-harness",
                "astroncode",
                "--entities",
                str(REPORT_DIR / "data/entities.yaml"),
                "--pricing-date",
                "2026-07-30",
                "--tasks-dir",
                str(self.comparison_tasks_dir),
                "--output-dir",
                str(output_dir),
            ],
            check=True,
            capture_output=True,
            text=True,
        )
        return next(output_dir.glob("report_4units_*.xlsx"))

    def test_validity_and_audit_support_selected_scope(self) -> None:
        result_root = self.create_comparison_fixture()
        validity = validity_check.scan_round(
            result_root,
            self.comparison_tasks_dir,
            models={"gpt-5.5"},
            harnesses={"astroncode"},
        )
        self.assertEqual(set(validity["units"]), {"gpt-5.5@astroncode"})
        self.assertEqual(validity["scope"]["models"], ["gpt-5.5"])
        self.assertEqual(validity["scope"]["harnesses"], ["astroncode"])

        excel_path = self.generate_comparison_excel(result_root)
        audit = report_audit.audit_report(
            result_root,
            excel_path,
            self.comparison_tasks_dir,
            models={"gpt-5.5", "xsparkx2agent"},
            harnesses={"astroncode", "opencode"},
            entities_path=REPORT_DIR / "data/entities.yaml",
            pricing_date=date(2026, 7, 30),
        )
        self.assertFalse(
            [item for item in audit["findings"] if item["severity"] == "error"]
        )

    def test_expected_tasks_prefers_structured_evaluation_scope(self) -> None:
        expected = {
            ("01_suite", "task_1"),
            ("07_Website_Generation", "task_web_1"),
            ("07_Website_Generation", "task_web_2"),
        }
        scope = {
            "schema_version": 1,
            "planned_tasks": [
                {"category": "07_Website_Generation", "task_id": "task_web_1"},
                {"category": "07_Website_Generation", "task_id": "task_web_2"},
            ],
        }

        selected, scoped = validity_check.expected_tasks_for_unit(
            expected, {}, "Category: 01_suite, 1 tasks", scope
        )

        self.assertTrue(scoped)
        self.assertEqual(selected, {
            ("07_Website_Generation", "task_web_1"),
            ("07_Website_Generation", "task_web_2"),
        })

    def test_expected_tasks_historical_log_limits_category_before_tag_filter(self) -> None:
        expected = {
            ("01_suite", "task_1"),
            ("07_Website_Generation", "task_web_1"),
            ("07_Website_Generation", "task_web_2"),
        }
        metadata = {
            ("01_suite", "task_1"): {"modality": "pure-text", "tags": {"common"}},
            ("07_Website_Generation", "task_web_1"): {
                "modality": "pure-text", "tags": {"web-site-gen"},
            },
            ("07_Website_Generation", "task_web_2"): {
                "modality": "pure-text", "tags": {"other"},
            },
        }
        run_log = (
            "Category: 07_Website_Generation, 2 tasks (official + extension), parallelism: 1\n"
            "Tag filter (any of ['web-site-gen']): 1/2 tasks kept in 07_Website_Generation\n"
        )

        selected, scoped = validity_check.expected_tasks_for_unit(
            expected, metadata, run_log, None
        )

        self.assertTrue(scoped)
        self.assertEqual(selected, {("07_Website_Generation", "task_web_1")})

    def test_structured_scope_still_detects_missing_planned_task(self) -> None:
        missing_task = self.tasks_dir / "01_suite" / "task_selected_but_missing.md"
        missing_task.write_text(
            "---\n"
            "id: task_selected_but_missing\n"
            "name: Missing selected task\n"
            "category: 01_suite\n"
            "difficulty: L2\n"
            "modality: pure-text\n"
            "timeout_seconds: 300\n"
            "grading_type: automated\n"
            "---\n\n## Prompt\nTest\n",
            encoding="utf-8",
        )
        (self.unit_dir / "evaluation_scope.json").write_text(json.dumps({
            "schema_version": 1,
            "planned_task_count": 2,
            "planned_tasks": [
                {"category": "01_suite", "task_id": "task_50"},
                {"category": "01_suite", "task_id": "task_selected_but_missing"},
            ],
        }), encoding="utf-8")

        report = validity_check.scan_round(self.round_dir, self.tasks_dir)
        missing = [
            item for item in report["findings"]
            if item["id"] == "TASK_MISSING"
        ]

        self.assertEqual(len(missing), 1)
        self.assertEqual(missing[0]["task_id"], "task_selected_but_missing")

    def test_audit_detects_recomputed_cost_regression(self) -> None:
        excel_path = self.generate_comparison_excel()
        workbook = load_workbook(excel_path)
        overview = workbook["总览"]
        headers = [cell.value for cell in overview[1]]
        overview.cell(2, headers.index("总成本(USD)") + 1).value = 99
        workbook.save(excel_path)

        report = report_audit.audit_report(
            excel_path.parent.parent,
            excel_path,
            self.comparison_tasks_dir,
            models={"gpt-5.5", "xsparkx2agent"},
            harnesses={"astroncode", "opencode"},
            entities_path=REPORT_DIR / "data/entities.yaml",
            pricing_date=date(2026, 7, 30),
        )
        self.assertTrue(any(
            item["id"] == "OVERVIEW_COST_MISMATCH"
            for item in report["findings"]
        ))

    def test_leader_extractor_reads_controlled_views(self) -> None:
        leader_extract = load_module("leader_extract", LEADER_EXTRACT_SCRIPT)
        payload = leader_extract.extract_workbook(self.generate_comparison_excel())

        self.assertEqual(
            payload["target"]["unit_display"], "Spark-X2-300B@AstronCode"
        )
        self.assertEqual(payload["overview"][0]["模型"], "GPT-5.5")
        self.assertEqual(
            {
                row["模型"]
                for row in payload["dimensions"]["分类对比"]["model_view"]
            },
            {"GPT-5.5", "Spark-X2-300B"},
        )
        self.assertEqual(
            {
                row["Harness"]
                for row in payload["dimensions"]["Agent能力对比"]["harness_view"]
            },
            {"AstronCode", "OpenCode"},
        )

    def selected_ids(self, selection, task_ids=None, task_paths=None):
        records = [dict(record) for record in self.records]
        selected = manifest.select_tasks(records, selection, task_ids or [], task_paths or [])
        return {item["task_id"] for item in selected}, selected

    def test_selection_modes_use_raw_score(self) -> None:
        ids, _ = self.selected_ids(self.selection())
        self.assertEqual(ids, {"task_50", "task_unscored"})

        ids, selected = self.selected_ids(self.selection(imperfect=True))
        self.assertEqual(ids, {"task_50", "task_70", "task_almost_full", "task_float_edge",
                               "task_unscored"})
        almost = next(item for item in selected if item["task_id"] == "task_almost_full")
        self.assertEqual(almost["score_pct"], 100.0)
        self.assertEqual(almost["analysis_type"], "failure")

        ids, selected = self.selected_ids(self.selection(select_all=True))
        self.assertEqual(ids, {record["task_id"] for record in self.records})
        full = next(item for item in selected if item["task_id"] == "task_full")
        self.assertEqual(full["analysis_type"], "success_control")

    def test_score_range_and_result_path_selection(self) -> None:
        selection = self.selection(score_min=60, score_max=80)
        self.assertEqual(selection["scope"], "gte60_lt80")
        ids, _ = self.selected_ids(selection)
        self.assertEqual(ids, {"task_70"})

        score_path = self.paths["task_70"] / "score.json"
        selection = self.selection(task_paths=[str(score_path)])
        ids, _ = self.selected_ids(selection, task_paths=[str(score_path)])
        self.assertEqual(ids, {"task_70"})
        self.assertEqual(manifest.find_unit_for_result_path(str(score_path)), self.unit_dir.resolve())

    def test_result_file_path_locks_the_selected_run(self) -> None:
        old_run = self.suite_dir / "task_70/run_000"
        old_run.mkdir()
        (old_run / "score.json").write_text(
            json.dumps({"overall_score": 0.65, "check": 0.65}), encoding="utf-8"
        )
        (old_run / "execution_status.json").write_text("{}", encoding="utf-8")
        records = manifest.apply_task_path_run_overrides(
            self.records, [str(old_run / "score.json")], None
        )
        record = next(item for item in records if item["task_id"] == "task_70")
        self.assertEqual(record["overall_score"], 0.65)
        self.assertEqual(Path(record["run_dir"]), old_run.resolve())

    def test_cli_places_manifest_in_round_workspace(self) -> None:
        result = subprocess.run(
            [sys.executable, str(MANIFEST_SCRIPT), "--result-root", str(self.unit_dir),
             "--threshold", "60", "--tasks-dir", str(Path(self.temp_dir.name) / "missing-tasks")],
            check=True,
            capture_output=True,
            text=True,
        )
        expected = self.round_dir / "report-workspace/_failed_tasks_model-x@harness-y__lt60.json"
        self.assertTrue(expected.is_file())
        self.assertFalse((self.unit_dir / "report-workspace").exists())
        workspace = self.round_dir.resolve() / "report-workspace"
        self.assertIn(f"WORKSPACE_DIR={workspace}", result.stdout)
        self.assertIn(f"ANALYSIS_PATH={workspace / 'analysis_model-x@harness-y__lt60.json'}",
                      result.stdout)

    def test_validity_cli_places_output_in_round_workspace(self) -> None:
        subprocess.run(
            [sys.executable, str(VALIDITY_SCRIPT), "--result-root", str(self.unit_dir),
             "--tasks-dir", str(self.tasks_dir)],
            check=True,
            capture_output=True,
            text=True,
        )
        expected = self.round_dir / "report-workspace/validity/eval_result_validity.json"
        self.assertTrue(expected.is_file())
        self.assertFalse((self.unit_dir / "report-workspace/validity").exists())
        self.assertEqual(
            validity_check.round_root_from_units(self.unit_dir, validity_check.discover_units(self.unit_dir)),
            self.round_dir.resolve(),
        )

    def test_validity_round_root_supports_three_level_layout(self) -> None:
        unit_dir = self.round_dir / "opencode" / "model-z" / "opencode"
        unit_dir.mkdir(parents=True)
        specs = [("model-z", "opencode", unit_dir)]
        self.assertEqual(
            validity_check.round_root_from_units(unit_dir, specs),
            self.round_dir.resolve(),
        )

    def test_validity_ignores_api_keywords_in_agent_log(self) -> None:
        run_dir = self.suite_dir / "task_50/run_001"
        (run_dir / "agent.log").write_text(
            'level=INFO message="task says rate limit and 500"\n'
            'level=ERROR small=true agent=title error="service unavailable"\n'
            'level=ERROR small=false agent=build error="HTTP 429 too many requests"\n',
            encoding="utf-8",
        )
        report = validity_check.scan_round(self.unit_dir, self.tasks_dir)
        ids = {item["id"] for item in report["findings"]}
        self.assertNotIn("MODEL_API_RATE_LIMIT", ids)
        self.assertNotIn("MODEL_API_SERVER_ERROR", ids)
        self.assertNotIn("AGENT_LOG_ERROR", ids)

    def test_validity_findings_redact_credentials(self) -> None:
        item = validity_check.finding(
            "TEST", "error",
            'request failed: AK="ak-ba7df6029dd8d8baae7b62983221a940" password=hunter2 '
            'OPENROUTER_API_KEY=fake-id:fake-secret',
            evidence={"authorization": "Bearer secret-token-value"},
        )
        serialized = json.dumps(item)
        self.assertNotIn("ba7df6029dd8d8baae7b62983221a940", serialized)
        self.assertNotIn("hunter2", serialized)
        self.assertNotIn("secret-token-value", serialized)
        self.assertNotIn("fake-id:fake-secret", serialized)
        self.assertIn("REDACTED", serialized)

    def test_validity_gate_distinguishes_harness_outcomes_from_framework_errors(self) -> None:
        for item, attribution in (
            ({"id": "TASK_TIMED_OUT", "attribution": "model"}, "model"),
            ({"id": "TOOL_CALLS_ALL_REJECTED", "attribution": "model"}, "model"),
            ({"id": "EXECUTION_ERROR", "attribution": "harness"}, "harness"),
        ):
            item.update({"validity_impact": "none", "rerun_action": "do_not_rerun"})
            severity, actual_attribution, _ = validity_check.classify_run_anomaly(item, {})
            self.assertEqual(severity, "info")
            self.assertEqual(actual_attribution, attribution)

        severity, attribution, _ = validity_check.classify_run_anomaly(
            {"id": "EXECUTION_ERROR", "severity": "error",
             "attribution": "evaluation_framework", "validity_impact": "fail",
             "rerun_action": "required_after_fix"},
            {},
        )
        self.assertEqual(severity, "error")
        self.assertEqual(attribution, "evaluation_framework")

        severity, attribution, _ = validity_check.classify_run_anomaly(
            {"id": "EXECUTION_ERROR", "severity": "error"},
            {},
        )
        self.assertEqual(severity, "warning")
        self.assertEqual(attribution, "undetermined")

        severity, attribution, recommendation = validity_check.classify_run_anomaly(
            {
                "id": "MODEL_API_RATE_LIMIT",
                "severity": "warning",
                "attribution": "external_service",
                "validity_impact": "review",
                "rerun_action": "review_first",
            },
            {},
        )
        self.assertEqual(severity, "warning")
        self.assertEqual(attribution, "external_service")
        self.assertIn("人工归因", recommendation)

    def test_model_api_errors_across_units_do_not_become_common_mode_failure(self) -> None:
        result_root = Path(self.temp_dir.name) / "api-review-round"
        for model in ("model-a", "model-b"):
            run_dir = result_root / model / "harness-y/01_suite/task_api/run_001"
            run_dir.mkdir(parents=True)
            (run_dir / "execution_status.json").write_text(json.dumps({
                "status": "finished", "timed_out": False, "elapsed_time": 30,
                "model": model, "harness": "harness-y",
            }), encoding="utf-8")
            (run_dir / "usage.json").write_text(json.dumps({
                "request_count": 1, "total_tokens": 100,
            }), encoding="utf-8")
            (run_dir / "score.json").write_text(
                json.dumps({"overall_score": 0.5}), encoding="utf-8"
            )
            events = [{"type": "event", "payload": {"index": index}} for index in range(5)]
            (run_dir / "chat.jsonl").write_text(
                "\n".join(json.dumps(event) for event in events), encoding="utf-8"
            )
            (run_dir / "runtime_events.jsonl").write_text(json.dumps({
                "stage": "model_inference", "model": model,
                "event_type": "request_failed", "http_status": 429,
            }) + "\n", encoding="utf-8")

        report = validity_check.scan_round(result_root, None)
        ids = [item["id"] for item in report["findings"]]
        self.assertEqual(ids.count("MODEL_API_RATE_LIMIT"), 2)
        self.assertNotIn("COMMON_MODE_ENV_FAILURE", ids)
        api_findings = [item for item in report["findings"]
                        if item["id"] == "MODEL_API_RATE_LIMIT"]
        self.assertTrue(all(item["severity"] == "warning" for item in api_findings))

    def test_scoped_batches_are_isolated_and_merge_without_overwrite(self) -> None:
        workspace = self.round_dir / "report-workspace"
        analysis_utils.save_batch_result(
            [{"task_id": "task_50", "result_analysis": "a", "root_cause_analysis": "b"}],
            workspace, "model-x@harness-y", 0, "lt60",
        )
        analysis_utils.save_batch_result(
            [{"task_id": "task_70", "result_analysis": "c", "root_cause_analysis": "d"}],
            workspace, "model-x@harness-y", 0, "gte60_lt80",
        )
        analysis_utils.save_batch_result(
            [{"task_id": "task_unscored", "result_analysis": "e", "root_cause_analysis": "f"}],
            workspace, "model-x@harness-y", 0, "lt60",
        )

        self.assertEqual(
            set(analysis_utils.load_completed_tasks(workspace, "model-x@harness-y", "lt60")),
            {"task_50", "task_unscored"},
        )
        self.assertEqual(
            set(analysis_utils.load_completed_tasks(workspace, "model-x@harness-y", "gte60_lt80")),
            {"task_70"},
        )
        final_path = analysis_utils.merge_all_batches(workspace, "model-x@harness-y", "lt60")
        self.assertEqual(final_path.name, "analysis_model-x@harness-y__lt60.json")

    def test_excel_loader_accepts_scoped_analysis_name(self) -> None:
        workspace = self.round_dir / "report-workspace"
        workspace.mkdir(parents=True, exist_ok=True)
        path = workspace / "analysis_model-x@harness-y__lt60.json"
        path.write_text(json.dumps({
            "task_50": {"result_analysis": "结果", "root_cause_analysis": "根因"}
        }), encoding="utf-8")

        unit = type("Unit", (), {"unit": "model-x@harness-y"})()
        loaded = excel_report.load_analysis([str(path)], [unit])
        self.assertEqual(loaded["model-x@harness-y::task_50"]["result_analysis"], "结果")
        self.assertEqual(excel_report.round_root_from_unit_dir(self.unit_dir), self.round_dir.resolve())

        three_level_unit = self.round_dir / "opencode" / "model-z" / "opencode"
        three_level_unit.mkdir(parents=True)
        self.assertEqual(excel_report.round_root_from_unit_dir(three_level_unit), self.round_dir.resolve())

    def test_excel_cli_backfills_scoped_analysis_in_detail_sheet(self) -> None:
        workspace = self.round_dir / "report-workspace"
        workspace.mkdir(parents=True, exist_ok=True)
        analysis_path = workspace / "analysis_model-x@harness-y__lt60.json"
        analysis_path.write_text(json.dumps({
            "task_50": {"result_analysis": "结果证据", "root_cause_analysis": "根因证据"}
        }), encoding="utf-8")

        subprocess.run(
            [sys.executable, str(EXCEL_SCRIPT), "--result-root", str(self.unit_dir),
             "--analysis", f"model-x@harness-y={analysis_path}",
             "--tasks-dir", str(Path(self.temp_dir.name) / "missing-tasks")],
            check=True,
            capture_output=True,
            text=True,
        )
        output_dir = self.round_dir / "report-workspace/output"
        workbooks = list(output_dir.glob("report_1units_*.xlsx"))
        self.assertEqual(len(workbooks), 1)
        workbook = load_workbook(workbooks[0], read_only=True)
        sheet = workbook["评分详情_model-x@harness-y"]
        headers = [cell.value for cell in sheet[1]]
        result_col = headers.index("结果分析") + 1
        root_col = headers.index("根因分析") + 1
        task_id_col = headers.index("用例ID") + 1
        row = next(row for row in range(2, sheet.max_row + 1)
                   if sheet.cell(row, task_id_col).value == "task_50")
        self.assertEqual(sheet.cell(row, result_col).value, "结果证据")
        self.assertEqual(sheet.cell(row, root_col).value, "根因证据")

    def generate_auditable_excel(self) -> Path:
        output_dir = self.round_dir / "audit-fixture-output"
        subprocess.run(
            [sys.executable, str(EXCEL_SCRIPT), "--result-root", str(self.unit_dir),
             "--tasks-dir", str(self.tasks_dir), "--output-dir", str(output_dir)],
            check=True,
            capture_output=True,
            text=True,
        )
        return next(output_dir.glob("report_1units_*.xlsx"))

    def test_audit_reconciles_generated_workbook(self) -> None:
        excel_path = self.generate_auditable_excel()
        report = report_audit.audit_report(self.unit_dir, excel_path, self.tasks_dir, None)
        self.assertFalse([item for item in report["findings"] if item["severity"] == "error"])
        self.assertEqual(report["verdict"], "REVIEW")  # 缺少可选 validity，仅需人工复核。

    def test_audit_normalizes_nested_harness_version_label(self) -> None:
        self.assertEqual(report_audit.normalize_harness("openclaw ((9f68ba9))"), "openclaw")
        self.assertEqual(report_audit.normalize_harness("opencode (1.18.4)"), "opencode")

    def test_excel_execution_errors_exclude_evaluation_failures(self) -> None:
        harness_run = self.paths["task_50"]
        framework_run = self.paths["task_70"]
        timeout_run = self.paths["task_almost_full"]
        (harness_run / "execution_status.json").write_text(json.dumps({
            "status": "error", "failure_stage": "astroncode_running",
            "error": "AstronCode run failed (rc=1)",
        }), encoding="utf-8")
        (framework_run / "execution_status.json").write_text(json.dumps({
            "status": "error", "failure_stage": "preparing_workspace",
            "error": "workspace preparation failed",
        }), encoding="utf-8")
        (timeout_run / "execution_status.json").write_text(json.dumps({
            "status": "timed_out", "timed_out": True, "error": "timed out",
        }), encoding="utf-8")

        excel_path = self.generate_auditable_excel()
        workbook = load_workbook(excel_path, read_only=True)
        sheet = workbook["总览"]
        headers = [cell.value for cell in sheet[1]]
        values = {header: sheet.cell(2, index + 1).value
                  for index, header in enumerate(headers)}
        self.assertEqual(values["正常完成数"], 2)
        self.assertEqual(values["执行错误数"], 1)
        self.assertEqual(values["超时数"], 1)
        # task_70 框架错误 + task_unscored 缺少 score.json。
        self.assertEqual(values["评测异常数"], 2)
        self.assertEqual(
            values["正常完成数"] + values["执行错误数"]
            + values["超时数"] + values["评测异常数"],
            values["用例数"],
        )

        report = report_audit.audit_report(self.unit_dir, excel_path, self.tasks_dir, None)
        self.assertFalse([item for item in report["findings"] if item["severity"] == "error"])

    def test_reliability_rerun_replaces_old_run_without_increasing_case_count(self) -> None:
        old_run = self.paths["task_50"]
        (old_run / "execution_status.json").write_text(json.dumps({
            "status": "error", "failure_stage": "astroncode_running",
            "error": "AstronCode run failed (rc=1)",
        }), encoding="utf-8")
        new_run = self.add_valid_run("task_50", "run_002", 0.9)
        write_rerun_metadata(
            new_run,
            supersedes_run=str(old_run),
            trigger="rerun_anomalous",
            task_id="task_50",
            model="model-x",
        )

        record = excel_report.TaskRecord("01_suite", old_run.parent, "harness-y")
        self.assertEqual(record.runs, 1)
        self.assertAlmostEqual(record.score, 0.9)
        manifest_record = next(
            item for item in manifest.scan_unit(
                "model-x", "harness-y", self.unit_dir, self.tasks_dir
            )
            if item["task_id"] == "task_50"
        )
        self.assertEqual(Path(manifest_record["run_dir"]), new_run)
        self.assertAlmostEqual(manifest_record["overall_score"], 0.9)

        excel_path = self.generate_auditable_excel()
        workbook = load_workbook(excel_path, read_only=True)
        overview = workbook["总览"]
        headers = [cell.value for cell in overview[1]]
        values = {header: overview.cell(2, index + 1).value
                  for index, header in enumerate(headers)}
        self.assertEqual(values["用例数"], 6)

        validity = validity_check.scan_round(self.unit_dir, self.tasks_dir)
        task = validity["units"]["model-x@harness-y"]["tasks"]["01_suite/task_50"]
        self.assertEqual(task["run_count"], 1)
        self.assertEqual(task["all_run_count"], 2)
        self.assertEqual(task["ignored_run_count"], 1)
        self.assertFalse(any(item["run_dir"] == str(old_run)
                             for item in validity["findings"]))

    def test_formal_valid_multirun_still_uses_all_runs(self) -> None:
        old_run = self.paths["task_50"]
        self.add_valid_run("task_50", "run_002", 0.9)

        record = excel_report.TaskRecord("01_suite", old_run.parent, "harness-y")
        self.assertEqual(record.runs, 2)
        self.assertAlmostEqual(record.score, 0.7)

    def test_task_record_reads_source_semantic_dimensions(self) -> None:
        run_dir = self.paths["task_50"]
        (run_dir / "score.json").write_text(json.dumps({
            "overall_score": 0.8,
            "_dimensions": {
                "metric_profile": "web-site-gen",
                "evidence_mode": "source_semantic",
                "primary": {
                    "content_structure": {
                        "score": 0.8, "weight": 0.2, "criterion_count": 2,
                    }
                },
                "secondary": {},
            },
        }), encoding="utf-8")

        record = excel_report.TaskRecord("01_suite", run_dir.parent, "harness-y")

        self.assertEqual(record.metric_dimensions["evidence_mode"], "source_semantic")
        self.assertEqual(
            record.metric_dimensions["primary"]["content_structure"]["score"], 0.8
        )

    def test_website_metrics_average_tasks_equally_and_write_semantic_scope(self) -> None:
        task_many_criteria = SimpleNamespace(
            task_id="website_many",
            metric_dimensions={
                "metric_profile": "web-site-gen",
                "evidence_mode": "source_semantic",
                "primary": {
                    "content_structure": {
                        "score": 1.0, "weight": 0.8, "criterion_count": 8,
                    }
                },
                "secondary": {
                    "basic_content": {
                        "score": 1.0, "weight": 0.8, "criterion_count": 8,
                        "primary": "content_structure",
                    }
                },
            },
        )
        task_one_criterion = SimpleNamespace(
            task_id="website_one",
            metric_dimensions={
                "metric_profile": "web-site-gen",
                "evidence_mode": "source_semantic",
                "primary": {
                    "content_structure": {
                        "score": 0.0, "weight": 0.2, "criterion_count": 1,
                    }
                },
                "secondary": {
                    "basic_content": {
                        "score": 0.0, "weight": 0.2, "criterion_count": 1,
                        "primary": "content_structure",
                    }
                },
            },
        )
        unit = SimpleNamespace(
            unit="model@harness",
            unit_display="Model@Harness",
            tasks=[task_many_criteria, task_one_criterion],
        )

        scores = excel_report._website_dimension_unit_scores(unit, "primary")
        self.assertEqual(scores["content_structure"], (50.0, 2))

        workbook = Workbook()
        workbook.remove(workbook.active)
        self.assertTrue(excel_report.write_website_metrics_sheet(workbook, [unit]))
        sheet = workbook["站点评测指标"]
        self.assertIn("源码语义评测", sheet[1][0].value)
        self.assertIn("不代表站点启动", sheet[1][0].value)
        first_column = [cell.value for cell in sheet["A"]]
        primary_title_row = first_column.index("一级维度汇总") + 1
        self.assertEqual(
            [sheet.cell(primary_title_row + 1, column).value for column in range(1, 5)],
            ["模型@Harness", "内容与结构", "交互与功能", "视觉与布局"],
        )
        self.assertEqual(
            [sheet.cell(primary_title_row + 2, column).value for column in range(1, 5)],
            ["Model@Harness", 50.0, "-", "-"],
        )

        empty_workbook = Workbook()
        self.assertFalse(excel_report.write_website_metrics_sheet(
            empty_workbook,
            [SimpleNamespace(unit="plain", unit_display="Plain", tasks=[])],
        ))
        self.assertNotIn("站点评测指标", empty_workbook.sheetnames)

        untyped_workbook = Workbook()
        self.assertFalse(excel_report.write_website_metrics_sheet(
            untyped_workbook,
            [SimpleNamespace(
                unit="untyped",
                unit_display="Untyped",
                tasks=[SimpleNamespace(
                    task_id="same_dimension_names",
                    metric_dimensions={
                        "evidence_mode": "source_semantic",
                        "primary": {"content_structure": {"score": 1.0}},
                    },
                )],
            )],
        ))
        self.assertNotIn("站点评测指标", untyped_workbook.sheetnames)

    def _create_website_task(
        self,
        task_id: str,
        runs: list[dict],
    ):
        task_dir = self.unit_dir / "07_Website_Generation" / task_id
        for index, run in enumerate(runs, 1):
            run_dir = task_dir / f"run_{index:03d}"
            run_dir.mkdir(parents=True)
            (run_dir / "score.json").write_text(json.dumps({
                "overall_score": run["score"],
                "_dimensions": {
                    "metric_profile": "web-site-gen",
                    "evidence_mode": "source_semantic",
                    "primary": {
                        "content_structure": {"score": run["score"]},
                        "interaction_function": {"score": run["score"]},
                        "visual_layout": {"score": run["score"]},
                    },
                    "secondary": {},
                },
            }), encoding="utf-8")
            (run_dir / "execution_status.json").write_text(json.dumps({
                "status": run.get("status", "completed"),
                "elapsed_time": run["elapsed_time"],
                "error": run.get("error", ""),
            }), encoding="utf-8")
            (run_dir / "usage.json").write_text(json.dumps({
                "input_tokens": run["input_tokens"],
                "cache_read_tokens": 0,
                "cache_write_tokens": 0,
                "output_tokens": run["output_tokens"],
                "total_tokens": run["input_tokens"] + run["output_tokens"],
                "request_count": 1,
            }), encoding="utf-8")
        registry = excel_report.report_entities.load_registry(
            REPORT_DIR / "data/entities.yaml"
        )
        task = excel_report.TaskRecord(
            "07_Website_Generation",
            task_dir,
            "astroncode",
            "xopglm52",
            registry,
            date(2026, 8, 12),
        )
        task.registry = registry
        return task

    def test_website_unit_metrics_include_result_and_run_efficiency(self) -> None:
        full_task = self._create_website_task("website_l1", [{
            "score": 1.0,
            "elapsed_time": 10,
            "input_tokens": 100,
            "output_tokens": 50,
        }])
        partial_task = self._create_website_task("website_l2", [{
            "score": 0.5,
            "status": "error",
            "error": "harness failed after producing output",
            "elapsed_time": 20,
            "input_tokens": 200,
            "output_tokens": 100,
        }])
        unit = SimpleNamespace(
            model="xopglm52",
            harness="astroncode",
            unit_display="GLM-5.2@AstronCode",
            registry=full_task.registry,
            pricing_date=date(2026, 8, 12),
            tasks=[full_task, partial_task],
        )

        metrics = excel_report._website_unit_metrics(unit, {
            "website_l1": {"difficulty": "L1"},
            "website_l2": {"difficulty": "L2"},
        })

        self.assertEqual(metrics["得分率"].value, 75.0)
        self.assertEqual(metrics["满分率"].value, 50.0)
        self.assertEqual(metrics["L1 题目得分率"].value, 100.0)
        self.assertEqual(metrics["L2 题目得分率"].value, 50.0)
        self.assertEqual(metrics["运行耗时平均值"].value, 15.0)
        self.assertEqual(metrics["运行耗时 P50"].value, 15.0)
        self.assertEqual(metrics["运行耗时 P90"].value, 19.0)
        self.assertEqual(metrics["单次运行平均总 Token"].value, 225.0)
        self.assertEqual(metrics["单次运行平均输入 Token"].value, 150.0)
        self.assertEqual(metrics["单次运行平均输出 Token"].value, 75.0)
        expected_cost = ((100 * 8 + 50 * 28) + (200 * 8 + 100 * 28)) \
            / 1_000_000 / 6.77 / 2
        self.assertAlmostEqual(metrics["单次运行平均成本"].value, expected_cost)
        self.assertEqual(metrics["单次运行平均成本"].sample, "2/2")

    def test_website_unit_metrics_exclude_superseded_runs(self) -> None:
        task = self._create_website_task("website_rerun", [
            {
                "score": 0.0,
                "elapsed_time": 100,
                "input_tokens": 900,
                "output_tokens": 100,
            },
            {
                "score": 1.0,
                "elapsed_time": 20,
                "input_tokens": 80,
                "output_tokens": 20,
            },
        ])
        write_rerun_metadata(
            task.run_dir,
            supersedes_run=str(task.run_dir.parent / "run_001"),
            trigger="regrade",
            task_id=task.task_id,
            model="xopglm52",
        )
        registry = excel_report.report_entities.load_registry(
            REPORT_DIR / "data/entities.yaml"
        )
        task = excel_report.TaskRecord(
            "07_Website_Generation",
            task.run_dir.parent,
            "astroncode",
            "xopglm52",
            registry,
            date(2026, 8, 12),
        )
        unit = SimpleNamespace(
            model="xopglm52",
            harness="astroncode",
            unit_display="GLM-5.2@AstronCode",
            registry=registry,
            pricing_date=date(2026, 8, 12),
            tasks=[task],
        )

        metrics = excel_report._website_unit_metrics(
            unit, {task.task_id: {"difficulty": "L1"}}
        )

        self.assertEqual(metrics["得分率"].value, 100.0)
        self.assertEqual(metrics["运行耗时平均值"].value, 20.0)
        self.assertEqual(metrics["单次运行平均总 Token"].value, 100.0)
        self.assertEqual(metrics["运行耗时平均值"].sample, 1)

    def test_website_unit_metrics_count_unscored_tagged_task_as_zero(self) -> None:
        full_task = self._create_website_task("website_full", [{
            "score": 1.0,
            "elapsed_time": 10,
            "input_tokens": 100,
            "output_tokens": 20,
        }])
        invalid_task = self._create_website_task("website_invalid", [{
            "score": 0.0,
            "status": "error",
            "error": "framework error before grading",
            "elapsed_time": 30,
            "input_tokens": 200,
            "output_tokens": 40,
        }])
        (invalid_task.run_dir / "score.json").unlink()
        registry = full_task.registry
        invalid_task = excel_report.TaskRecord(
            "07_Website_Generation",
            invalid_task.run_dir.parent,
            "astroncode",
            "xopglm52",
            registry,
            date(2026, 8, 12),
        )
        unit = SimpleNamespace(
            model="xopglm52",
            harness="astroncode",
            unit_display="GLM-5.2@AstronCode",
            registry=registry,
            pricing_date=date(2026, 8, 12),
            tasks=[full_task, invalid_task],
        )

        metrics = excel_report._website_unit_metrics(unit, {
            "website_full": {"difficulty": "L1", "tags": "web-site-gen"},
            "website_invalid": {"difficulty": "L2", "tags": "web-site-gen"},
        })

        self.assertEqual(metrics["得分率"].value, 50.0)
        self.assertEqual(metrics["满分率"].value, 50.0)
        self.assertEqual(metrics["L2 题目得分率"].value, 0.0)
        self.assertEqual(metrics["运行耗时平均值"].value, 20.0)
        self.assertEqual(metrics["单次运行平均总 Token"].value, 180.0)
        self.assertEqual(metrics["运行耗时平均值"].sample, 2)

    def test_website_metrics_sheet_keeps_all_unscored_tagged_tasks(self) -> None:
        task = self._create_website_task("website_invalid", [{
            "score": 0.0,
            "status": "error",
            "error": "framework error before grading",
            "elapsed_time": 30,
            "input_tokens": 200,
            "output_tokens": 40,
        }])
        (task.run_dir / "score.json").unlink()
        registry = task.registry
        task = excel_report.TaskRecord(
            "07_Website_Generation",
            task.run_dir.parent,
            "astroncode",
            "xopglm52",
            registry,
            date(2026, 8, 12),
        )
        unit = SimpleNamespace(
            model="xopglm52",
            harness="astroncode",
            unit_display="GLM-5.2@AstronCode",
            registry=registry,
            pricing_date=date(2026, 8, 12),
            tasks=[task],
        )
        workbook = Workbook()
        workbook.remove(workbook.active)

        written = excel_report.write_website_metrics_sheet(
            workbook,
            [unit],
            {"website_invalid": {"difficulty": "L1", "tags": "web-site-gen"}},
        )

        self.assertTrue(written)
        sheet = workbook["站点评测指标"]
        headers = [sheet.cell(4, column).value for column in range(2, 16)]
        self.assertEqual(sheet.cell(5, headers.index("得分率") + 2).value, 0.0)
        self.assertEqual(
            sheet.cell(5, headers.index("运行耗时平均值") + 2).value, 30.0
        )

    def test_website_metrics_sheet_uses_horizontal_grouped_headers(self) -> None:
        task = self._create_website_task("website_l1", [{
            "score": 1.0,
            "elapsed_time": 12,
            "input_tokens": 120,
            "output_tokens": 30,
        }])
        unit = SimpleNamespace(
            model="xopglm52",
            harness="astroncode",
            unit="xopglm52@astroncode",
            unit_display="GLM-5.2@AstronCode",
            registry=task.registry,
            pricing_date=date(2026, 8, 12),
            tasks=[task],
        )
        workbook = Workbook()
        workbook.remove(workbook.active)

        self.assertTrue(excel_report.write_website_metrics_sheet(
            workbook, [unit], {"website_l1": {"difficulty": "L1"}}
        ))

        sheet = workbook["站点评测指标"]
        merged = {str(item) for item in sheet.merged_cells.ranges}
        self.assertTrue({"A3:A4", "B3:C3", "D3:H3", "I3:O3"}.issubset(merged))
        self.assertEqual(sheet["A3"].value, "模型@Harness")
        self.assertEqual(sheet["B3"].value, "结果指标")
        self.assertEqual(sheet["D3"].value, "分层分析")
        self.assertEqual(sheet["I3"].value, "效率指标")
        self.assertEqual(
            [sheet.cell(4, column).value for column in range(1, 16)],
            [
                None, "得分率", "满分率", "L1 题目得分率", "L2 题目得分率",
                "内容与结构得分率", "交互与功能得分率", "视觉与布局得分率",
                "运行耗时平均值", "运行耗时 P50", "运行耗时 P90", "单次运行平均成本",
                "单次运行平均总 Token", "单次运行平均输入 Token", "单次运行平均输出 Token",
            ],
        )
        self.assertEqual(sheet.cell(5, 1).value, "GLM-5.2@AstronCode")
        self.assertEqual(sheet.cell(5, 2).value, 100.0)
        self.assertAlmostEqual(
            sheet.cell(5, 12).value,
            (120 * 8 + 30 * 28) / 1_000_000 / 6.77,
        )
        self.assertTrue(all(
            isinstance(sheet.cell(5, column).value, (int, float))
            or sheet.cell(5, column).value == "-"
            for column in range(2, 16)
        ))
        self.assertEqual(sheet.cell(5, 5).value, "-")
        self.assertEqual(sheet.freeze_panes, "B5")

        self.assertIn("_站点评测指标口径", workbook.sheetnames)
        self.assertEqual(workbook["_站点评测指标口径"].sheet_state, "hidden")
        self.assertEqual(workbook["_站点评测指标口径"].max_row, 15)

    def test_website_dimension_summaries_use_horizontal_grouped_tables(self) -> None:
        task = SimpleNamespace(
            task_id="website_dimensions",
            score=0.9,
            effective_score=0.9,
            effective_run_dirs=[],
            metric_dimensions={
                "metric_profile": "web-site-gen",
                "evidence_mode": "source_semantic",
                "primary": {
                    "content_structure": {"score": 0.9},
                    "interaction_function": {"score": 0.8},
                    "visual_layout": {"score": 0.7},
                },
                "secondary": {
                    "basic_content": {"score": 0.9, "primary": "content_structure"},
                    "information_organization": {
                        "score": 0.8, "primary": "content_structure"
                    },
                    "page_navigation": {
                        "score": 0.7, "primary": "interaction_function"
                    },
                    "operation_feedback": {
                        "score": 0.6, "primary": "interaction_function"
                    },
                    "visual_style": {"score": 0.95, "primary": "visual_layout"},
                    "page_layout": {"score": 0.85, "primary": "visual_layout"},
                    "future_metric": {"score": 0.5, "primary": "future_primary"},
                },
            },
        )
        unit = SimpleNamespace(
            unit="model@harness",
            unit_display="Model@Harness",
            tasks=[task],
        )
        workbook = Workbook()
        workbook.remove(workbook.active)

        self.assertTrue(excel_report.write_website_metrics_sheet(workbook, [unit]))

        sheet = workbook["站点评测指标"]
        first_column = [cell.value for cell in sheet["A"]]
        primary_title_row = first_column.index("一级维度汇总") + 1
        self.assertEqual(
            [sheet.cell(primary_title_row + 1, column).value for column in range(1, 5)],
            ["模型@Harness", "内容与结构", "交互与功能", "视觉与布局"],
        )
        self.assertEqual(
            [sheet.cell(primary_title_row + 2, column).value for column in range(1, 5)],
            ["Model@Harness", 90.0, 80.0, 70.0],
        )

        secondary_title_row = first_column.index("二级维度汇总") + 1
        group_row = secondary_title_row + 1
        label_row = secondary_title_row + 2
        self.assertEqual(sheet.cell(group_row, 1).value, "模型@Harness")
        merged = {str(item) for item in sheet.merged_cells.ranges}
        self.assertIn(f"A{group_row}:A{label_row}", merged)
        self.assertIn(f"B{group_row}:C{group_row}", merged)
        self.assertIn(f"D{group_row}:E{group_row}", merged)
        self.assertIn(f"F{group_row}:G{group_row}", merged)
        self.assertEqual(
            [sheet.cell(group_row, column).value for column in (2, 4, 6, 8)],
            ["内容与结构", "交互与功能", "视觉与布局", "其他"],
        )
        self.assertEqual(
            [sheet.cell(label_row, column).value for column in range(2, 9)],
            [
                "基础内容", "信息组织", "页面导航", "操作反馈",
                "视觉风格", "页面布局", "future_metric",
            ],
        )
        self.assertEqual(
            [sheet.cell(label_row + 1, column).value for column in range(1, 9)],
            ["Model@Harness", 90.0, 80.0, 70.0, 60.0, 95.0, 85.0, 50.0],
        )
        conditional_ranges = {str(item.sqref) for item in sheet.conditional_formatting}
        self.assertIn(
            f"B{primary_title_row + 2}:D{primary_title_row + 2}",
            conditional_ranges,
        )
        self.assertIn(
            f"B{label_row + 1}:H{label_row + 1}",
            conditional_ranges,
        )

    def test_website_metrics_sheet_writes_complete_summary(self) -> None:
        task = self._create_website_task("website_l1", [{
            "score": 1.0,
            "elapsed_time": 12,
            "input_tokens": 120,
            "output_tokens": 30,
        }])
        unit = SimpleNamespace(
            model="xopglm52",
            harness="astroncode",
            unit="xopglm52@astroncode",
            unit_display="GLM-5.2@AstronCode",
            registry=task.registry,
            pricing_date=date(2026, 8, 12),
            tasks=[task],
        )
        workbook = Workbook()
        workbook.remove(workbook.active)

        self.assertTrue(excel_report.write_website_metrics_sheet(
            workbook, [unit], {"website_l1": {"difficulty": "L1"}}
        ))

        sheet = workbook["站点评测指标"]
        values = list(sheet.iter_rows(values_only=True))
        metric_names = [sheet.cell(4, column).value for column in range(2, 16)]
        self.assertEqual(len(metric_names), 14)
        self.assertNotIn("美观度", metric_names)
        self.assertIn("满分率", metric_names)
        self.assertIn("单次运行平均总 Token", metric_names)
        self.assertIn("单次运行平均输入 Token", metric_names)
        self.assertIn("单次运行平均输出 Token", metric_names)
        self.assertTrue(all(
            isinstance(sheet.cell(5, column).value, (int, float))
            or sheet.cell(5, column).value == "-"
            for column in range(2, 16)
        ))
        self.assertTrue(any(row[0] == "一级维度汇总" for row in values))
        primary_title_row = next(
            index for index, row in enumerate(values, 1)
            if row[0] == "一级维度汇总"
        )
        self.assertEqual(
            [sheet.cell(primary_title_row + 1, column).value for column in range(1, 5)],
            ["模型@Harness", "内容与结构", "交互与功能", "视觉与布局"],
        )

    def test_leader_extractor_reads_horizontal_website_metrics(self) -> None:
        task = self._create_website_task("website_l1", [{
            "score": 1.0,
            "elapsed_time": 12,
            "input_tokens": 120,
            "output_tokens": 30,
        }])
        unit = SimpleNamespace(
            model="xopglm52",
            harness="astroncode",
            unit="xopglm52@astroncode",
            unit_display="GLM-5.2@AstronCode",
            registry=task.registry,
            pricing_date=date(2026, 8, 12),
            tasks=[task],
        )
        workbook = Workbook()
        workbook.remove(workbook.active)
        excel_report.write_website_metrics_sheet(
            workbook, [unit], {"website_l1": {"difficulty": "L1"}}
        )
        web_path = Path(self.temp_dir.name) / "website.xlsx"
        workbook.save(web_path)

        leader_extract = load_module("leader_extract_website", LEADER_EXTRACT_SCRIPT)
        web_metrics = leader_extract._extract_website_metrics(
            load_workbook(web_path, data_only=True)
        )

        self.assertEqual(
            web_metrics["GLM-5.2@AstronCode"]["满分率"]["value"], 100.0
        )
        self.assertEqual(
            web_metrics["GLM-5.2@AstronCode"]["满分率"]["category"], "结果指标"
        )
        self.assertEqual(
            web_metrics["GLM-5.2@AstronCode"]["满分率"]["sample"], 1
        )
        self.assertIn(
            "overall_score = 1.0",
            web_metrics["GLM-5.2@AstronCode"]["满分率"]["method"],
        )
        self.assertEqual(len(web_metrics["GLM-5.2@AstronCode"]), 14)
        self.assertNotIn("美观度", web_metrics["GLM-5.2@AstronCode"])
        self.assertEqual(
            leader_extract._extract_website_metrics(Workbook()), {}
        )

    def test_leader_extractor_keeps_legacy_vertical_website_compatibility(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "站点评测指标"
        sheet.append(["结果与效率指标汇总"])
        sheet.append([
            "模型@Harness", "指标分类", "指标名称", "数值", "样本数", "计算方法",
        ])
        sheet.append([
            "GLM-5.2@AstronCode", "结果指标", "得分率", 82.5, 6,
            "各任务 overall_score 按任务等权平均",
        ])

        leader_extract = load_module("leader_extract_legacy_website", LEADER_EXTRACT_SCRIPT)

        self.assertEqual(
            leader_extract._extract_website_metrics(workbook),
            {
                "GLM-5.2@AstronCode": {
                    "得分率": {
                        "category": "结果指标",
                        "value": 82.5,
                        "sample": 6,
                        "method": "各任务 overall_score 按任务等权平均",
                    }
                }
            },
        )

    def test_report_skill_documents_website_metrics(self) -> None:
        template = (REPORT_DIR / "skills/eval-report/references/report_template.md").read_text(
            encoding="utf-8"
        )
        skill = (REPORT_DIR / "skills/eval-report/SKILL.md").read_text(encoding="utf-8")

        self.assertIn("站点评测指标", template)
        self.assertIn("leader_data.website_metrics", template)
        self.assertIn("无 Web 指标时整节省略", template)
        self.assertIn("结果与效率指标汇总", skill)
        self.assertIn("不展示美观度", skill)
        self.assertIn("_站点评测指标口径", skill)
        self.assertIn("正式表只展示", skill)

    def test_audit_detects_request_count_regression(self) -> None:
        excel_path = self.generate_auditable_excel()
        workbook = load_workbook(excel_path)
        sheet = workbook["总览"]
        headers = [cell.value for cell in sheet[1]]
        request_column = headers.index("总请求数") + 1
        sheet.cell(2, request_column).value = 999
        workbook.save(excel_path)

        report = report_audit.audit_report(self.unit_dir, excel_path, self.tasks_dir, None)
        errors = [item for item in report["findings"] if item["severity"] == "error"]
        self.assertTrue(any(item["id"] == "OVERVIEW_VALUE_MISMATCH" for item in errors))

    def test_audit_propagates_upstream_validity_failure(self) -> None:
        excel_path = self.generate_auditable_excel()
        validity_path = self.round_dir / "report-workspace/validity/eval_result_validity.json"
        validity_path.parent.mkdir(parents=True, exist_ok=True)
        validity_path.write_text(json.dumps({
            "verdict": "FAIL", "summary": {"errors": 1, "warnings": 0}
        }), encoding="utf-8")
        report = report_audit.audit_report(
            self.unit_dir, excel_path, self.tasks_dir, validity_path
        )
        self.assertEqual(report["verdict"], "FAIL")
        self.assertTrue(any(item["id"] == "UPSTREAM_VALIDITY_FAILED"
                            for item in report["findings"]))

    def test_difficulty_inversion_is_review_not_failure(self) -> None:
        units = {}
        for index in range(5):
            tasks = {}
            for task_index in range(3):
                tasks[f"easy_{task_index}"] = {"score": 0.2}
                tasks[f"hard_{task_index}"] = {"score": 0.8}
            units[f"model-{index}@harness"] = {"tasks": tasks}
        groups = {
            "L2": {f"easy_{index}" for index in range(3)},
            "L3": {f"hard_{index}" for index in range(3)},
        }
        findings = []
        report_audit.audit_difficulty_inversion(units, groups, findings)
        self.assertEqual(len(findings), 1)
        self.assertEqual(findings[0]["id"], "DIFFICULTY_INVERSION")
        self.assertEqual(findings[0]["severity"], "warning")

    def test_eval_report_skill_contract_for_controlled_views(self) -> None:
        skill = (REPORT_DIR / "skills/eval-report/SKILL.md").read_text(encoding="utf-8")
        template = (REPORT_DIR / "skills/eval-report/references/report_template.md").read_text(
            encoding="utf-8"
        )
        readme = (REPORT_DIR / "README.md").read_text(encoding="utf-8")

        self.assertLess(skill.index("二、分类维度"), skill.index("三、Agent能力"))
        self.assertIn("固定目标 Harness", skill)
        self.assertIn("固定目标模型", skill)
        self.assertIn("总览保留全部列", skill)
        self.assertIn("extract_leader_report_data.py", skill)
        self.assertIn("最多 2 句话", skill)
        self.assertIn("3 个关键数字", skill)
        self.assertIn("优先排查", skill)
        self.assertIn("L3/L4", skill)
        self.assertIn("不得进入典型低分案例", skill)
        self.assertIn("先说明已有优势", skill)
        self.assertIn("数字用于支撑判断", skill)
        self.assertIn("对 Harness 选择较敏感", skill)
        self.assertIn("避免使用", skill)

        self.assertLess(template.index("## 二、分类维度分析"),
                        template.index("## 三、Agent 能力维度分析"))
        self.assertGreaterEqual(template.count("固定目标 Harness"), 4)
        self.assertGreaterEqual(template.count("固定目标模型"), 4)
        self.assertIn("全部 Excel 总览列", template)
        self.assertIn("评测有效性与剔除说明", template)
        self.assertIn("先写已有基础或接近项", template)
        self.assertIn("分差均在 5 分以内", template)

        for option in (
            "--target-model", "--target-harness", "--models", "--harnesses",
            "--entities", "--pricing-date",
        ):
            self.assertIn(option, readme)


if __name__ == "__main__":
    unittest.main()
