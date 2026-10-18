import importlib.util
import json
import sys
import tempfile
import time
import unittest
from datetime import datetime, timedelta, timezone
from pathlib import Path
from unittest.mock import patch

import pyarrow.parquet as pq
from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[1]
SCRIPT_PATH = (
    REPO_ROOT / "tools" / "astroncode_usage" / "analyze_astroncode_prod_usage.py"
)
SPEC = importlib.util.spec_from_file_location("astroncode_prod_usage", SCRIPT_PATH)
assert SPEC and SPEC.loader
MODULE = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = MODULE
SPEC.loader.exec_module(MODULE)


class AstronCodeProdUsageTests(unittest.TestCase):
    def test_safe_excel_text_removes_invalid_xml_unicode_before_formula_guard(self) -> None:
        cleaned, truncated = MODULE.safe_excel_text("示例A\ud800\ufffeB\uffff\x01")
        self.assertEqual(cleaned, "示例AB")
        self.assertFalse(truncated)
        cleaned, truncated = MODULE.safe_excel_text("\ufffe=1+1\uffff")
        self.assertEqual(cleaned, "'=1+1")
        self.assertFalse(truncated)

    def test_client_disables_ambient_proxy_config(self) -> None:
        client = MODULE.ElasticsearchClient(
            "http://127.0.0.1:9200",
            timeout_seconds=1.0,
            retries=0,
            ca_cert=None,
        )

        self.assertFalse(client.session.trust_env)

    def test_extracts_otlp_envelope_record(self) -> None:
        attributes = {
            "gen_ai.input.messages": json.dumps(
                [
                    {
                        "role": "user",
                        "parts": [
                            {"type": "text", "content": "请修复登录接口的 500 错误"}
                        ],
                    }
                ],
                ensure_ascii=False,
            ),
            "gen_ai.request.model": "gpt-5.5",
            "gen_ai.request.reasoning_effort": "high",
        }
        hit = {
            "_index": "prod-1",
            "_id": "doc-1",
            "_source": {
                "@timestamp": "2026-09-01T08:00:00Z",
                "event": "agent.turn",
                "Attributes": json.dumps(attributes, ensure_ascii=False),
                "Duration": "1234000",
                "TraceID": "trace-1",
                "SpanID": "span-1",
            },
        }

        row = MODULE.extract_usage_row(hit, "Asia/Shanghai")

        self.assertEqual(row.query, "请修复登录接口的 500 错误")
        self.assertEqual(row.query_length, len(row.query))
        self.assertEqual(row.occurred_at.isoformat(), "2026-09-01T16:00:00")
        self.assertEqual(row.model, "gpt-5.5")
        self.assertEqual(row.reasoning_effort, "high")
        self.assertAlmostEqual(row.duration_seconds, 1.234)
        self.assertEqual(row.fingerprint, "trace-1|span-1")

    def test_missing_reasoning_effort_is_not_fabricated(self) -> None:
        hit = {
            "_index": "prod-1",
            "_id": "doc-2",
            "_source": {
                "Attributes": json.dumps(
                    {
                        "gen_ai.input.messages": [
                            {
                                "role": "user",
                                "parts": [
                                    {"type": "text", "content": "生成一个 Python 脚本"}
                                ],
                            }
                        ],
                        "gen_ai.request.model": "model-a",
                    },
                    ensure_ascii=False,
                ),
                "span": {"duration_us": 2_000_000},
            },
        }

        row = MODULE.extract_usage_row(hit, "Asia/Shanghai")

        self.assertEqual(row.reasoning_effort, "未记录")
        self.assertEqual(row.duration_seconds, 2.0)

    def test_date_only_end_includes_full_day(self) -> None:
        now = datetime(2026, 9, 14, 12, tzinfo=timezone(timedelta(hours=8)))
        window = MODULE.requested_window(
            "2026-09-01", "2026-09-14", "Asia/Shanghai", now=now
        )

        self.assertEqual(window.start.isoformat(), "2026-09-01T00:00:00+08:00")
        self.assertEqual(window.end.isoformat(), "2026-09-15T00:00:00+08:00")

    def test_large_month_falls_back_to_two_weeks(self) -> None:
        requested = MODULE.TimeWindow(
            datetime(2026, 8, 1, tzinfo=timezone.utc),
            datetime(2026, 9, 1, tzinfo=timezone.utc),
        )
        seen = []

        def count_records(window):
            seen.append(window)
            return 150_000 if len(seen) == 1 else 60_000

        effective, count, changed = MODULE.choose_effective_window(
            requested,
            count_records=count_records,
            max_records=100_000,
            fallback_days=14,
            auto_fallback=True,
        )

        self.assertTrue(changed)
        self.assertEqual(count, 60_000)
        self.assertEqual(effective.start, requested.end - timedelta(days=14))

    def test_category_rules_are_bounded_and_classify(self) -> None:
        version, rules = MODULE.load_category_rules(MODULE.DEFAULT_CATEGORY_RULES)

        self.assertEqual(version, "astroncode-scene-v3")
        self.assertEqual(len(rules), 16)
        self.assertEqual(
            MODULE.classify_query("请修复 React 页面白屏问题", rules),
            "问题排查与修复",
        )
        self.assertEqual(
            MODULE.classify_query("帮我写一个 React 响应式页面", rules),
            "前端、网站与交互开发",
        )
        self.assertEqual(
            MODULE.classify_query("把这份技术方案整理成 Markdown 文档", rules),
            "文档写作与办公处理",
        )
        self.assertEqual(
            MODULE.classify_query("调研主流 RAG 评测方法并给出资料来源", rules),
            "信息检索与研究",
        )
        self.assertEqual(
            MODULE.classify_query("整理产品需求并拆分用户故事和验收标准", rules),
            "产品需求与项目协作",
        )
        self.assertEqual(
            MODULE.classify_query("把下载目录中的文件移动到归档文件夹", rules),
            "文件与本地资源操作",
        )
        self.assertEqual(
            MODULE.classify_query("给这张图片去除水印并生成视频封面", rules),
            "图像音视频与设计",
        )
        self.assertEqual(
            MODULE.classify_query("请修复图片上传错误", rules),
            "问题排查与修复",
        )
        self.assertEqual(
            MODULE.classify_query("继续", rules),
            "未分类",
        )
        self.assertEqual(
            MODULE.classify_query("按照上面调整", rules),
            "未分类",
        )
        self.assertEqual(
            MODULE.classify_query("", rules),
            "未分类",
        )

    def test_v1_category_rules_remain_available(self) -> None:
        version, rules = MODULE.load_category_rules(
            MODULE.DEFAULT_CATEGORY_RULES.with_name("scene_categories_v1.json")
        )

        self.assertEqual(version, "astroncode-scene-v1")
        self.assertEqual(len(rules), 13)
        self.assertEqual(rules[-1].name, "其他或复合任务")
        self.assertEqual(MODULE.classify_query("", rules), "其他或复合任务")

    def test_context_dependent_queries_are_kept_unclassified(self) -> None:
        self.assertTrue(MODULE.is_context_dependent_query("继续"))
        self.assertTrue(MODULE.is_context_dependent_query("继续任务"))
        self.assertTrue(MODULE.is_context_dependent_query("继续执行"))
        self.assertTrue(MODULE.is_context_dependent_query("继续完成任务"))
        self.assertTrue(MODULE.is_context_dependent_query("重新执行"))
        self.assertTrue(MODULE.is_context_dependent_query("重新生成"))
        self.assertTrue(MODULE.is_context_dependent_query("再看下"))
        self.assertTrue(MODULE.is_context_dependent_query("提交"))
        self.assertTrue(MODULE.is_context_dependent_query("按照建议实施"))
        self.assertTrue(MODULE.is_context_dependent_query("这个怎么修改？"))
        self.assertFalse(MODULE.is_context_dependent_query("按照接口文档实现登录功能"))
        self.assertFalse(MODULE.is_context_dependent_query("重新生成项目月度报告"))
        self.assertFalse(MODULE.is_context_dependent_query("继续分析登录错误原因"))

    def test_secondary_candidate_requires_explicit_task_intent(self) -> None:
        self.assertTrue(
            MODULE.is_secondary_classification_candidate(
                "测试一下 web fetch 工具(get)"
            )
        )
        self.assertTrue(
            MODULE.is_secondary_classification_candidate("你有pdf解析工具吗?")
        )
        self.assertTrue(
            MODULE.is_secondary_classification_candidate(
                "列出来你现在能看见的插件版本"
            )
        )
        self.assertFalse(MODULE.is_secondary_classification_candidate("你好你好你好你好"))
        self.assertFalse(MODULE.is_secondary_classification_candidate("12345678"))
        self.assertFalse(MODULE.is_secondary_classification_candidate("继续完成任务"))

    def test_semantic_response_parser_validates_structured_output(self) -> None:
        content = """```json
        {"items":[{"id":"Q001","category":"图像音视频与设计","confidence":0.92,"reason":"图像编辑"}]}
        ```"""

        parsed = MODULE.parse_semantic_response(
            content,
            {"Q001"},
            {"图像音视频与设计", "未分类"},
        )

        self.assertEqual(parsed["Q001"], ("图像音视频与设计", 0.92, "图像编辑"))

    def test_semantic_client_keeps_single_rejected_query_unclassified(self) -> None:
        _, rules = MODULE.load_category_rules(MODULE.DEFAULT_CATEGORY_RULES)

        class FakeResponse:
            status_code = 400

            def raise_for_status(self):
                raise MODULE.requests.HTTPError("bad request", response=self)

        class FakeSession:
            trust_env = True
            post_count = 0

            def post(self, *args, **kwargs):
                del args, kwargs
                self.post_count += 1
                return FakeResponse()

            def close(self):
                return None

        session = FakeSession()
        client = MODULE.SemanticSceneClient(
            base_url="https://example.invalid/v1",
            api_key="test-key",
            model="test-model",
            rules=rules,
            timeout_seconds=1.0,
            retries=3,
            request_interval=0.0,
        )
        with patch.object(MODULE.requests, "Session", return_value=session):
            parsed = client.classify_batch([("Q001", "test", False)])

        self.assertEqual(session.post_count, 1)
        self.assertEqual(
            parsed["Q001"],
            ("未分类", 0.0, "分类服务拒绝单条请求"),
        )

    def test_semantic_client_splits_rejected_batch(self) -> None:
        _, rules = MODULE.load_category_rules(MODULE.DEFAULT_CATEGORY_RULES)

        class FakeResponse:
            def __init__(self, status_code, items):
                self.status_code = status_code
                self.items = items

            def raise_for_status(self):
                if self.status_code >= 400:
                    raise MODULE.requests.HTTPError("bad request", response=self)

            def json(self):
                content = {
                    "items": [
                        {
                            "id": item["id"],
                            "category": "通用代码与功能开发",
                            "confidence": 0.9,
                            "reason": "测试",
                        }
                        for item in self.items
                    ]
                }
                return {"choices": [{"message": {"content": json.dumps(content)}}]}

        class FakeSession:
            trust_env = True
            post_count = 0

            def post(self, *args, **kwargs):
                del args
                self.post_count += 1
                request_items = json.loads(kwargs["json"]["messages"][1]["content"])[
                    "items"
                ]
                return FakeResponse(
                    400 if self.post_count == 1 else 200,
                    request_items,
                )

            def close(self):
                return None

        session = FakeSession()
        client = MODULE.SemanticSceneClient(
            base_url="https://example.invalid/v1",
            api_key="test-key",
            model="test-model",
            rules=rules,
            timeout_seconds=1.0,
            retries=3,
            request_interval=0.0,
        )
        with patch.object(MODULE.requests, "Session", return_value=session):
            parsed = client.classify_batch(
                [("Q001", "alpha", False), ("Q002", "beta", False)]
            )

        self.assertEqual(session.post_count, 3)
        self.assertEqual(set(parsed), {"Q001", "Q002"})
        self.assertTrue(
            all(item[0] == "通用代码与功能开发" for item in parsed.values())
        )

    def test_semantic_failure_does_not_drain_all_queued_batches(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_dir:
            directory = Path(temporary_dir)
            input_path = directory / "normalized.parquet"
            output_path = directory / "classifications.parquet"
            checkpoint_path = directory / "classifications.checkpoint.jsonl"
            pq.write_table(
                MODULE.pa.table({"query": ["alpha", "beta", "gamma"]}),
                input_path,
            )

            class FakeClient:
                call_count = 0

                def __init__(self, **kwargs):
                    del kwargs

                def classify_batch(self, items):
                    type(self).call_count += 1
                    if type(self).call_count == 1:
                        raise RuntimeError("synthetic failure")
                    time.sleep(0.1)
                    return {
                        item_id: ("通用代码与功能开发", 0.9, "测试")
                        for item_id, _, _ in items
                    }

            with (
                patch.object(MODULE, "SemanticSceneClient", FakeClient),
                self.assertRaisesRegex(RuntimeError, "synthetic failure"),
            ):
                MODULE.classify_scenes_to_parquet(
                    input_path,
                    output_path,
                    checkpoint_path,
                    MODULE.DEFAULT_CATEGORY_RULES,
                    base_url="https://example.invalid/v1",
                    api_key="test-key",
                    model="test-model",
                    confidence_threshold=0.7,
                    batch_size=1,
                    batch_char_limit=100,
                    workers=2,
                    request_interval=0.0,
                    timeout_seconds=1.0,
                    retries=0,
                    resume=False,
                )

            self.assertEqual(FakeClient.call_count, 2)

    def test_refinement_reuses_base_and_runs_second_pass_only_for_candidates(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as temporary_dir:
            directory = Path(temporary_dir)
            input_path = directory / "normalized.parquet"
            base_path = directory / "base.parquet"
            output_path = directory / "refined.parquet"
            checkpoint_path = directory / "refined.checkpoint.jsonl"
            queries = [
                "继续任务",
                "测试一下 web fetch 工具(get)",
                "你好",
                "请修复登录错误",
            ]
            pq.write_table(MODULE.pa.table({"query": queries}), input_path)
            rule_version, rules = MODULE.load_category_rules(
                MODULE.DEFAULT_CATEGORY_RULES
            )
            base_classifications = {}
            for query in queries:
                _, query_key = MODULE.normalize_query(query)
                query_hash = MODULE.query_key_hash(query_key)
                is_rule = query == "请修复登录错误"
                category = (
                    "问题排查与修复"
                    if is_rule
                    else MODULE.LEGACY_FALLBACK_CATEGORY
                )
                base_classifications[query_hash] = MODULE.SceneClassification(
                    query_hash=query_hash,
                    category=category,
                    proposed_category=category,
                    method="规则" if is_rule else "大模型未判定",
                    confidence=1.0 if is_rule else 0.9,
                    reason="初次分类",
                    model="" if is_rule else "xopglm52",
                    prompt_version=MODULE.SCENE_PROMPT_VERSION,
                    rule_version="astroncode-scene-v2",
                    query_length=len(query),
                    input_truncated=False,
                )
            MODULE.write_scene_classification_parquet(
                base_classifications,
                base_path,
                rule_version="astroncode-scene-v2",
                model="xopglm52",
                confidence_threshold=0.70,
            )

            class FakeClient:
                seen_queries = []

                def __init__(self, **kwargs):
                    del kwargs

                def classify_batch(self, items):
                    type(self).seen_queries.extend(query for _, query, _ in items)
                    return {
                        item_id: ("测试、评审与质量", 0.92, "工具测试")
                        for item_id, _, _ in items
                    }

            with patch.object(MODULE, "SemanticSceneClient", FakeClient):
                stats = MODULE.refine_scene_classifications_to_parquet(
                    input_path,
                    base_path,
                    output_path,
                    checkpoint_path,
                    MODULE.DEFAULT_CATEGORY_RULES,
                    base_url="https://example.invalid/v1",
                    api_key="test-key",
                    model="xopglm52",
                    confidence_threshold=0.70,
                    minimum_query_length=8,
                    batch_size=40,
                    batch_char_limit=24_000,
                    workers=2,
                    request_interval=0.0,
                    timeout_seconds=1.0,
                    retries=0,
                    resume=False,
                )

            self.assertEqual(FakeClient.seen_queries, ["测试一下 web fetch 工具(get)"])
            self.assertEqual(stats["new_context_queries"], 1)
            self.assertEqual(stats["secondary_candidate_queries"], 1)
            self.assertEqual(stats["secondary_classified_queries"], 1)
            self.assertEqual(stats["unclassified_queries"], 2)

            refined = MODULE.load_scene_classifications(output_path, rules)

            def result_for(query):
                _, query_key = MODULE.normalize_query(query)
                return refined.classifications[MODULE.query_key_hash(query_key)]

            self.assertEqual(result_for("继续任务").method, "上下文依赖")
            self.assertEqual(
                result_for("测试一下 web fetch 工具(get)").method,
                "大模型二次分类",
            )
            self.assertEqual(
                result_for("测试一下 web fetch 工具(get)").category,
                "测试、评审与质量",
            )
            self.assertEqual(result_for("你好").category, "未分类")
            self.assertEqual(result_for("请修复登录错误").category, "问题排查与修复")
            self.assertEqual(
                refined.classification_version,
                MODULE.REFINED_CLASSIFICATION_VERSION,
            )
            self.assertEqual(
                refined.secondary_prompt_version,
                MODULE.SECONDARY_SCENE_PROMPT_VERSION,
            )
            self.assertEqual(refined.secondary_min_query_length, 8)
            self.assertFalse(checkpoint_path.exists())

    def test_percentile_uses_linear_interpolation(self) -> None:
        self.assertEqual(MODULE.linear_percentile([], 0.9), None)
        self.assertAlmostEqual(MODULE.linear_percentile([1, 2, 3, 4], 0.9), 3.7)

    def test_pit_pagination_is_sequential_and_uses_search_after(self) -> None:
        class FakeClient:
            def __init__(self):
                self.calls = []
                self.search_count = 0

            def request_json(self, method, path, *, body=None, params=None):
                self.calls.append((method, path, body, params))
                if path.endswith("/_pit") and method == "POST":
                    return {"id": "pit-1"}
                if path == "/_search":
                    self.search_count += 1
                    if self.search_count == 1:
                        return {
                            "hits": {
                                "hits": [
                                    {"_id": "1", "sort": [1, 1]},
                                    {"_id": "2", "sort": [2, 2]},
                                ]
                            }
                        }
                    return {"hits": {"hits": []}}
                if path == "/_pit" and method == "DELETE":
                    return {"succeeded": True}
                raise AssertionError((method, path))

        client = FakeClient()
        hits = list(
            MODULE.iter_matching_hits(
                client,
                index=MODULE.DEFAULT_INDEX,
                query={"match_all": {}},
                time_field="@timestamp",
                page_size=2,
                request_interval=0.0,
                pit_keep_alive="2m",
                max_records=10,
            )
        )

        self.assertEqual([hit["_id"] for hit in hits], ["1", "2"])
        search_calls = [call for call in client.calls if call[1] == "/_search"]
        self.assertEqual(len(search_calls), 2)
        self.assertNotIn("search_after", search_calls[0][2])
        self.assertEqual(search_calls[1][2]["search_after"], [2, 2])
        self.assertEqual(client.calls[-1][0:2], ("DELETE", "/_pit"))

    def test_pit_falls_back_to_opensearch_compatible_api(self) -> None:
        class FakeClient:
            def __init__(self):
                self.calls = []

            def request_json(self, method, path, *, body=None, params=None):
                self.calls.append((method, path, body, params))
                if path.endswith("/_pit") and method == "POST":
                    raise RuntimeError("unsupported Elasticsearch PIT API")
                if path.endswith("/_search/point_in_time") and method == "POST":
                    return {"pit_id": "pit-open-search"}
                if path == "/_search":
                    return {"hits": {"hits": []}}
                if path == "/_search/point_in_time" and method == "DELETE":
                    return {"succeeded": 1, "failed": 0}
                raise AssertionError((method, path, body, params))

        client = FakeClient()
        hits = list(
            MODULE.iter_matching_hits(
                client,
                index="prod-*",
                query={"match_all": {}},
                time_field="@timestamp",
                page_size=100,
                request_interval=0.05,
                pit_keep_alive="2m",
                max_records=1_000,
            )
        )

        self.assertEqual(hits, [])
        search_call = next(call for call in client.calls if call[1] == "/_search")
        self.assertEqual(
            search_call[2]["sort"],
            [{"@timestamp": "asc"}, {"_index": "asc"}, {"_id": "asc"}],
        )
        self.assertIn(
            (
                "DELETE",
                "/_search/point_in_time",
                {"pit_id": "pit-open-search"},
                None,
            ),
            client.calls,
        )

    def test_analysis_outputs_must_use_distinct_paths(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_dir:
            directory = Path(temporary_dir)
            parquet_path = directory / "normalized.parquet"
            output_path = directory / "output.xlsx"
            with self.assertRaisesRegex(ValueError, "路径必须不同"):
                MODULE.analyze_to_workbooks(
                    parquet_path,
                    output_path,
                    output_path,
                    MODULE.DEFAULT_CATEGORY_RULES,
                )

    def test_output_suffix_validation(self) -> None:
        with self.assertRaisesRegex(ValueError, "必须使用 .xlsx"):
            MODULE.ensure_suffix(Path("result.xls"), ".xlsx", "结果路径")
        with self.assertRaisesRegex(ValueError, "必须使用 .jsonl.gz"):
            MODULE.ensure_suffix(Path("result.jsonl"), ".jsonl.gz", "快照路径")

    def test_excel_text_is_formula_safe_and_reports_truncation(self) -> None:
        value, truncated = MODULE.safe_excel_text("=" + "x" * 40_000)

        self.assertTrue(truncated)
        self.assertTrue(value.startswith("'="))
        self.assertLessEqual(len(value), MODULE.EXCEL_CELL_LIMIT)
        self.assertTrue(value.endswith("…[Excel单元格截断]"))

    def test_query_redaction_and_normalization(self) -> None:
        display, key = MODULE.normalize_query("  Hello\n  WORLD  ")
        self.assertEqual(display, "Hello WORLD")
        self.assertEqual(key, "hello world")
        redacted, labels = MODULE.redact_query(
            "email=a@example.com password=secret 10.1.2.3 /Users/alice/project"
        )
        self.assertNotIn("a@example.com", redacted)
        self.assertNotIn("secret", redacted)
        self.assertNotIn("10.1.2.3", redacted)
        self.assertNotIn("alice", redacted)
        self.assertGreaterEqual(len(labels), 4)

    def test_model_token_comparison_uses_paired_weighted_cache_rate(self) -> None:
        analysis = MODULE.DatasetAnalysis(
            total_rows=0, timezone_name="Asia/Shanghai", schema_version="test"
        )
        records = [
            dict(input_tokens=100, cache_read_tokens=80, output_tokens=20, reasoning_tokens=5, total_tokens=120),
            dict(input_tokens=900, cache_read_tokens=90, output_tokens=100, reasoning_tokens=50, total_tokens=1000),
            dict(input_tokens=0, cache_read_tokens=0, output_tokens=0, reasoning_tokens=0, total_tokens=0),
            dict(input_tokens=None, cache_read_tokens=9, output_tokens=7, reasoning_tokens=None, total_tokens=7),
            dict(input_tokens=10, cache_read_tokens=11, output_tokens=4, reasoning_tokens=5, total_tokens=14),
            dict(input_tokens=float("inf"), cache_read_tokens=-1, output_tokens=float("nan"), reasoning_tokens=-1, total_tokens=-1),
        ]
        for record in records:
            record["usage_scope"] = "turn_total"
            MODULE.update_group(analysis.model_metrics["model-a"], record)
            MODULE.update_model_tokens(analysis.model_token_metrics["model-a"], record)
            analysis.total_rows += 1
        MODULE.update_group(analysis.model_metrics["model-b"], {})
        MODULE.update_model_tokens(analysis.model_token_metrics["model-b"], {})
        analysis.total_rows += 1
        metrics = analysis.model_token_metrics["model-a"]
        self.assertEqual(metrics.cache_exceeds_input_count, 1)
        self.assertEqual(metrics.reasoning_exceeds_output_count, 1)
        self.assertEqual(metrics.non_turn_usage_count, 0)
        layout = MODULE.model_token_layout(analysis)
        self.assertEqual(len(layout["summary_rows"]), 2)
        self.assertEqual(len(layout["detail_rows"]), 12)
        cache_row = layout["detail_rows"][5]
        self.assertEqual(cache_row[2], "缓存命中率")
        self.assertEqual(cache_row[3], 2)
        self.assertAlmostEqual(cache_row[4], 2 / 6)
        self.assertAlmostEqual(cache_row[5], 170 / 1000)
        self.assertAlmostEqual(cache_row[6], 0.45)
        self.assertNotAlmostEqual(cache_row[5], cache_row[6])
        self.assertEqual(layout["detail_rows"][1][3], 4)
        self.assertEqual(layout["detail_rows"][1][5], 252.5)
        self.assertEqual(layout["detail_rows"][3][5], 15)
        self.assertTrue(all(row[5] is None for row in layout["detail_rows"][6:]))
        self.assertIn("缓存输入>输入：1条", layout["summary_rows"][0][-1])
        self.assertIn("推理Token>输出Token：1条", layout["summary_rows"][0][-1])
        with tempfile.TemporaryDirectory() as temporary_dir:
            path = Path(temporary_dir) / "tokens.xlsx"
            workbook = MODULE.Workbook()
            MODULE.write_model_token_sheet(workbook.active, analysis)
            workbook.save(path)
            saved = load_workbook(path, data_only=False)
            try:
                sheet = saved.active
                self.assertEqual(sheet.freeze_panes, "C5")
                self.assertEqual(len(sheet.tables), 2)
                self.assertIn('SUMIFS($D$11:$D$22', sheet["H5"].value)
                self.assertIn('$A5,$C$11:$C$22,"缓存命中率"', sheet["H5"].value)
                self.assertEqual(sheet["H5"].number_format, "0.0%")
                self.assertAlmostEqual(sheet["F16"].value, 0.17)
                self.assertEqual(sheet["F16"].number_format, "0.0%")
                self.assertEqual(sheet["A6"].value, "model-b")
                self.assertIn('$A6,$C$11:$C$22,"总Token"', sheet["C6"].value)
            finally:
                saved.close()

    def test_model_averages_exclude_missing_and_preserve_zero(self) -> None:
        analysis = MODULE.DatasetAnalysis(
            total_rows=0, timezone_name="Asia/Shanghai", schema_version="test"
        )
        category_version, rules = MODULE.load_category_rules(
            MODULE.DEFAULT_CATEGORY_RULES
        )
        for model, values in {
            "model-a": [100, 201, 0, None, -1, float("nan"), float("inf")],
            "model-b": [None, -1],
            "model-c": [0],
        }.items():
            for value in values:
                duration = value / 100 if value is not None else None
                if model == "model-a" and value is None:
                    duration = 6
                record = {
                    "total_tokens": value,
                    "duration_seconds": duration,
                }
                MODULE.update_group(analysis.model_metrics[model], record)
                MODULE.update_group(analysis.category_metrics[rules[-1].name], record)
                MODULE.update_group(
                    analysis.classification_method_metrics["规则未命中"], record
                )
                MODULE.update_group(
                    analysis.length_metrics[MODULE.query_length_band(0)], record
                )
                for product, source_field, _ in MODULE.VERSION_DIMENSIONS:
                    MODULE.update_group(
                        analysis.version_metrics[product, source_field, "未记录"], record
                    )
                analysis.total_rows += 1
        with tempfile.TemporaryDirectory() as temporary_dir:
            path = Path(temporary_dir) / "analysis.xlsx"
            MODULE.write_analysis_workbook(
                analysis,
                path,
                input_path=Path("test.parquet"),
                category_version=category_version,
                rules=rules,
            )
            workbook = load_workbook(path, read_only=False, data_only=True)
            try:
                sheet = workbook["模型分布"]
                self.assertEqual(sheet.max_column, 14)
                self.assertEqual(sheet.auto_filter.ref, "A4:N7")
                headers = [cell.value for cell in sheet[4]]
                self.assertEqual(headers, [
                    "维度值", "任务数", "占比", "用户数", "会话数", "有效耗时数",
                    "平均耗时", "耗时P50", "耗时P90", "有效总Token记录数",
                    "平均总Token", "总Token P50", "总Token P90", "Turn成功率",
                ])
                self.assertNotIn("总Token合计", headers)
                rows = {
                    row[0]: dict(zip(headers, row))
                    for row in sheet.iter_rows(min_row=5, values_only=True)
                }
                self.assertEqual(rows["model-a"]["有效总Token记录数"], 3)
                self.assertEqual(rows["model-a"]["有效耗时数"], 4)
                self.assertAlmostEqual(rows["model-a"]["平均耗时"], 9.01 / 4)
                self.assertAlmostEqual(rows["model-a"]["平均总Token"], 301 / 3)
                self.assertEqual(rows["model-a"]["总Token P50"], 100)
                self.assertAlmostEqual(rows["model-a"]["总Token P90"], 180.8)
                self.assertIsNone(rows["model-b"]["平均耗时"])
                self.assertIsNone(rows["model-b"]["平均总Token"])
                self.assertEqual(rows["model-c"]["平均耗时"], 0)
                self.assertEqual(rows["model-c"]["平均总Token"], 0)
                self.assertEqual(sheet["G5"].number_format, '0.000" 秒"')
                self.assertEqual(sheet["K5"].number_format, "#,##0.00")
                self.assertEqual(sheet["N5"].number_format, "0.0%")
                self.assertEqual(workbook["场景分类"].max_column, 13)
                self.assertEqual(workbook["版本分布"].max_column, 14)
                self.assertEqual(workbook["推理强度"].max_column, 12)
            finally:
                workbook.close()

    def test_snapshot_parquet_and_workbooks_reconcile(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_dir:
            directory = Path(temporary_dir)
            snapshot_path = directory / "snapshot.jsonl.gz"
            parquet_path = directory / "normalized.parquet"
            classifications_path = directory / "classifications.parquet"
            analysis_path = directory / "analysis.xlsx"
            candidates_path = directory / "candidates.xlsx"
            window = MODULE.TimeWindow(
                datetime(2026, 9, 1, tzinfo=timezone.utc),
                datetime(2026, 9, 2, tzinfo=timezone.utc),
            )
            hits = [
                self._make_hit(
                    "doc-1",
                    "trace-1",
                    "span-1",
                    "请修复登录接口，联系 a@example.com",
                    "model-a",
                    "high",
                    1_000,
                    100,
                ),
                self._make_hit(
                    "doc-2",
                    "trace-2",
                    "span-2",
                    '=HYPERLINK("https://example.com")',
                    "model-b",
                    None,
                    3_000,
                    300,
                ),
                self._make_hit(
                    "doc-3",
                    "trace-2",
                    "span-2",
                    "这条是重复日志",
                    "model-b",
                    None,
                    3_000,
                    300,
                ),
            ]
            with patch.dict("os.environ", {"ASTRONCODE_HASH_SALT": "unit-test-only"}):
                fetch_stats = MODULE.write_dataset(
                    hits,
                    MODULE.DatasetPaths(snapshot_path, parquet_path),
                    requested=window,
                    effective=window,
                    estimated_count=3,
                    fallback_applied=False,
                    index=MODULE.DEFAULT_INDEX,
                    event_field="event",
                    event_value="agent.turn",
                    timezone_name="Asia/Shanghai",
                )
            category_version, _ = MODULE.load_category_rules(
                MODULE.DEFAULT_CATEGORY_RULES
            )
            semantic_query = '=HYPERLINK("https://example.com")'
            classifications = {}
            for query, category, method, confidence in (
                (
                    "请修复登录接口，联系 a@example.com",
                    "问题排查与修复",
                    "规则",
                    1.0,
                ),
                (
                    semantic_query,
                    "通用代码与功能开发",
                    "大模型",
                    0.91,
                ),
            ):
                _, query_key = MODULE.normalize_query(query)
                query_hash = MODULE.query_key_hash(query_key)
                classifications[query_hash] = MODULE.SceneClassification(
                    query_hash=query_hash,
                    category=category,
                    proposed_category=category,
                    method=method,
                    confidence=confidence,
                    reason="测试分类",
                    model="xopglm52" if method == "大模型" else "",
                    prompt_version=MODULE.SCENE_PROMPT_VERSION,
                    rule_version=category_version,
                    query_length=len(query),
                    input_truncated=False,
                )
            MODULE.write_scene_classification_parquet(
                classifications,
                classifications_path,
                rule_version=category_version,
                model="xopglm52",
                confidence_threshold=0.70,
            )
            analysis_stats = MODULE.analyze_to_workbooks(
                parquet_path,
                analysis_path,
                candidates_path,
                MODULE.DEFAULT_CATEGORY_RULES,
                classifications_path,
            )

            self.assertEqual(fetch_stats.fetched_hits, 3)
            self.assertEqual(fetch_stats.exported_rows, 2)
            self.assertEqual(fetch_stats.duplicate_rows, 1)
            self.assertEqual(analysis_stats.total_rows, 2)
            self.assertEqual(len(analysis_stats.candidates), 2)
            self.assertEqual(
                sum(
                    metric.count for metric in analysis_stats.category_metrics.values()
                ),
                2,
            )
            self.assertEqual(
                sum(metric.count for metric in analysis_stats.length_metrics.values()),
                2,
            )
            self.assertEqual(
                sum(metric.count for metric in analysis_stats.model_metrics.values()),
                2,
            )
            self.assertEqual(
                analysis_stats.model_token_metrics["model-a"].values["input_tokens"],
                [90],
            )
            self.assertEqual(
                analysis_stats.model_token_metrics["model-b"].values["output_tokens"],
                [10],
            )
            self.assertEqual(
                sum(
                    metric.count
                    for metric in analysis_stats.classification_method_metrics.values()
                ),
                2,
            )
            for product, source_field, _ in MODULE.VERSION_DIMENSIONS:
                self.assertEqual(
                    sum(
                        metric.count
                        for key, metric in analysis_stats.version_metrics.items()
                        if key[0] == product and key[1] == source_field
                    ),
                    2,
                )

            table = pq.read_table(parquet_path)
            self.assertEqual(table.num_rows, 2)
            self.assertNotIn("用户ID", table.schema.names)
            first = table.to_pylist()[0]
            self.assertEqual(len(first["user_hash"]), 64)
            self.assertNotEqual(first["user_hash"], "user-1")
            self.assertEqual(first["total_tokens"], 100)

            workbook = load_workbook(analysis_path, read_only=True, data_only=True)
            try:
                self.assertEqual(
                    workbook.sheetnames,
                    [
                        "分析总览",
                        "场景分类",
                        "分类方法",
                        "Query长度",
                        "耗时分布",
                        "模型分布",
                        "版本分布",
                        "推理强度",
                        "Token分布",
                        "状态分布",
                        "日期趋势",
                        "数据质量",
                        "口径说明",
                    ],
                )
                version_rows = list(
                    workbook["版本分布"].iter_rows(
                        min_row=5, min_col=1, max_col=4, values_only=True
                    )
                )
                self.assertIn(
                    ("AstronStudio", "astron.desktop.version", "1.2.3", 2),
                    version_rows,
                )
                self.assertIn(
                    ("AstronCode", "acode.cli_version", "0.0.42", 2),
                    version_rows,
                )
                classification_rows = list(
                    workbook["分类方法"].iter_rows(
                        min_row=5, min_col=1, max_col=2, values_only=True
                    )
                )
                self.assertIn(("规则", 1), classification_rows)
                self.assertIn(("大模型", 1), classification_rows)
            finally:
                workbook.close()

            workbook = load_workbook(candidates_path, read_only=True, data_only=False)
            try:
                candidate_sheet = workbook["候选Query"]
                headers = [
                    cell.value
                    for cell in next(candidate_sheet.iter_rows(min_row=4, max_row=4))
                ]
                self.assertEqual(headers, MODULE.CANDIDATE_HEADERS)
                classification_info = {
                    row[0]: row[1:4]
                    for row in candidate_sheet.iter_rows(
                        min_row=5,
                        min_col=5,
                        max_col=8,
                        values_only=True,
                    )
                }
                self.assertEqual(
                    classification_info["通用代码与功能开发"],
                    ("大模型", 0.91, "xopglm52"),
                )
                queries = [
                    row[0]
                    for row in candidate_sheet.iter_rows(
                        min_row=5, min_col=2, max_col=2, values_only=True
                    )
                ]
                self.assertTrue(any("已脱敏邮箱" in query for query in queries))
                self.assertTrue(any(query.startswith("'=") for query in queries))
            finally:
                workbook.close()

    @staticmethod
    def _make_hit(
        document_id,
        trace_id,
        span_id,
        query,
        model,
        effort,
        duration_ms,
        total_tokens,
    ):
        attributes = {
            "gen_ai.input.messages": [
                {"role": "user", "parts": [{"type": "text", "content": query}]}
            ],
            "gen_ai.request.model": model,
            "gen_ai.usage.total_tokens": total_tokens,
            "gen_ai.usage.input_tokens": total_tokens - 10,
            "gen_ai.usage.output_tokens": 10,
            "turn.duration_ms": duration_ms,
            "turn.status": "completed",
            "acode.uid": "user-1",
            "session.id": "session-1",
            "host.name": "workstation-1",
            "acode.cwd": "/Users/alice/project",
            "astron.desktop.version": "1.2.3",
            "acode.cli_version": "0.0.42",
        }
        if effort is not None:
            attributes["gen_ai.request.reasoning_effort"] = effort
        return {
            "_index": "prod-1",
            "_id": document_id,
            "_source": {
                "@timestamp": "2026-09-01T08:00:00Z",
                "event": "agent.turn",
                "Attributes": json.dumps(attributes, ensure_ascii=False),
                "TraceID": trace_id,
                "SpanID": span_id,
            },
        }


if __name__ == "__main__":
    unittest.main()
