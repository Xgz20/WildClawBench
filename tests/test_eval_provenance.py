from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook

from src.utils.eval_provenance import build_task_provenance, write_provenance_file


ROOT_DIR = Path(__file__).resolve().parents[1]
REPORT_SCRIPT = ROOT_DIR / "tools/report/scripts/generate_eval_report.py"


class EvalProvenanceHashTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.task_file = self.root / "task.md"
        self.task_file.write_text("# task\n", encoding="utf-8")
        self.workspace = self.root / "workspace"
        for component in ("exec", "tmp", "gt", "eval"):
            component_dir = self.workspace / component
            component_dir.mkdir(parents=True)
            (component_dir / "input.txt").write_text(component, encoding="utf-8")
        self.skills_root = self.root / "skills"
        skill_dir = self.skills_root / "sample-skill"
        (skill_dir / "scripts").mkdir(parents=True)
        (skill_dir / "SKILL.md").write_text("# sample skill\n", encoding="utf-8")
        (skill_dir / "scripts/helper.sh").write_text("echo ok\n", encoding="utf-8")
        self.task = {
            "task_id": "task_001",
            "file_path": str(self.task_file),
            "workspace_path": str(self.workspace),
            "skills_path": str(self.skills_root),
            "prompt": "Build the artifact",
            "env": "export MODE=test",
            "skills": "sample-skill",
            "warmup": "echo warmup",
            "timeout_seconds": 300,
            "automated_checks": "def grade(**kwargs): return {'overall_score': 1}",
            "llm_judge_rubric": "Judge quality",
            "rubric_criteria": [{"key": "quality", "weight": 1.0}],
            "grading_type": "hybrid",
            "grading_weights": {"automated": 0.5, "llm_judge": 0.5},
            "metric_profile": "",
            "judge_evidence": {},
        }

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_hashes_are_stable_and_written_as_metadata(self) -> None:
        first = build_task_provenance(self.task)
        second = build_task_provenance(dict(self.task))
        self.assertEqual(first, second)
        for key in (
            "task_sha256",
            "execution_contract_sha256",
            "scoring_contract_sha256",
            "workspace_exec_sha256",
            "workspace_tmp_sha256",
            "skill_bundles_sha256",
            "ground_truth_sha256",
            "workspace_eval_sha256",
        ):
            self.assertRegex(first[key], r"^[0-9a-f]{64}$")

        output = self.root / "run_001/provenance.json"
        write_provenance_file(output, first)
        self.assertEqual(json.loads(output.read_text(encoding="utf-8")), first)
        self.assertEqual(list(output.parent.glob(".*.tmp")), [])

    def test_task_source_change_is_informational_when_contracts_do_not_change(self) -> None:
        before = build_task_provenance(self.task)
        self.task_file.write_text("# task\n\n", encoding="utf-8")
        after = build_task_provenance(self.task)

        self.assertNotEqual(before["task_sha256"], after["task_sha256"])
        self.assertEqual(
            before["execution_contract_sha256"],
            after["execution_contract_sha256"],
        )
        self.assertEqual(
            before["scoring_contract_sha256"],
            after["scoring_contract_sha256"],
        )

    def test_execution_inputs_only_change_execution_contract(self) -> None:
        mutations = (
            (
                lambda task: task.update(prompt="Changed prompt"),
                lambda task: task.update(prompt="Build the artifact"),
            ),
            (
                lambda task: (self.workspace / "exec/input.txt").write_text(
                    "changed exec", encoding="utf-8"
                ),
                lambda task: (self.workspace / "exec/input.txt").write_text(
                    "exec", encoding="utf-8"
                ),
            ),
            (
                lambda task: (self.workspace / "tmp/input.txt").write_text(
                    "changed tmp", encoding="utf-8"
                ),
                lambda task: (self.workspace / "tmp/input.txt").write_text(
                    "tmp", encoding="utf-8"
                ),
            ),
            (
                lambda task: (
                    self.skills_root / "sample-skill/scripts/helper.sh"
                ).write_text("echo changed\n", encoding="utf-8"),
                lambda task: (
                    self.skills_root / "sample-skill/scripts/helper.sh"
                ).write_text("echo ok\n", encoding="utf-8"),
            ),
        )
        for mutate, restore in mutations:
            with self.subTest(mutation=mutate):
                before = build_task_provenance(self.task)
                mutate(self.task)
                after = build_task_provenance(self.task)
                self.assertNotEqual(
                    before["execution_contract_sha256"],
                    after["execution_contract_sha256"],
                )
                self.assertEqual(
                    before["scoring_contract_sha256"],
                    after["scoring_contract_sha256"],
                )
                restore(self.task)
                self.assertEqual(build_task_provenance(self.task), before)

    def test_scoring_inputs_only_change_scoring_contract(self) -> None:
        mutations = (
            (
                lambda task: task.update(automated_checks="def grade(): return 0"),
                lambda task: task.update(
                    automated_checks=(
                        "def grade(**kwargs): return {'overall_score': 1}"
                    )
                ),
            ),
            (
                lambda task: (self.workspace / "gt/input.txt").write_text(
                    "changed gt", encoding="utf-8"
                ),
                lambda task: (self.workspace / "gt/input.txt").write_text(
                    "gt", encoding="utf-8"
                ),
            ),
            (
                lambda task: (self.workspace / "eval/input.txt").write_text(
                    "changed eval", encoding="utf-8"
                ),
                lambda task: (self.workspace / "eval/input.txt").write_text(
                    "eval", encoding="utf-8"
                ),
            ),
            (
                lambda task: task.update(judge_evidence={
                    "required": [{"path": "results/answer.md", "role": "deliverable"}],
                    "references": [],
                }),
                lambda task: task.update(judge_evidence={}),
            ),
        )
        for mutate, restore in mutations:
            with self.subTest(mutation=mutate):
                before = build_task_provenance(self.task)
                mutate(self.task)
                after = build_task_provenance(self.task)
                self.assertEqual(
                    before["execution_contract_sha256"],
                    after["execution_contract_sha256"],
                )
                self.assertNotEqual(
                    before["scoring_contract_sha256"],
                    after["scoring_contract_sha256"],
                )
                restore(self.task)
                self.assertEqual(build_task_provenance(self.task), before)

    def test_symlink_target_content_is_not_followed(self) -> None:
        external = self.root / "external.txt"
        external.write_text("first", encoding="utf-8")
        (self.workspace / "exec/link.txt").symlink_to(external)
        before = build_task_provenance(self.task)
        external.write_text("second", encoding="utf-8")
        after = build_task_provenance(self.task)
        self.assertEqual(
            before["execution_contract_sha256"],
            after["execution_contract_sha256"],
        )


class EvalProvenanceReportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        import importlib.util

        spec = importlib.util.spec_from_file_location("provenance_report", REPORT_SCRIPT)
        cls.report = importlib.util.module_from_spec(spec)
        assert spec.loader is not None
        sys.modules["provenance_report"] = cls.report
        spec.loader.exec_module(cls.report)

    @staticmethod
    def _unit(unit_id: str, task_id: str, observations: list[dict]):
        task = SimpleNamespace(run_provenance=observations)
        return SimpleNamespace(unit=unit_id, task_map={task_id: task})

    def test_compatibility_uses_split_contract_hashes(self) -> None:
        base = {
            "run_dir": "run_001",
            "provenance_status": "complete",
            "task_sha256": "task-a",
            "execution_contract_sha256": "exec-a",
            "scoring_contract_sha256": "score-a",
        }
        cases = (
            ("consistent", dict(base)),
            ("task_source_only_changed", {**base, "task_sha256": "task-b"}),
            ("execution_mismatch", {**base, "execution_contract_sha256": "exec-b"}),
            ("scoring_mismatch", {**base, "scoring_contract_sha256": "score-b"}),
            (
                "execution_and_scoring_mismatch",
                {
                    **base,
                    "execution_contract_sha256": "exec-b",
                    "scoring_contract_sha256": "score-b",
                },
            ),
        )
        for expected, second in cases:
            with self.subTest(expected=expected):
                result = self.report.build_provenance_consistency([
                    self._unit("model-a@harness", "task_001", [base]),
                    self._unit("model-b@harness", "task_001", [second]),
                ])
                self.assertEqual(result["items"][0]["status"], expected)
                self.assertEqual(result["mode"], "informational_non_blocking")
                self.assertTrue(result["report_scores_unchanged"])
                self.assertEqual(
                    result["compatibility_hash_fields"],
                    ["execution_contract_sha256", "scoring_contract_sha256"],
                )
                self.assertTrue(result["task_sha256_informational_only"])

    def test_legacy_results_still_generate_report_without_changing_scores(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "round"
            run_dir = root / "model-x/harness-y/01_suite/task_001/run_001"
            run_dir.mkdir(parents=True)
            (run_dir / "score.json").write_text(
                json.dumps({"overall_score": 0.5, "check": 0.5}), encoding="utf-8"
            )
            (run_dir / "execution_status.json").write_text(
                json.dumps({"status": "completed"}), encoding="utf-8"
            )
            (run_dir / "usage.json").write_text(
                json.dumps({"request_count": 1, "total_tokens": 10}), encoding="utf-8"
            )
            output_dir = Path(tmp) / "output"

            completed = subprocess.run(
                [
                    sys.executable,
                    str(REPORT_SCRIPT),
                    "--result-root",
                    str(root / "model-x/harness-y"),
                    "--tasks-dir",
                    str(Path(tmp) / "missing-tasks"),
                    "--output-dir",
                    str(output_dir),
                    "--emit",
                    "summary_json,md,html",
                ],
                check=True,
                capture_output=True,
                text=True,
            )

            self.assertIn("非阻断", completed.stdout)
            self.assertFalse((run_dir / "provenance.json").exists())
            workbook_path = next(output_dir.glob("report_1units_*.xlsx"))
            workbook = load_workbook(workbook_path, read_only=True, data_only=True)
            self.assertIn("评测契约一致性", workbook.sheetnames)
            self.assertLess(
                workbook.sheetnames.index("评测契约一致性"),
                workbook.sheetnames.index("评分详情_model-x@harness-y"),
            )
            overview = workbook["总览"]
            overview_headers = [cell.value for cell in overview[1]]
            total_score_column = overview_headers.index("总平均分") + 1
            self.assertEqual(overview.cell(2, total_score_column).value, 50)
            provenance_sheet = workbook["评测契约一致性"]
            self.assertEqual(provenance_sheet["B1"].value.split("；", 1)[0], "非阻断检查")
            self.assertEqual(provenance_sheet["B3"].value, "legacy_missing")

            summary_path = next(output_dir.glob("*.summary.json"))
            summary = json.loads(summary_path.read_text(encoding="utf-8"))
            self.assertEqual(summary["overview"]["average_score"], 0.5)
            self.assertEqual(
                summary["provenance_consistency"]["items"][0]["status"],
                "legacy_missing",
            )
            markdown = next(output_dir.glob("*.md")).read_text(encoding="utf-8")
            html = next(output_dir.glob("*.html")).read_text(encoding="utf-8")
            self.assertIn("评测契约一致性（非阻断）", markdown)
            self.assertIn("task_sha256` 只用于完整任务追溯", markdown)
            self.assertIn("评测契约一致性（非阻断）", html)


if __name__ == "__main__":
    unittest.main()
